"""Cash & Expenses -- the cash book (owner, 25 Oct 2026): every cash movement of the day in one
numbered list in the order it happened -- 1. Opening balance, 2. Sale, 3. tea, ... -- with In,
Out and the running balance, the totals at the end, Counted and Difference once the day is
closed. The book is built from the register's own sections, so its totals are the register's
and a deleted bill leaves it. A date range numbers on across the days with a sub-total per day.

Temp store files only; no server.
"""
from __future__ import annotations

import base64
import os
import sqlite3
import tempfile
import unittest

from core import cash_register as cr
from core import db_setup

D = "2026-10-24"
D2 = "2026-10-25"


class CashBook(unittest.TestCase):
    def setUp(self):
        fd, self.path = tempfile.mkstemp(suffix=".db", prefix="cashbook_")
        os.close(fd)
        c = self.conn = sqlite3.connect(self.path)
        db_setup.initialise(c)
        self.cust = c.execute("INSERT INTO customers (name) VALUES ('RAM')").lastrowid
        self.sup = c.execute("INSERT INTO suppliers (name) VALUES ('TULJAI')").lastrowid
        c.commit()
        self.addCleanup(self._cleanup)

    def _cleanup(self):
        self.conn.close()
        for s in ("", "-wal", "-shm"):
            try:
                os.remove(self.path + s)
            except OSError:
                pass

    def sale(self, day, at, bill, cash, online=0.0, customer=None):
        sid = self.conn.execute(
            "INSERT INTO sales (bill_no, bill_date, customer_id, total_amount, amount_paid, cash_paid, online_paid, "
            "created_at) VALUES (?,?,?,?,?,?,?,?)",
            (bill, day, customer, cash + online, cash + online, cash, online, f"{day}T{at}:00")).lastrowid
        self.conn.commit()
        return sid

    def fill_day(self):
        c = self.conn
        cr.set_opening(c, D, 1000)
        self.sale(D, "09:05", "SB1", 300)                            # walk-in
        self.sale(D, "09:30", "SB2", 200, 100, customer=self.cust)   # only the 200 cash
        cr.save_expense(c, {"date": D, "description": "tea", "category": "Tea & snacks", "amount": 20, "at": "09:15"})
        cr.save_expense(c, {"date": D, "description": "nashta", "amount": 30, "at": "09:40"})
        cr.save_expense(c, {"date": D, "description": "mama la dile", "category": "Other", "amount": 250, "at": "10:00"})
        cr.save_expense(c, {"date": D, "category": "Transport", "amount": 15, "at": "10:05"})  # nothing typed
        c.execute("INSERT INTO customer_payments (customer_id, payment_date, amount, payment_mode, cash_amount, "
                  "online_amount, created_at) VALUES (?,?,100,'cash',100,0,?)", (self.cust, D, f"{D}T10:10:00"))
        c.execute("INSERT INTO purchases (purchase_no, supplier_id, purchase_date, bill_number, total_amount, "
                  "final_amount, cash_paid_at_entry, online_paid_at_entry, amount_paid_at_entry, created_at) "
                  "VALUES ('P1',?,?,'T-77',900,900,500,0,500,?)", (self.sup, D, f"{D}T10:20:00"))
        c.commit()
        self.gone = self.sale(D, "10:30", "SB3", 150)
        c.execute("INSERT INTO supplier_payments (payment_no, supplier_id, payment_date, amount, mode, created_at) "
                  "VALUES ('PAY1',?,?,200,'Cash',?)", (self.sup, D, f"{D}T11:00:00"))
        c.execute("INSERT INTO sales_returns (return_no, customer_id, return_date, refund_amount, settle_mode) "
                  "VALUES ('SR1',?,?,50,'cash')", (self.cust, D))
        c.execute("INSERT INTO customer_payments (customer_id, payment_date, amount, payment_mode, cash_amount, "
                  "online_amount, reference_no, note, created_at) VALUES (?,?,-50,'cash',-50,0,'SR1','Refund SR1',?)",
                  (self.cust, D, f"{D}T11:10:00"))
        c.execute("INSERT INTO supplier_payments (payment_no, supplier_id, payment_date, amount, mode, reference, "
                  "created_at) VALUES ('PAY2',?,?,-40,'Cash','PR1',?)", (self.sup, D, f"{D}T11:20:00"))
        c.commit()
        cr.save_expense(c, {"date": D, "kind": "in", "description": "Owner added cash", "amount": 500, "at": "11:30"})

    def test_one_numbered_list_in_time_order_with_running_balance(self):
        self.fill_day()
        rep = cr.day_report(self.conn, D)
        book = rep["book"]
        rows = book["rows"]
        self.assertEqual([r["sr"] for r in rows], list(range(1, len(rows) + 1)))
        self.assertEqual(
            [(r["time"], r["entry"]) for r in rows],
            [("", "Opening balance"),
             ("09:05", "Sale SB1"),
             ("09:15", "tea"),
             ("09:30", "Sale SB2 – RAM"),
             ("09:40", "nashta"),
             ("10:00", "mama la dile"),
             ("10:05", "Transport"),
             ("10:10", "Received RAM"),
             ("10:20", "Purchase payment TULJAI"),
             ("10:30", "Sale SB3"),
             ("11:00", "Purchase payment TULJAI"),
             ("11:10", "Refund SR1"),
             ("11:20", "Purchase return TULJAI"),
             ("11:30", "Owner added cash")])
        self.assertEqual([r["type"] for r in rows][1:4], ["Sale", "Expense", "Sale"])
        self.assertEqual(rows[8]["type"], "Purchase")
        self.assertEqual(rows[10]["type"], "Supplier payment")
        self.assertEqual(rows[13]["type"], "Cash added")
        self.assertEqual(rows[0]["balance"], 1000.0)
        self.assertEqual((rows[1]["cash_in"], rows[1]["cash_out"], rows[1]["balance"]), (300.0, None, 1300.0))
        self.assertEqual((rows[2]["cash_in"], rows[2]["cash_out"], rows[2]["balance"]), (None, 20.0, 1280.0))
        self.assertEqual(rows[3]["cash_in"], 200.0)                 # the online part is not cash
        # Each balance is the one before plus In less Out.
        for a, b in zip(rows, rows[1:]):
            self.assertAlmostEqual(b["balance"], a["balance"] + (b["cash_in"] or 0) - (b["cash_out"] or 0), places=2)
        # The totals are the register's, and the last balance is the expected cash.
        t = rep["totals"]
        self.assertEqual(book["total_in"], t["cash_in"])
        self.assertEqual(book["total_out"], round(t["cash_out"] + t["expenses"], 2))
        self.assertEqual(round(sum(r["cash_in"] or 0 for r in rows), 2), book["total_in"])
        self.assertEqual(round(sum(r["cash_out"] or 0 for r in rows), 2), book["total_out"])
        self.assertEqual(rows[-1]["balance"], t["expected"])
        self.assertEqual(book["expected"], t["expected"])
        lines = sum(len(s["lines"]) for s in rep["sections"]) + len(rep["expenses"]) + len(rep["cash_added"])
        self.assertEqual(len(rows), lines + 1)
        self.assertIsNone(book["counted"])

    def test_a_deleted_bill_drops_out_and_a_closed_day_shows_counted_and_difference(self):
        self.fill_day()
        before = cr.day_report(self.conn, D)["book"]
        self.conn.execute("UPDATE sales SET deleted=1 WHERE id=?", (self.gone,))
        self.conn.commit()
        rep = cr.day_report(self.conn, D)
        book = rep["book"]
        self.assertNotIn("Sale SB3", [r["entry"] for r in book["rows"]])
        self.assertEqual(len(book["rows"]), len(before["rows"]) - 1)
        self.assertEqual([r["sr"] for r in book["rows"]], list(range(1, len(book["rows"]) + 1)))
        self.assertEqual(book["expected"], round(before["expected"] - 150, 2))
        self.assertEqual(book["rows"][-1]["balance"], book["expected"])

        cr.close_day(self.conn, D, book["expected"] - 10, "short 10")
        rep = cr.day_report(self.conn, D)
        self.assertEqual(rep["book"]["counted"], rep["counted"])
        self.assertEqual(rep["book"]["difference"], -10.0)
        sec = cr._day_sections(rep)[0]
        self.assertEqual(sec["columns"], ["Sr", "Time", "Entry", "In", "Out", "Balance"])
        self.assertEqual(sec["rows"][0], [1, "", "Opening balance", "", "", 1000.0])
        labels = [r[2] for r in sec["rows"][len(rep["book"]["rows"]):]]
        self.assertEqual(labels[:4], ["Total", "Expected closing cash", "Counted cash", "Difference (counted - expected)"])
        tot = sec["rows"][len(rep["book"]["rows"])]
        self.assertEqual((tot[3], tot[4]), (rep["book"]["total_in"], rep["book"]["total_out"]))

    def test_a_range_numbers_on_across_days_with_a_sub_total_per_day(self):
        self.fill_day()
        cr.set_opening(self.conn, D2, 700)
        self.sale(D2, "09:00", "SB9", 80, customer=self.cust)
        cr.save_expense(self.conn, {"date": D2, "description": "tea", "amount": 10, "at": "08:30"})
        rng = cr.range_report(self.conn, "2026-10-20", "2026-10-26")
        book = rng["book"]
        self.assertEqual([d["date"] for d in book["days"]], [D, D2])
        day1, day2 = book["days"]
        self.assertEqual(day2["rows"][0]["sr"], day1["rows"][-1]["sr"] + 1)
        self.assertEqual([(r["entry"], r["balance"]) for r in day2["rows"]],
                         [("Opening balance", 700.0), ("tea", 690.0), ("Sale SB9 – RAM", 770.0)])
        for d, line in zip(book["days"], rng["days"]):
            self.assertEqual(d["expected"], line["expected"])
            self.assertEqual(d["total_in"], line["cash_in"])
            self.assertEqual(d["total_out"], round(line["cash_out"] + line["expenses"], 2))
        self.assertEqual(book["total_in"], round(rng["totals"]["cash_in"], 2))
        sec = cr._range_sections(rng)[0]
        self.assertEqual(sec["columns"], ["Sr", "Date", "Time", "Entry", "In", "Out", "Balance"])
        entries = [r[3] for r in sec["rows"]]
        self.assertIn(f"Total {cr._dmy(D)}", entries)
        self.assertLess(entries.index(f"Total {cr._dmy(D)}"), entries.index("Opening balance", 1))
        self.assertEqual(entries[-1], "Total of all days")
        srs = [r[0] for r in sec["rows"] if r[0] != ""]
        self.assertEqual(srs, list(range(1, len(srs) + 1)))

    def test_files_lead_with_the_cash_book(self):
        self.fill_day()
        out = cr.export(self.conn, {"date": D, "format": "csv"})
        self.assertTrue(out["ok"], out)
        text = base64.b64decode(out["content_base64"]).decode("utf-8-sig")
        first = text.splitlines()
        self.assertTrue(first[0].startswith("Cash book"))
        self.assertEqual(first[1], "Sr,Time,Entry,In,Out,Balance")
        self.assertIn("2,09:05,Sale SB1,300.0,,1300.0", text)
        out = cr.export(self.conn, {"date": D, "format": "xlsx"})
        from io import BytesIO

        from openpyxl import load_workbook

        wb = load_workbook(BytesIO(base64.b64decode(out["content_base64"])))
        self.assertEqual(wb.sheetnames, ["Cash book", "Details"])
        self.assertEqual([c.value for c in wb["Cash book"][4]], ["Sr", "Time", "Entry", "In", "Out", "Balance"])
        out = cr.export(self.conn, {"from": "2026-10-20", "to": "2026-10-26", "format": "csv"})
        self.assertTrue(out["ok"], out)
        self.assertIn("Sr,Date,Time,Entry,In,Out,Balance", base64.b64decode(out["content_base64"]).decode("utf-8-sig"))


if __name__ == "__main__":
    unittest.main()
