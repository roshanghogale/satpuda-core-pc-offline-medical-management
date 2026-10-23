"""Doctor margin by rules (owner, 22 + 23 Oct 2026).

core/doctor_margin_rules.py  the rules, version 2 (per type: none / on profit by groups -- only
                             the products in a group, a product's own group, batch-limited
                             items, an "All products" group -- / on total sales at a %), version
                             1 read and turned into 2, the per-line margin and the commission
core/doctor_margin_store.py  the doctor report's own saved lines, written when a bill is saved
core/doctor_margin_rules_report.py  the report (only lines with a margin), Excel / PDF

Every store here is built in the test; no real data.
"""
import io
import os
import sqlite3
import sys
import tempfile
import unittest
from unittest import mock

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from core import db_setup  # noqa: E402
from core import doctor_margin_rules as dmr  # noqa: E402
from core import doctor_margin_rules_report as rr  # noqa: E402
from core import doctor_margin_store as st  # noqa: E402

FROM, TO = "2026-10-11", "2026-10-20"


def fixed(v):
    return {"kind": "fixed", "value": v}


RULES = {
    "rate_basis": "purchase",
    "sets": [
        {"doctor": "", "types": [
            {"type": "Injection", "mode": "groups", "groups": [
                {"name": "Batch B1", "margin": fixed(30), "items": [{"product": "zifi 1gm", "batches": ["b1"]}]},
                {"name": "Zifi own", "margin": fixed(20), "items": [{"product": "ZIFI 1GM"}]},
                {"name": "Cef group", "margin": fixed(10),
                 "items": [{"product": "ZIFI 1GM"}, {"product": "LYCEFT 1GM"}, {"product": "XONE 1GM"}]},
                {"name": "Not these", "margin": {"kind": "none"}, "items": [{"product": "XONE 1GM"}]},
            ]},
            {"type": "Tablet", "mode": "type", "margin": {"kind": "mrp_rate"}},
            {"type": "Syrup", "mode": "none"},
        ]},
        {"doctor": "Dr B", "rate_basis": "purchase_free", "types": [{"type": "Injection", "mode": "none"}]},
    ],
}


class TheRules(unittest.TestCase):
    def setUp(self):
        self.res = dmr.Resolver(RULES)

    def kind_value(self, doctor, mtype, product, batch):
        got = self.res.resolve(doctor, mtype, product, batch)
        return None if got is None else (got["kind"], got["value"])

    def test_precedence_batch_product_group_all_none(self):
        self.assertEqual(self.kind_value("DR A", "Injection", "ZIFI 1GM", "B1"), ("fixed", 30))  # batch-limited
        self.assertEqual(self.kind_value("DR A", "Injection", "Zifi  1gm", "B2"), ("fixed", 20))  # its own group
        self.assertEqual(self.kind_value("DR A", "Injection", "LYCEFT 1GM", "L1"), ("fixed", 10))  # a group
        self.assertIsNone(self.kind_value("DR A", "Injection", "PIPZO 4.5", "P1"))  # in no group: not in the report
        self.assertEqual(self.kind_value("DR A", "TABLET", "ANY TAB", "T1"), ("mrp_rate", 0))  # v1 whole type
        self.assertIsNone(self.kind_value("DR A", "Syrup", "ANY SYRUP", "S1"))  # type: no margin
        self.assertIsNone(self.kind_value("DR A", "Drops", "ANY DROPS", "D1"))  # type not in the rules

    def test_a_no_margin_group_takes_products_out(self):
        # XONE is in "Cef group" (3 products) and alone in "Not these": the one-product group wins.
        self.assertIsNone(self.kind_value("DR A", "Injection", "XONE 1GM", "X1"))

    def test_ungrouped_products_give_nothing_even_with_a_v1_rest(self):
        r = {"sets": [{"doctor": "", "types": [{"type": "Injection", "mode": "groups",
                                                "rest": {"kind": "pct_mrp", "value": 5}, "groups": [
                                                    {"name": "G", "margin": fixed(5), "items": ["ZIFI 1GM"]}]}]}]}
        res = dmr.Resolver(r)
        self.assertIsNotNone(res.resolve("X", "Injection", "ZIFI 1GM", ""))
        self.assertIsNone(res.resolve("X", "Injection", "OTHER", ""))  # the v1 rest is gone
        self.assertNotIn("rest", dmr.clean_rules(r)["sets"][0]["types"][0])

    def test_an_all_products_group_is_below_every_picked_group(self):
        r = {"sets": [{"doctor": "", "types": [{"type": "Injection", "mode": "groups", "groups": [
            {"name": "Every injection", "all": True, "margin": {"kind": "pct_mrp", "value": 5}},
            {"name": "Two", "margin": fixed(9), "items": ["ZIFI 1GM", "XONE 1GM"]},
            {"name": "Not this one", "margin": {"kind": "none"}, "items": ["XONE 1GM"]}]}]}]}
        res = dmr.Resolver(r)
        self.assertEqual(res.resolve("X", "Injection", "ZIFI 1GM", "")["value"], 9)
        self.assertEqual(res.resolve("X", "Injection", "PIPZO", "")["kind"], "pct_mrp")
        self.assertIsNone(res.resolve("X", "Injection", "XONE 1GM", ""))
        self.assertEqual(dmr.clean_rules(r)["sets"][0]["types"][0]["groups"][0]["items"], [])

    def test_total_sales_way_and_its_percent(self):
        r = {"sets": [
            {"doctor": "", "sales_pct": 10, "types": [
                {"type": "Tablet", "mode": "sales"},
                {"type": "Injection", "mode": "groups", "groups": [{"name": "G", "margin": fixed(5), "items": ["ZIFI"]}]}]},
            {"doctor": "Dr B", "sales_pct": 12.5, "types": [{"type": "Injection", "mode": "sales"}]},
            {"doctor": "Dr C", "sales_pct": None, "types": [{"type": "Syrup", "mode": "sales"}]}]}
        res = dmr.Resolver(r)
        self.assertEqual(res.resolve("DR A", "tablet", "ANY", "")["way"], "sales")
        self.assertEqual(res.resolve("DR A", "tablet", "ANY", "")["value"], 10)
        self.assertEqual(res.resolve("DR A", "Injection", "ZIFI", "")["way"], "profit")
        self.assertEqual((res.resolve("DR B", "Injection", "ANYTHING", "")["value"],
                          res.resolve("DR B", "Tablet", "ANY", "")["value"]), (12.5, 12.5))  # his own %
        self.assertEqual(res.resolve("DR C", "Syrup", "S", "")["value"], 10)  # no % of his own: All doctors'
        self.assertIsNone(res.resolve("DR A", "Syrup", "S", ""))
        with self.assertRaises(ValueError):
            dmr.clean_rules({"sets": [{"doctor": "", "sales_pct": 101, "types": []}]})
        self.assertTrue(dmr.has_rules({"sets": [{"doctor": "", "types": [{"type": "Tablet", "mode": "sales"}]}]}))

    def test_version_1_is_read_as_version_2(self):
        v1 = {"version": 1, "rate_basis": "purchase_free", "updated": "2026-10-22T10:00:00", "sets": [
            {"doctor": "", "rate_basis": None, "types": [
                {"type": "Tablet", "mode": "type", "margin": {"kind": "pct_mrp", "value": 7}, "groups": [], "rest": None},
                {"type": "Injection", "mode": "groups", "margin": None, "rest": {"kind": "fixed", "value": 2},
                 "groups": [{"name": "Z", "margin": fixed(4), "items": [{"product": "ZIFI", "batches": []}]}]},
                {"type": "Syrup", "mode": "none", "margin": None, "groups": [], "rest": None}]},
            {"doctor": "DR B", "rate_basis": "purchase", "types": []}]}
        v2 = dmr.clean_rules(v1)
        self.assertEqual(v2, {"version": 2, "rate_basis": "purchase_free", "updated": "2026-10-22T10:00:00", "sets": [
            {"doctor": "", "rate_basis": None, "sales_pct": 0.0, "types": [
                {"type": "Tablet", "mode": "groups", "groups": [
                    {"name": "All Tablet", "margin": {"kind": "pct_mrp", "value": 7.0}, "all": True, "items": []}]},
                {"type": "Injection", "mode": "groups", "groups": [
                    {"name": "Z", "margin": {"kind": "fixed", "value": 4.0}, "all": False,
                     "items": [{"product": "ZIFI", "batches": []}]}]},
                {"type": "Syrup", "mode": "none", "groups": []}]},
            {"doctor": "DR B", "rate_basis": "purchase", "sales_pct": None, "types": []}]})
        self.assertEqual(dmr.clean_rules(v2), v2)  # version 2 stays as it is
        # the same figures as version 1 gave, except the rest of type
        sale = _sale(1, "Dr A", [
            {"medicine_id": 1, "name": "ANY TAB", "type": "Tablet", "batch": "T", "mrp_unit": 10, "qty": 3},
            {"medicine_id": 2, "name": "ZIFI", "type": "Injection", "batch": "Z", "mrp_unit": 50, "qty": 2},
            {"medicine_id": 3, "name": "OTHER INJ", "type": "Injection", "batch": "O", "mrp_unit": 50, "qty": 2}])
        rep = dmr.compute([sale], [], v1)
        self.assertEqual([(r["name"], r["value"]) for r in rep["rows"]], [("ANY TAB", 2.1), ("ZIFI", 8.0)])

    def test_a_doctors_own_rules_override_per_type(self):
        self.assertIsNone(self.kind_value("DR B", "Injection", "ZIFI 1GM", "B1"))  # his Injection: none
        self.assertEqual(self.kind_value("DR B", "Tablet", "ANY TAB", "T1"), ("mrp_rate", 0))  # follows default
        self.assertEqual(self.res.rate_basis("dr b"), "purchase_free")
        self.assertEqual(self.res.rate_basis("DR A"), "purchase")

    def test_cleaning(self):
        clean = dmr.clean_rules(RULES)
        self.assertEqual(clean["version"], 2)
        self.assertEqual(clean["sets"][0]["doctor"], "")
        self.assertEqual(clean["sets"][1]["doctor"], "DR B")
        g = clean["sets"][0]["types"][0]["groups"][0]
        self.assertEqual(g["items"], [{"product": "ZIFI 1GM", "batches": ["B1"]}])
        with self.assertRaises(ValueError):
            dmr.clean_rules({"sets": [{"doctor": "", "types": [
                {"type": "Tablet", "mode": "type", "margin": {"kind": "pct_mrp", "value": 140}}]}]})
        self.assertFalse(dmr.has_rules({}))
        self.assertTrue(dmr.has_rules(RULES))
        # a group with no products is dropped
        none = dmr.clean_rules({"sets": [{"doctor": "", "types": [
            {"type": "Injection", "mode": "groups", "groups": [{"name": "E", "margin": fixed(1), "items": []}]}]}]})
        self.assertEqual(none["sets"][0]["types"][0]["groups"], [])


class TheRate(unittest.TestCase):
    def test_free_included(self):
        # Rs 100 a unit-pack, 10 bought + 5 free: as billed Rs 100 a unit; free included 1000 / 15.
        r = dmr.batch_rates([{"medicine_id": 1, "qty": 10, "free_qty": 5, "rate": 100, "units": 1}])[1]
        self.assertAlmostEqual(r["purchase"], 100.0)
        self.assertAlmostEqual(r["purchase_free"], 1000 / 15)
        # the owner's words: "the rate is 100 for 10 items and 5 are free" -> 100 per 15 units
        r = dmr.batch_rates([{"medicine_id": 2, "qty": 10, "free_qty": 5, "rate": 10, "units": 1}])[2]
        self.assertEqual(round(r["purchase_free"], 4), 6.6667)

    def test_strips_per_tablet(self):
        # 2 strips of 10 + 1 free strip at Rs 50 a strip
        r = dmr.batch_rates([{"medicine_id": 3, "qty": 2, "free_qty": 1, "rate": 50, "units": 10}])[3]
        self.assertAlmostEqual(r["purchase"], 5.0)
        self.assertAlmostEqual(r["purchase_free"], 100 / 30)


def _sale(sid, doctor, items, returns=(), date="2026-10-12"):
    return {"id": sid, "bill_no": f"B{sid}", "date": date, "patient": f"PATIENT {sid}", "doctor": doctor,
            "items": [dict(id=sid * 100 + i, **it) for i, it in enumerate(items)], "returns": list(returns)}


class TheCalculation(unittest.TestCase):
    def setUp(self):
        self.purchases = [
            {"medicine_id": 1, "qty": 10, "free_qty": 5, "rate": 10, "units": 1},  # ZIFI B1
            {"medicine_id": 2, "qty": 10, "free_qty": 0, "rate": 40, "units": 1},  # ZIFI B2
            {"medicine_id": 4, "qty": 2, "free_qty": 0, "rate": 50, "units": 10},  # TAB, Rs 5 a tablet
        ]
        zifi_b1 = {"medicine_id": 1, "name": "ZIFI 1GM", "type": "Injection", "batch": "B1", "mrp_unit": 66.64}
        zifi_b2 = {"medicine_id": 2, "name": "ZIFI 1GM", "type": "Injection", "batch": "B2", "mrp_unit": 66.64}
        syrup = {"medicine_id": 3, "name": "SYR", "type": "Syrup", "batch": "S1", "mrp_unit": 90}
        tab = {"medicine_id": 4, "name": "TAB", "type": "Tablet", "batch": "T1", "mrp_unit": 10}
        self.sales = [
            _sale(1, "Dr A", [dict(zifi_b1, qty=2), dict(syrup, qty=1), dict(tab, qty=10)]),
            _sale(2, "Dr A", [dict(zifi_b2, qty=4)], returns=[{"medicine_id": 2, "qty": 1}]),
            _sale(3, "", [dict(tab, qty=5)]),  # no doctor: nobody's margin
            _sale(4, "Dr B", [dict(zifi_b1, qty=1), dict(tab, qty=2)]),
        ]

    def test_only_margin_lines_and_totals(self):
        rules = {"sets": [{"doctor": "", "types": [
            {"type": "Injection", "mode": "type", "margin": {"kind": "mrp_rate"}},
            {"type": "Tablet", "mode": "type", "margin": {"kind": "mrp_rate"}},
        ]}], "rate_basis": "purchase_free"}
        rep = dmr.compute(self.sales, self.purchases, rules)
        names = {(r["doctor"], r["name"]) for r in rep["rows"]}
        self.assertNotIn(("DR A", "SYR"), names)  # Syrup gives no margin: not in the report
        self.assertFalse(any(r["doctor"] == "" for r in rep["rows"]))
        b1 = next(r for r in rep["rows"] if r["bill_no"] == "B1" and r["batch"] == "B1")
        self.assertEqual(b1["rate"], 6.6667)  # 100 for 10 + 5 free
        self.assertEqual(b1["value"], round((66.64 - 100 / 15) * 2, 2))
        b2 = next(r for r in rep["rows"] if r["bill_no"] == "B2")
        self.assertEqual(b2["qty"], 3)  # 4 sold, 1 returned
        self.assertEqual(b2["value"], round((66.64 - 40) * 3, 2))
        tab = next(r for r in rep["rows"] if r["bill_no"] == "B1" and r["name"] == "TAB")
        self.assertEqual(tab["value"], 50.0)  # (10 - 5) x 10 tablets
        self.assertAlmostEqual(rep["totals"]["value"], round(sum(r["value"] for r in rep["rows"]), 2), places=9)
        self.assertAlmostEqual(sum(d["value"] for d in rep["doctors"]), rep["totals"]["value"], places=9)
        self.assertAlmostEqual(sum(g["value"] for g in rep["summary"]), rep["totals"]["value"], places=9)
        self.assertEqual(sum(g["qty"] for g in rep["summary"]), rep["totals"]["qty"])

    def test_report_rate_choice(self):
        rules = {"sets": [{"doctor": "", "types": [{"type": "Injection", "mode": "type",
                                                     "margin": {"kind": "mrp_rate"}}]}]}
        billed = dmr.compute(self.sales, self.purchases, rules, {"rate_basis": "purchase", "doctors": ["dr a"]})
        free = dmr.compute(self.sales, self.purchases, rules, {"rate_basis": "purchase_free", "doctors": ["DR A"]})
        r1 = next(r for r in billed["rows"] if r["batch"] == "B1")
        r2 = next(r for r in free["rows"] if r["batch"] == "B1")
        self.assertEqual(r1["rate"], 10.0)
        self.assertEqual(r2["rate"], 6.6667)
        self.assertEqual({r["doctor"] for r in billed["rows"]}, {"DR A"})
        # a saved line seen with the other rate gives what computing with that rate gives
        self.assertEqual(dmr.with_basis(r1, "purchase_free")["value"], r2["value"])

    def test_returns_take_lines_without_margin_first_in_bill_order(self):
        rules = {"sets": [{"doctor": "", "types": [{"type": "Injection", "mode": "groups", "rest": None, "groups": [
            {"name": "G", "margin": fixed(10), "items": [{"product": "ZIFI 1GM", "batches": ["B2"]}]}]}]}]}
        line = {"medicine_id": 2, "name": "ZIFI 1GM", "type": "Injection", "batch": "B2", "mrp_unit": 66.64}
        s = _sale(9, "Dr A", [dict(line, qty=2), dict(line, qty=3)], returns=[{"medicine_id": 2, "qty": 4}])
        rep = dmr.compute([s], self.purchases, rules)
        self.assertEqual([r["qty"] for r in rep["rows"]], [1])  # 2 then 2 of 3 came back
        self.assertEqual(rep["totals"]["value"], 10.0)

    def test_pct_of_mrp_and_rate_over_mrp(self):
        rules = {"sets": [{"doctor": "", "types": [
            {"type": "Injection", "mode": "type", "margin": {"kind": "pct_mrp", "value": 10}},
            {"type": "Tablet", "mode": "type", "margin": {"kind": "mrp_rate"}}]}]}
        tab = {"medicine_id": 4, "name": "TAB", "type": "Tablet", "batch": "T1", "mrp_unit": 4}  # below Rs 5 rate
        inj = {"medicine_id": 1, "name": "ZIFI 1GM", "type": "Injection", "batch": "B1", "mrp_unit": 66.64}
        rep = dmr.compute([_sale(5, "Dr A", [dict(tab, qty=2), dict(inj, qty=1)])], self.purchases, rules)
        t = next(r for r in rep["rows"] if r["name"] == "TAB")
        self.assertEqual(t["value"], 0.0)
        self.assertIn("rate_over_mrp", t["flags"])
        self.assertEqual(next(r for r in rep["rows"] if r["name"] == "ZIFI 1GM")["value"], 6.66)


# ── the saved lines, on a store database ───────────────────────────────────────────────────

def _store():
    conn = sqlite3.connect(tempfile.mktemp(suffix=".db"))
    db_setup.initialise(conn)
    return conn


def _table(conn, table):
    return conn.execute(f"SELECT * FROM {table} ORDER BY id").fetchall()


class SavedLines(unittest.TestCase):
    def setUp(self):
        p = mock.patch("core.sync_prefs.is_online_mode", lambda *a, **k: False)
        p.start()
        self.addCleanup(p.stop)
        self.conn = c = _store()
        self.addCleanup(c.close)
        sup = c.execute("INSERT INTO suppliers (name) VALUES ('TEST SUPPLIER')").lastrowid

        def med(name, mtype, batch, unit, mrp, rate=0):
            return c.execute("INSERT INTO medicines (name, type, batch_no, unit, mrp, rate, stock_qty) "
                             "VALUES (?,?,?,?,?,?,50)", (name, mtype, batch, unit, mrp, rate)).lastrowid

        self.zifi = med("ZIFI 1GM", "Injection", "Z1", "1", 66.64)
        self.lyc = med("LYCEFT 1GM", "Injection", "L1", "1", 66.40)
        self.tab = med("TAB S", "Tablet", "T1", "10", 100)  # Rs 10 a tablet
        self.syr = med("SYR Z", "Syrup", "S1", "1", 90)
        self.old = med("OPENING INJ", "Injection", "O1", "1", 50, rate=20)  # no purchase bill
        pid = c.execute("INSERT INTO purchases (purchase_no, supplier_id, purchase_date) VALUES ('P1',?, '2026-10-01')",
                        (sup,)).lastrowid
        for mid, mtype, qty, free, rate, stocked in (
            (self.zifi, "Injection", 10, 5, 10, 15),   # Rs 10 billed, 6.6667 with free
            (self.lyc, "Injection", 10, 0, 21.47, 10),
            (self.tab, "Tablet", 2, 0, 50, 20),        # Rs 5 a tablet
            (self.syr, "Syrup", 5, 0, 40, 5),
        ):
            c.execute("INSERT INTO purchase_items (purchase_id, medicine_id, type, qty, free_qty, rate, stock_units) "
                      "VALUES (?,?,?,?,?,?,?)", (pid, mid, mtype, qty, free, rate, stocked))
        c.commit()
        dmr.save_rules(c, {"rate_basis": "purchase_free", "sets": [{"doctor": "", "types": [
            {"type": "Injection", "mode": "type", "margin": {"kind": "mrp_rate"}},
            {"type": "Tablet", "mode": "type", "margin": {"kind": "mrp_rate"}},
        ]}]})

    def sale(self, doctor, lines, date="2026-10-12", bill="S1"):
        c = self.conn
        cust = c.execute("INSERT INTO customers (name) VALUES (?)", (f"PATIENT {bill}",)).lastrowid
        sid = c.execute("INSERT INTO sales (bill_no, bill_date, doctor_name, customer_id, customer_name, total_amount) "
                        "VALUES (?,?,?,?,?,0)", (bill, date, doctor, cust, f"PATIENT {bill}")).lastrowid
        for mid, qty in lines:
            c.execute("INSERT INTO sales_items (sale_id, medicine_id, qty, rate, amount) VALUES (?,?,?,?,?)",
                      (sid, mid, qty, 1, qty))
        c.commit()
        return sid

    def bill(self, doctor, lines, date="2026-10-12", bill="S1"):
        """lines: (medicine id, qty, sale rate a unit, line discount Rs) -- amount = qty x rate - disc."""
        c = self.conn
        cust = c.execute("INSERT INTO customers (name) VALUES (?)", (f"PATIENT {bill}",)).lastrowid
        sid = c.execute("INSERT INTO sales (bill_no, bill_date, doctor_name, customer_id, customer_name, total_amount) "
                        "VALUES (?,?,?,?,?,0)", (bill, date, doctor, cust, f"PATIENT {bill}")).lastrowid
        for mid, qty, rate, disc in lines:
            c.execute("INSERT INTO sales_items (sale_id, medicine_id, qty, rate, amount, item_discount) "
                      "VALUES (?,?,?,?,?,?)", (sid, mid, qty, rate, round(qty * rate - disc, 2), disc))
        c.commit()
        return sid

    def saved(self, sid):
        return self.conn.execute(
            "SELECT name, qty_sold, qty_returned, qty, rate, value FROM doctor_margin_lines WHERE sale_id=? "
            "ORDER BY line_no", (sid,)).fetchall()

    def ret(self, sid, mid, qty):
        c = self.conn
        rid = c.execute("INSERT INTO sales_returns (return_no, sale_id, return_date) VALUES (?,?,?)",
                        (f"R{sid}-{mid}-{qty}", sid, "2026-10-15")).lastrowid
        c.execute("INSERT INTO sales_return_items (return_id, medicine_id, qty) VALUES (?,?,?)", (rid, mid, qty))
        c.commit()
        return rid

    def test_save_writes_only_margin_lines(self):
        sid = self.sale("Dr A", [(self.zifi, 3), (self.syr, 1), (self.tab, 10)])
        self.assertTrue(st.sync_sale(self.conn, sid)["ok"])
        rows = self.saved(sid)
        self.assertEqual([r[0] for r in rows], ["ZIFI 1GM", "TAB S"])  # the syrup gives no margin: not saved
        self.assertEqual(rows[0][4], 6.6667)
        self.assertEqual(rows[0][5], round((66.64 - 6.6667) * 3, 2))
        self.assertEqual(rows[1][5], 50.0)  # (10 - 5) x 10 tablets
        nodoc = self.sale("", [(self.zifi, 1)], bill="S2")
        st.sync_sale(self.conn, nodoc)
        self.assertEqual(self.saved(nodoc), [])

    def test_edit_delete_and_return_keep_the_rows_right(self):
        sid = self.sale("Dr A", [(self.zifi, 3), (self.tab, 10)])
        st.sync_sale(self.conn, sid)
        # an edit that changes a line: worked out again
        self.conn.execute("UPDATE sales_items SET qty=5 WHERE sale_id=? AND medicine_id=?", (sid, self.zifi))
        self.conn.commit()
        st.sync_sale(self.conn, sid)
        self.assertEqual(self.saved(sid)[0][3], 5)
        # a return takes qty off (rate stays as saved)
        rid = self.ret(sid, self.zifi, 2)
        st.refresh_returns(self.conn, [sid])
        self.assertEqual(self.saved(sid)[0][1:4], (5, 2, 3))
        self.assertEqual(self.saved(sid)[0][5], round((66.64 - 6.6667) * 3, 2))
        # the return deleted: the qty comes back
        self.conn.execute("UPDATE sales_returns SET deleted=1 WHERE id=?", (rid,))
        self.conn.commit()
        st.refresh_returns(self.conn, None)
        self.assertEqual(self.saved(sid)[0][1:4], (5, 0, 5))
        # the bill deleted: its lines go
        self.conn.execute("UPDATE sales SET deleted=1 WHERE id=?", (sid,))
        self.conn.commit()
        st.sync_sale(self.conn, sid)
        self.assertEqual(self.saved(sid), [])
        self.assertIsNone(self.conn.execute("SELECT 1 FROM doctor_margin_bills WHERE sale_id=?", (sid,)).fetchone())

    def test_lines_rewritten_under_new_ids_keep_their_figures(self):
        sid = self.sale("Dr A", [(self.zifi, 3)])
        st.sync_sale(self.conn, sid)
        self.conn.execute("DELETE FROM sales_items WHERE sale_id=?", (sid,))
        self.conn.execute("INSERT INTO sales_items (sale_id, medicine_id, qty, rate, amount) VALUES (?,?,3,1,3)",
                          (sid, self.zifi))
        self.conn.commit()
        dmr.save_rules(self.conn, {"sets": [{"doctor": "", "types": [
            {"type": "Injection", "mode": "type", "margin": {"kind": "fixed", "value": 1}}]}]})
        self.assertTrue(st.sync_sale(self.conn, sid).get("kept"))
        self.ret(sid, self.zifi, 1)
        st.refresh_returns(self.conn, [sid])
        self.assertEqual(self.saved(sid)[0][1:4], (3, 1, 2))
        self.assertEqual(self.saved(sid)[0][5], round((66.64 - 6.6667) * 2, 2))

    def test_rule_change_waits_for_recalculate(self):
        sid = self.sale("Dr A", [(self.zifi, 3)])
        st.sync_sale(self.conn, sid)
        before = self.saved(sid)
        dmr.save_rules(self.conn, {"sets": [{"doctor": "", "types": [
            {"type": "Injection", "mode": "type", "margin": {"kind": "fixed", "value": 2}}]}]})
        rep = rr.build_report(self.conn, {"from": FROM, "to": TO})
        self.assertEqual(self.saved(sid), before)  # the report opening does not re-price a saved bill
        self.assertEqual(rep["totals"]["value"], before[0][5])
        rep = rr.build_report(self.conn, {"from": FROM, "to": TO, "recalculate": True})
        self.assertEqual(rep["totals"]["value"], 6.0)  # Rs 2 x 3 with today's rules
        self.assertEqual(self.saved(sid)[0][5], 6.0)

    def test_recalculate_leaves_a_marker_for_the_server(self):
        """The server's report (web, phone) follows a period recalculated here."""
        import json

        from core.desktop_settings_service import _setting_get

        self.assertEqual(_setting_get(self.conn, dmr.RECALC_SETTING, ""), "")
        rr.build_report(self.conn, {"from": FROM, "to": TO})  # opening the report: no marker
        self.assertEqual(_setting_get(self.conn, dmr.RECALC_SETTING, ""), "")
        rules = dmr.load_rules(self.conn)
        rr.build_report(self.conn, {"from": FROM, "to": TO, "recalculate": True})
        rr.build_report(self.conn, {"from": "2026-10-01", "to": "2026-10-05", "recalculate": True})
        data = json.loads(_setting_get(self.conn, dmr.RECALC_SETTING, ""))
        self.assertEqual(data["version"], 1)
        self.assertEqual([(r["from"], r["to"]) for r in data["runs"]], [(FROM, TO), ("2026-10-01", "2026-10-05")])
        self.assertEqual({r["rules_updated"] for r in data["runs"]}, {rules["updated"]})
        self.assertTrue(all(len(r["at"]) == 19 for r in data["runs"]))
        # only the last RECALC_KEEP runs are kept
        with mock.patch.object(dmr, "RECALC_KEEP", 2):
            rr.build_report(self.conn, {"from": TO, "to": TO, "recalculate": True})
        data = json.loads(_setting_get(self.conn, dmr.RECALC_SETTING, ""))
        self.assertEqual([r["from"] for r in data["runs"]], ["2026-10-01", TO])
        # it travels with the shop settings (offline-first triggers, mode switch carry)
        from core import store_kv_carry
        from core.offline_first import schema

        self.assertIn(dmr.RECALC_SETTING, schema.SHOP_WIDE_SETTING_NAMES)
        self.assertTrue(store_kv_carry._carried(dmr.RECALC_SETTING))

    def test_report_fills_bills_that_arrived_by_sync(self):
        a = self.sale("Dr A", [(self.zifi, 2)])  # saved on another device: no hook ran here
        b = self.sale("Dr B", [(self.lyc, 1), (self.syr, 2)], bill="S2")
        out_of_period = self.sale("Dr A", [(self.zifi, 1)], date="2026-09-01", bill="S3")
        rep = rr.build_report(self.conn, {"from": FROM, "to": TO})
        self.assertEqual(rep["sync"]["worked_out"], 2)
        self.assertEqual({r["doctor"] for r in rep["rows"]}, {"DR A", "DR B"})
        self.assertEqual([r["name"] for r in rep["rows"] if r["doctor"] == "DR B"], ["LYCEFT 1GM"])
        self.assertEqual(self.saved(out_of_period), [])
        # the next opening has nothing to do
        self.assertEqual(rr.build_report(self.conn, {"from": FROM, "to": TO})["sync"]["worked_out"], 0)
        # a bill deleted on another device drops out
        self.conn.execute("UPDATE sales SET deleted=1 WHERE id=?", (a,))
        self.conn.commit()
        rep = rr.build_report(self.conn, {"from": FROM, "to": TO})
        self.assertEqual({r["doctor"] for r in rep["rows"]}, {"DR B"})
        self.assertTrue(b)

    def test_report_layouts_totals_and_files(self):
        self.sale("Dr A", [(self.zifi, 3), (self.tab, 10), (self.old, 2)])
        self.sale("Dr A", [(self.zifi, 1)], bill="S2", date="2026-10-13")
        self.sale("Dr B", [(self.lyc, 6)], bill="S3")
        rep = rr.build_report(self.conn, {"from": FROM, "to": TO, "view": "detail"})
        self.assertAlmostEqual(rep["totals"]["value"], round(sum(r["value"] for r in rep["rows"]), 2), places=9)
        self.assertAlmostEqual(sum(p["value"] for p in rep["products"]), rep["totals"]["value"], places=9)
        old = next(r for r in rep["rows"] if r["name"] == "OPENING INJ")
        self.assertEqual((old["rate"], old["flags"]), (20.0, ["estimated"]))  # the batch's own rate
        zifi = next(g for g in rep["summary"] if g["name"] == "ZIFI 1GM")
        self.assertEqual(zifi["qty"], 4)  # two bills, one row: same batch, MRP and rate
        only_b = rr.build_report(self.conn, {"from": FROM, "to": TO, "doctors": ["dr b"]})
        self.assertEqual(only_b["totals"]["value"], round((66.40 - 21.47) * 6, 2))
        for view in ("doctor", "detail"):
            out = rr.export_report(self.conn, {"from": FROM, "to": TO, "view": view, "format": "xlsx"})
            self.assertTrue(out["ok"], out)
            import base64

            import openpyxl

            wb = openpyxl.load_workbook(io.BytesIO(base64.b64decode(out["content_base64"])))
            cells = [c for row in wb.active.iter_rows(values_only=True) for c in row if c is not None]
            self.assertIn("TOTAL value" if view == "doctor" else "TOTAL value (on profit)", cells)
            self.assertIn(rep["totals"]["value"], cells) if view == "detail" else None
        csv = rr.export_report(self.conn, {"from": FROM, "to": TO, "format": "csv"})
        self.assertTrue(csv["ok"])
        # the existing route keeps working: the dialog's mode 'rules' goes through it
        from core import doctor_margin_report as old_rep

        self.assertEqual(old_rep.build_report(self.conn, {"from": FROM, "to": TO, "mode": "rules"})["totals"],
                         rep["totals"])

    def test_nothing_outside_the_doctor_report_changes(self):
        """Owner: "this rate is only for the doctor not for the sale"."""
        from core import desktop_sales_service as sales_svc

        sid = self.sale("Dr A", [(self.zifi, 3), (self.tab, 10)])
        books = ("sales", "sales_items", "purchases", "purchase_items", "medicines", "sales_returns",
                 "sales_return_items")
        before = {t: _table(self.conn, t) for t in books}
        screen_before = sales_svc.load_sale(self.conn, sid)
        st.sync_sale(self.conn, sid)
        for basis in ("purchase", "purchase_free"):
            dmr.save_rules(self.conn, dict(dmr.load_rules(self.conn), rate_basis=basis))
            rr.build_report(self.conn, {"from": FROM, "to": TO, "recalculate": True})
            rr.build_report(self.conn, {"from": FROM, "to": TO, "rate_basis": basis})
        self.assertEqual({t: _table(self.conn, t) for t in books}, before)
        # The Sales screen's lines and margin column for that bill, as before.
        after = sales_svc.load_sale(self.conn, sid)
        pick = ("qty", "rate", "amount", "margin", "gross_margin", "net_margin", "purchase_rate", "margin_pct")
        meds = lambda d: [{k: m.get(k) for k in pick} for m in (d.get("form") or {}).get("items") or []]  # noqa: E731
        self.assertTrue(meds(after), screen_before)
        self.assertEqual(meds(after), meds(screen_before))

    def test_the_rules_travel_with_the_shop_settings(self):
        from core import store_kv_carry
        from core.offline_first import schema

        self.assertIn("doctor_margin_rules", schema.SHOP_WIDE_SETTING_NAMES)
        self.assertTrue(store_kv_carry._carried(dmr.SETTING_NAME))
        schema.ensure_schema(self.conn)
        names = {r[0] for r in self.conn.execute("SELECT name FROM sqlite_master WHERE type='trigger'")}
        self.assertIn("of_insert_settings_v7", names)
        self.assertNotIn("of_insert_settings_v6", names)
        self.assertNotIn("of_insert_settings_v4", names)
        self.assertEqual(dmr.load_rules(self.conn)["rate_basis"], "purchase_free")
        # the saved lines are this PC's working copy: no sync trigger on them
        self.assertFalse(any("doctor_margin" in n for n in names))



# ── the second way: a % of the total sales (owner, 23 Oct 2026) ──────────────────────────

SALES_RULES = {"rate_basis": "purchase_free", "sets": [
    {"doctor": "", "sales_pct": 10, "types": [
        {"type": "Tablet", "mode": "sales"},
        {"type": "Syrup", "mode": "sales"},
        {"type": "Injection", "mode": "groups", "groups": [
            {"name": "Zifi", "margin": {"kind": "mrp_rate"}, "items": [{"product": "ZIFI 1GM"}]}]}]},
    {"doctor": "Dr B", "sales_pct": 5, "types": []},
]}


def _line(mid, name, mtype, qty, rate, disc=0.0, pack="", mrp=None):
    return {"medicine_id": mid, "name": name, "type": mtype, "batch": f"B{mid}", "qty": qty,
            "mrp_unit": rate if mrp is None else mrp, "sale_rate": rate, "disc": disc,
            "amount": round(qty * rate - disc, 2), "pack": pack}


class TotalSalesCalculation(unittest.TestCase):
    def test_percent_of_the_sale_value_net_of_line_discount_and_returns(self):
        s1 = _sale(1, "Dr A", [_line(1, "TAB", "Tablet", 10, 10, disc=5, pack="10"),
                               _line(2, "SYR", "Syrup", 2, 90),
                               _line(3, "ZIFI 1GM", "Injection", 3, 66.64),
                               _line(4, "LYCEFT 1GM", "Injection", 1, 66.40)],  # in no group: nowhere
                   returns=[{"medicine_id": 1, "qty": 4}])
        s2 = _sale(2, "Dr B", [_line(1, "TAB", "Tablet", 2, 10)])
        purchases = [{"medicine_id": 3, "qty": 10, "free_qty": 0, "rate": 40, "units": 1}]
        rep = dmr.compute([s1, s2], purchases, SALES_RULES)
        self.assertEqual([(r["name"], r["way"]) for r in rep["rows"]], [("ZIFI 1GM", "profit")])
        sales = rep["sales"]
        a = next(d for d in sales["doctors"] if d["doctor"] == "DR A")
        tab = next(r for r in sales["rows"] if r["doctor"] == "DR A" and r["name"] == "TAB")
        self.assertEqual((tab["qty_sold"], tab["qty_returned"], tab["qty"]), (10, 4, 6))
        self.assertEqual((tab["sale_rate"], tab["disc"], tab["amount"], tab["value"]), (10, 5, 95, 57.0))
        self.assertEqual(a["value"], 57.0 + 180.0)
        self.assertEqual(a["by_pct"], [{"pct": 10.0, "value": 237.0, "commission": 23.7}])
        self.assertEqual(a["types"], ["Syrup", "Tablet"])
        self.assertEqual([(p["name"], p["pack"], p["qty"], p["value"]) for p in a["products"]],
                         [("SYR", "", 2, 180.0), ("TAB", "10", 6, 57.0)])
        b = next(d for d in sales["doctors"] if d["doctor"] == "DR B")
        self.assertEqual((b["value"], b["commission"]), (20.0, 1.0))  # his own 5%
        g = {d["doctor"]: d for d in rep["grand"]["doctors"]}
        self.assertEqual(g["DR A"]["profit"], round((66.64 - 40) * 3, 2))
        self.assertEqual(g["DR A"]["total"], round(g["DR A"]["profit"] + 23.7, 2))
        self.assertEqual(rep["grand"]["totals"]["total"], round(rep["totals"]["value"] + 24.7, 2))
        self.assertEqual(sales["totals"]["commission"], 24.7)

    def test_commission_on_the_total_not_line_by_line(self):
        lines = [_line(i, f"T{i}", "Tablet", 1, 0.33) for i in range(1, 4)]
        rep = dmr.compute([_sale(1, "Dr A", lines)], [], SALES_RULES)
        d = rep["sales"]["doctors"][0]
        self.assertEqual(d["value"], 0.99)
        self.assertEqual(d["commission"], 0.1)  # 0.099 -> 0.10, not 3 x 0.03

    def test_a_rate_choice_leaves_sales_lines_alone(self):
        rep = dmr.compute([_sale(1, "Dr A", [_line(1, "TAB", "Tablet", 3, 10)])], [], SALES_RULES)
        row = rep["sales"]["rows"][0]
        self.assertIs(dmr.with_basis(row, "purchase"), row)


class TotalSalesSaved(unittest.TestCase):
    setUp = SavedLines.setUp
    bill = SavedLines.bill
    ret = SavedLines.ret

    def lines(self, sid):
        return self.conn.execute(
            "SELECT name, way, qty_sold, qty, sale_rate, disc, amount, value, kind, kind_value, rules_at "
            "FROM doctor_margin_lines WHERE sale_id=? ORDER BY line_no", (sid,)).fetchall()

    def test_saved_at_bill_save_kept_until_recalculate(self):
        dmr.save_rules(self.conn, SALES_RULES)
        rules_at = dmr.load_rules(self.conn)["updated"]
        sid = self.bill("Dr A", [(self.tab, 10, 10, 5), (self.syr, 2, 90, 0), (self.zifi, 3, 66.64, 0),
                                 (self.lyc, 1, 66.40, 0)])
        st.sync_sale(self.conn, sid)
        got = self.lines(sid)
        self.assertEqual([(r[0], r[1]) for r in got], [("TAB S", "sales"), ("SYR Z", "sales"), ("ZIFI 1GM", "profit")])
        self.assertEqual(got[0][2:], (10, 10, 10, 5, 95, 95, "sales_pct", 10, rules_at))
        rep = rr.build_report(self.conn, {"from": FROM, "to": TO})
        d = rep["sales"]["doctors"][0]
        self.assertEqual((d["value"], d["commission"]), (275.0, 27.5))
        # a return: the sale value comes down with the qty
        self.ret(sid, self.tab, 4)
        st.refresh_returns(self.conn, [sid])
        self.assertEqual(self.lines(sid)[0][3], 6)
        self.assertEqual(self.lines(sid)[0][7], 57.0)
        # the price edited (same medicines and qty): the bill's value, the saved %
        self.conn.execute("UPDATE sales_items SET amount=90, item_discount=10 WHERE sale_id=? AND medicine_id=?",
                          (sid, self.tab))
        self.conn.commit()
        self.assertTrue(st.sync_sale(self.conn, sid).get("kept"))
        self.assertEqual(self.lines(sid)[0][5:8], (10, 90, 54.0))
        # a new % waits for Recalculate
        dmr.save_rules(self.conn, dict(SALES_RULES, sets=[dict(SALES_RULES["sets"][0], sales_pct=20)]))
        rep = rr.build_report(self.conn, {"from": FROM, "to": TO})
        self.assertEqual(rep["sales"]["doctors"][0]["by_pct"][0]["pct"], 10)
        rep = rr.build_report(self.conn, {"from": FROM, "to": TO, "recalculate": True})
        d = rep["sales"]["doctors"][0]
        self.assertEqual((d["value"], d["by_pct"][0]["pct"], d["commission"]), (234.0, 20, 46.8))
        self.assertEqual(rep["grand"]["doctors"][0]["total"], round(rep["totals"]["value"] + 46.8, 2))

    def test_two_percents_in_one_period(self):
        dmr.save_rules(self.conn, SALES_RULES)
        a = self.bill("Dr A", [(self.tab, 10, 10, 0)])
        st.sync_sale(self.conn, a)
        dmr.save_rules(self.conn, dict(SALES_RULES, sets=[dict(SALES_RULES["sets"][0], sales_pct=20)]))
        b = self.bill("Dr A", [(self.tab, 5, 10, 0)], bill="S2")
        st.sync_sale(self.conn, b)
        d = rr.build_report(self.conn, {"from": FROM, "to": TO})["sales"]["doctors"][0]
        self.assertEqual(d["by_pct"], [{"pct": 20.0, "value": 50.0, "commission": 10.0},
                                       {"pct": 10.0, "value": 100.0, "commission": 10.0}])
        self.assertEqual(d["commission"], 20.0)

    def test_a_table_made_before_gets_the_new_columns(self):
        c = self.conn
        c.execute("DROP TABLE IF EXISTS doctor_margin_lines")
        c.execute("CREATE TABLE doctor_margin_lines (id INTEGER PRIMARY KEY AUTOINCREMENT, sale_id INTEGER NOT NULL, "
                  "sale_item_id INTEGER, line_no INTEGER, bill_no TEXT, bill_date TEXT, patient TEXT, doctor TEXT, "
                  "medicine_id INTEGER, name TEXT, type TEXT, batch TEXT, qty_sold REAL DEFAULT 0, "
                  "qty_returned REAL DEFAULT 0, qty REAL DEFAULT 0, mrp REAL DEFAULT 0, rate_purchase REAL DEFAULT 0, "
                  "rate_free REAL DEFAULT 0, rate_basis TEXT, rate REAL DEFAULT 0, kind TEXT, kind_value REAL DEFAULT 0, "
                  "margin_unit REAL DEFAULT 0, value REAL DEFAULT 0, rule TEXT, flags TEXT, saved_at TEXT)")
        c.execute("INSERT INTO doctor_margin_lines (sale_id, line_no, bill_no, bill_date, doctor, name, type, qty_sold, "
                  "qty, mrp, rate, rate_basis, kind, margin_unit, value, rule, flags) VALUES "
                  "(99, 1, 'OLD', '2026-10-12', 'DR A', 'OLD INJ', 'Injection', 2, 2, 50, 20, 'purchase', 'mrp_rate', "
                  "30, 60, 'Injection > rest of type', '[]')")
        c.commit()
        rows = st.read_lines(c, FROM, TO)
        self.assertEqual([(r["name"], r["way"], r["value"]) for r in rows], [("OLD INJ", "profit", 60)])
        self.assertTrue({"way", "pack", "sale_rate", "disc", "amount", "rules_at"} <= st._cols(c, "doctor_margin_lines"))
        rep = dmr.assemble(rows)
        self.assertEqual((rep["totals"]["value"], rep["sales"]["rows"]), (60, []))  # kept as saved

    def test_files_carry_both_parts(self):
        import base64

        import openpyxl

        dmr.save_rules(self.conn, SALES_RULES)
        st.sync_sale(self.conn, self.bill("Dr A", [(self.tab, 10, 10, 5), (self.syr, 2, 90, 0), (self.zifi, 3, 66.64, 0)]))
        st.sync_sale(self.conn, self.bill("Dr B", [(self.syr, 1, 90, 0)], bill="S2"))
        body = {"from": FROM, "to": TO}
        out = rr.export_report(self.conn, dict(body, format="xlsx"))
        self.assertTrue(out["ok"], out)
        wb = openpyxl.load_workbook(io.BytesIO(base64.b64decode(out["content_base64"])))
        cells = [c for row in wb.active.iter_rows(values_only=True) for c in row if c is not None]
        for want in ("TAB S 10", "SYR Z", "Sub Total:", "Total sales value (SYRUP, TABLET)", "Commission 10% of 275.00",
                     27.5, "Commission 5% of 90.00", 4.5, "TOTAL value", "All doctors", "Doctor's margin"):
            self.assertIn(want, cells)
        page = base64.b64decode(rr.export_report(self.conn, dict(body, format="html"))["content_base64"]).decode()
        self.assertIn("Customer-Wise Product-Batch Wise Sales From 11/10/2026 To 20/10/2026", page)
        self.assertIn("size:A4 portrait", page)
        self.assertIn("counter(pages)", page)
        csv = base64.b64decode(rr.export_report(self.conn, dict(body, format="csv"))["content_base64"]).decode("utf-8-sig")
        self.assertIn("Sub Total:", csv)
        self.assertIn("27.50", csv)
        # only total-sales lines (no profit part): a file all the same
        only = rr.export_report(self.conn, dict(body, doctors=["DR B"], format="csv"))
        self.assertTrue(only["ok"], only)


if __name__ == "__main__":
    unittest.main()
