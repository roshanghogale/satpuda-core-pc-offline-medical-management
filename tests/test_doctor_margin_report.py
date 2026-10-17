"""Doctor Margin Report (owner, 17 Oct 2026; Sales History -> Export -> Doctor Margin Report).

core/doctor_margin.py is the PC copy of the server's rule (server-live
src/services/doctorMargin.js); tests/fixtures/doctor_margin_vectors.json is shared with the
server and the phone. core/doctor_margin_report.py reads this PC's books for it and lays the
result out for the dialog and the files. Tests A-E are the owner's.
"""
import base64
import csv
import io
import json
import os
import sqlite3
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from core import db_setup  # noqa: E402
from core import doctor_margin as dm  # noqa: E402
from core import doctor_margin_report as rep  # noqa: E402

VECTORS = Path(__file__).parent / "fixtures" / "doctor_margin_vectors.json"
FROM, TO = "2026-10-01", "2026-10-06"
PATIENTS = ("PATIENT ONE", "PATIENT TWO", "PATIENT THREE", "PATIENT FOUR", "PATIENT FIVE",
            "PATIENT NINE", "PATIENT TEN")


class TheSharedVectors(unittest.TestCase):
    def test_every_vector(self):
        cases = json.loads(VECTORS.read_text(encoding="utf-8"))
        self.assertGreaterEqual(len(cases), 11)
        for c in cases:
            with self.subTest(c["name"]):
                self.assertEqual(dm.rounded(dm.doctor_margin(c["sales"], c["purchases"], c["opts"])), c["expect"])

    def test_the_two_styles_reconcile(self):
        for c in json.loads(VECTORS.read_text(encoding="utf-8")):
            got = dm.doctor_margin(c["sales"], c["purchases"], c["opts"])
            for k in ("qty", "amount", "cost", "commission"):
                with self.subTest(c["name"], k=k):
                    self.assertAlmostEqual(sum(r[k] for r in got["rows"]), sum(r[k] for r in got["summary"]), places=6)
                    self.assertAlmostEqual(sum(r[k] for r in got["rows"]), got["totals"][k], places=6)
                    self.assertAlmostEqual(sum(d[k] for d in got["doctors"]), got["totals"][k], places=6)

    def test_rounding_is_javascripts(self):
        self.assertEqual(dm.r2(0.125), 0.13)
        self.assertEqual(dm.r2(-0.125), -0.12)  # a half goes up, also below zero
        self.assertEqual(dm.r2(1.005), 1.0)  # 100.49999999999999 in binary
        self.assertEqual(dm.r2("12.345"), 12.35)
        self.assertEqual(dm.r2(None), 0.0)
        self.assertEqual(dm._math_round(0.49999999999999994), 0.0)  # Math.round, not floor(x + .5)
        self.assertEqual(dm._to_fixed6(0.0078125), "0.007813")  # toFixed: the exact half goes away from 0

    def test_nothing_is_rounded_before_rounded(self):
        out = dm.doctor_margin(
            [{"id": 1, "bill_no": "B", "date": "2026-10-01", "doctor": "x", "total_amount": 10,
              "items": [{"medicine_id": 1, "type": "Tablet", "qty": 3, "amount": 10}], "returns": []}],
            [{"medicine_id": 1, "qty": 3, "free_qty": 0, "item_amount": 1, "units": 1}], {"pct": 10})
        self.assertEqual(out["rows"][0]["sale_rate"], 10 / 3)
        self.assertEqual(out["rows"][0]["doctor"], "X")


def _store() -> sqlite3.Connection:
    conn = sqlite3.connect(tempfile.mktemp(suffix=".db"))
    db_setup.initialise(conn)
    return conn


class Books(unittest.TestCase):
    """A small store: schemes, a strip medicine, a bill discount, a return, no-doctor,
    deleted / draft bills and purchases, a loss, an injection."""

    def setUp(self):
        p = mock.patch("core.sync_prefs.is_online_mode", lambda *a, **k: False)
        p.start()
        self.addCleanup(p.stop)
        self.conn = c = _store()
        self.addCleanup(c.close)
        self.cust2 = c.execute("INSERT INTO customers (name) VALUES ('PATIENT TWO')").lastrowid
        c.execute("INSERT INTO doctors (name) VALUES ('Dr Zed')")
        gone = c.execute("INSERT INTO doctors (name) VALUES ('Dr Gone')").lastrowid
        c.execute("UPDATE doctors SET deleted=1 WHERE id=?", (gone,))
        sup = c.execute("INSERT INTO suppliers (name) VALUES ('TEST SUPPLIER')").lastrowid

        def med(name, mtype, batch, unit="1"):
            return c.execute("INSERT INTO medicines (name, type, batch_no, unit) VALUES (?,?,?,?)",
                             (name, mtype, batch, unit)).lastrowid

        self.g1 = med("GEL X", "Gel", "G1")
        self.g2 = med("GEL X", "Gel", "G2")
        self.t1 = med("TAB S", "Tablet", "T1", "10")
        self.t2 = med("TAB S", "Tablet", "T2", "10")
        self.inj = med("INJ Y", "Injection", "I1")
        self.syr = med("SYR Z", "Syrup", "S1")
        self.cream = med("CREAM W", "Cream", "C1")

        def purchase(no, lines, deleted=0, autosave=0):
            pid = c.execute(
                "INSERT INTO purchases (purchase_no, supplier_id, purchase_date, deleted, is_autosave) "
                "VALUES (?,?,?,?,?)", (no, sup, "2026-09-20", deleted, autosave)).lastrowid
            for mid, mtype, qty, free, amount, stocked in lines:
                c.execute("INSERT INTO purchase_items (purchase_id, medicine_id, type, qty, free_qty, "
                          "item_amount, amount, stock_units) VALUES (?,?,?,?,?,?,?,?)",
                          (pid, mid, mtype, qty, free, amount, amount, stocked))

        # Test B: 10 + 10 free for Rs 100 = Rs 5 a gel
        purchase("P1", [(self.g1, "Gel", 10, 10, 100, 20), (self.g2, "Gel", 10, 0, 100, 10),
                        # 3 strips + 1 free of 10 for Rs 120, as stocked: 40 tablets -> Rs 3 a tablet
                        (self.t1, "Tablet", 3, 1, 120, 40),
                        # an older line without stock_units: the purchase-save rule, 10 a strip -> Rs 5
                        (self.t2, "Tablet", 2, 0, 100, None),
                        (self.inj, "Injection", 2, 0, 200, 2)])
        # neither may move the rate
        purchase("P-DEL", [(self.g1, "Gel", 10, 0, 900, 10)], deleted=1)
        purchase("P-DRAFT", [(self.g1, "Gel", 10, 0, 900, 10)], autosave=1)

        def sale(no, day, doctor, patient, total, lines, deleted=0, autosave=0, customer_id=None):
            sid = c.execute(
                "INSERT INTO sales (bill_no, customer_id, bill_date, total_amount, doctor_name, customer_name, "
                "deleted, is_autosave) VALUES (?,?,?,?,?,?,?,?)",
                (no, customer_id, day, total, doctor, patient, deleted, autosave)).lastrowid
            for mid, qty, amount, cost in lines:
                c.execute("INSERT INTO sales_items (sale_id, medicine_id, qty, rate, amount, cost_price) "
                          "VALUES (?,?,?,?,?,?)", (sid, mid, qty, amount / qty, amount, cost))
            return sid

        sale("SB1/FY2026-27", "2026-10-01", "DR A", "PATIENT ONE", 40, [(self.g1, 2, 40, 0)])
        sale("SB2/FY2026-27", "2026-10-02", "DR A", None, 100, [(self.g1, 4, 80, 0), (self.g2, 1, 20, 0)],
             customer_id=self.cust2)
        sale("SB3/FY2026-27", "2026-10-03", " dr b ", "PATIENT THREE", 240,
             [(self.g1, 2, 40, 0), (self.t1, 10, 50, 0), (self.inj, 1, 150, 0)])
        # no doctor; lines 110, bill 99 after the bill discount; no purchase of either
        sale("SB4/FY2026-27", "2026-10-04", "", "PATIENT FOUR", 99, [(self.syr, 1, 60, 30), (self.cream, 1, 50, 0)])
        s5 = sale("SB5/FY2026-27", "2026-10-05", "DR A", "PATIENT FIVE", 80, [(self.t2, 10, 80, 0)])
        sale("SB6/FY2026-27", "2026-10-05", "DR A", "DELETED PATIENT", 100, [(self.g1, 5, 100, 0)], deleted=1)
        sale("SB7/FY2026-27", "2026-10-05", "DR A", "DRAFT PATIENT", 100, [(self.g1, 5, 100, 0)], autosave=1)
        sale("SB8/FY2026-27", "2026-09-30", "DR A", "EARLY PATIENT", 100, [(self.g1, 5, 100, 0)])
        sale("SB9/FY2026-27", "2026-10-06", "DR C", "PATIENT NINE", 15, [(self.g2, 2, 15, 0)])  # a loss
        sale("SB10/FY2026-27", "2026-10-06", "Dr D", "PATIENT TEN", 200, [(self.g2, 10, 200, 0)])  # Test A

        def sale_return(no, sid, mid, qty, deleted=0):
            rid = c.execute("INSERT INTO sales_returns (return_no, sale_id, return_date, deleted) VALUES (?,?,?,?)",
                            (no, sid, "2026-10-06", deleted)).lastrowid
            c.execute("INSERT INTO sales_return_items (return_id, medicine_id, qty) VALUES (?,?,?)", (rid, mid, qty))

        sale_return("SR1", s5, self.t2, 4)  # 4 of the 10 tablets came back
        sale_return("SR2", s5, self.t2, 5, deleted=1)
        c.commit()

    def report(self, **kw):
        body = {"from": FROM, "to": TO, "mode": "regular", "pct": 10}
        body.update(kw)
        return rep.build_report(self.conn, body)

    def rows_by(self, r):
        return {(x["bill_no"], x["name"], x["batch"]): x for x in r["rows"]}


class TheOwnersTests(Books):
    def test_a_regular_medicine(self):
        r = self.report(doctor="dr d")
        self.assertEqual(r["totals"], {"qty": 10, "amount": 200, "cost": 100, "commission": 20,
                                       "profit": 100, "remaining": 80})
        r = self.report(doctor="dr d", pct=15)
        self.assertEqual((r["totals"]["commission"], r["totals"]["remaining"]), (30, 70))
        self.assertEqual(r["rows"][0]["pct"], 15)

    def test_b_supplier_scheme(self):
        row = self.rows_by(self.report())[("SB1", "GEL X", "G1")]
        self.assertEqual((row["purchase_rate"], row["qty"], row["cost"], row["basis"]), (5, 2, 10, "scheme"))
        summary = [s for s in self.report(style="summary")["summary"] if s["batch"] == "G1"]
        self.assertTrue(all(s["purchase_rate"] == 5 for s in summary))

    def test_c_injection(self):
        r = self.report(mode="injection", pct=37)
        self.assertEqual(len(r["rows"]), 1)
        row = r["rows"][0]
        self.assertEqual((row["name"], row["amount"], row["cost"], row["profit"], row["commission"], row["pct"]),
                         ("INJ Y", 150, 100, 50, 50, 100))
        self.assertEqual(r["pct"], 100)
        self.assertEqual(self.report(mode="injection", pct=5)["totals"], r["totals"])
        self.assertNotIn("INJ Y", {x["name"] for x in self.report()["rows"]})

    def test_d_display_styles(self):
        r = self.report()
        mine = {(s["doctor"], s["name"], s["batch"]): s for s in r["summary"]}
        g1 = mine[("DR A", "GEL X", "G1")]  # two bills, two patients, one row
        self.assertEqual((g1["qty"], g1["amount"], g1["cost"], g1["avg_sale_rate"], g1["purchase_rate"],
                          g1["commission"]), (6, 120, 30, 20, 5, 12))
        self.assertIn(("DR A", "GEL X", "G2"), mine)  # another batch, its own row
        self.assertIn(("DR B", "GEL X", "G1"), mine)  # another doctor, its own row
        self.assertEqual(len([x for x in r["rows"] if x["doctor"] == "DR A" and x["batch"] == "G1"]), 2)
        for k in ("qty", "amount", "cost", "commission", "remaining"):
            self.assertAlmostEqual(sum(x[k] for x in r["rows"]), sum(x[k] for x in r["summary"]), places=6)
            self.assertAlmostEqual(sum(x[k] for x in r["rows"]), r["totals"][k], places=6)
        self.assertNotIn("patient", r["summary"][0])
        self.assertNotIn("bill_no", r["summary"][0])

    def test_e_filters(self):
        r = self.report()
        got = self.rows_by(r)
        # deleted, draft and out-of-range bills are not there; the injection is not regular
        self.assertEqual({k[0] for k in got}, {"SB1", "SB2", "SB3", "SB4", "SB5", "SB9", "SB10"})
        self.assertEqual(r["bill_count"], 7)
        # the name on the bill, else the customer's
        self.assertEqual(got[("SB2", "GEL X", "G1")]["patient"], "PATIENT TWO")
        self.assertEqual(got[("SB3", "TAB S", "T1")]["doctor"], "DR B")
        # strip medicine: tablets sold, the purchase's tablets per strip (stocked 40 -> Rs 3)
        t1 = got[("SB3", "TAB S", "T1")]
        self.assertEqual((t1["qty"], t1["purchase_rate"], t1["cost"], t1["profit"]), (10, 3, 30, 20))
        # a return of 4 of 10 tablets (the deleted return does not count), rate from pack 10
        t2 = got[("SB5", "TAB S", "T2")]
        self.assertEqual((t2["qty"], t2["amount"], t2["purchase_rate"], t2["cost"], t2["commission"]),
                         (6, 48, 5, 30, 4.8))
        # the bill discount shared out; estimated and unavailable cost flagged; no doctor
        syr, cream = got[("SB4", "SYR Z", "S1")], got[("SB4", "CREAM W", "C1")]
        self.assertEqual((syr["doctor"], syr["amount"], syr["cost"], syr["basis"], syr["flags"]),
                         (dm.NO_DOCTOR, 54, 30, "estimated", ["estimated"]))
        self.assertEqual((cream["amount"], cream["cost"], cream["flags"]), (45, 0, ["unavailable"]))
        # a loss is flagged, the commission is still on the sale
        loss = got[("SB9", "GEL X", "G2")]
        self.assertEqual((loss["profit"], loss["commission"], loss["remaining"], loss["flags"]),
                         (-5, 1.5, -6.5, ["loss"]))
        self.assertEqual(r["flag_counts"], {"estimated": 1, "unavailable": 1, "loss": 1})
        self.assertEqual(r["totals"], {"qty": 39, "amount": 592, "cost": 260, "commission": 59.2,
                                       "profit": 332, "remaining": 272.8})
        self.assertEqual([d["doctor"] for d in r["doctors"]], [dm.NO_DOCTOR, "DR A", "DR B", "DR C", "DR D"])
        dr_a = r["doctors"][1]
        self.assertEqual((dr_a["amount"], dr_a["cost"], dr_a["commission"]), (188, 70, 18.8))

        none = self.report(doctor=dm.NO_DOCTOR)
        self.assertEqual({x["bill_no"] for x in none["rows"]}, {"SB4"})
        one_day = self.report(**{"from": "2026-10-02", "to": "2026-10-02"})
        self.assertEqual({x["bill_no"] for x in one_day["rows"]}, {"SB2"})
        typed = self.report(types=["Syrup", "Gel"])
        self.assertEqual({x["type"] for x in typed["rows"]}, {"Syrup", "Gel"})
        self.assertEqual({x["type"] for x in self.report(types=["syrup"])["rows"]}, {"Syrup"})

    def test_the_books_are_read_once_per_table(self):
        seen = []
        self.conn.set_trace_callback(seen.append)
        rep.read_books(self.conn, FROM, TO)
        self.conn.set_trace_callback(None)
        selects = [s for s in seen if s.lstrip().upper().startswith("SELECT")]
        self.assertEqual(len(selects), 4, selects)  # bills, lines, returns, purchase lines
        for table in ("sales_items", "sales_return_items", "purchase_items"):
            self.assertTrue(any(table in s for s in selects), table)

    def test_nothing_is_written(self):
        before = self.conn.total_changes
        self.report()
        rep.report_options(self.conn)
        self.assertEqual(self.conn.total_changes, before)

    def test_options(self):
        with mock.patch("core.layout_config.get_med_types", lambda: ["Tablet", "Syrup", "Injection"]):
            o = rep.report_options(self.conn)
        self.assertEqual(o["doctors"], ["DR A", "DR B", "DR C", "DR D", "DR ZED", dm.NO_DOCTOR])
        self.assertEqual(o["types"][:3], ["Tablet", "Syrup", "Injection"])  # the configured list first
        self.assertEqual(set(o["types"]), {"Tablet", "Syrup", "Injection", "Cream", "Gel"})
        self.assertEqual(o["injection_types"], ["Injection"])
        self.assertEqual(o["method"], "Scheme-adjusted effective purchase rate")

    def test_filters_are_checked(self):
        with self.assertRaises(rep.ReportError):
            self.report(**{"from": "2026-10-06", "to": "2026-10-01"})
        with self.assertRaises(rep.ReportError):
            self.report(pct=150)
        with self.assertRaises(rep.ReportError):
            self.report(pct="ten")
        with self.assertRaises(rep.ReportError):
            self.report(**{"from": "01/10/2026"})

    def test_an_online_store_is_sent_to_the_web(self):
        with mock.patch("core.sync_prefs.is_online_mode", lambda *a, **k: True):
            with self.assertRaises(rep.ReportError):
                self.report()


class TheFiles(Books):
    def export(self, **kw):
        body = {"from": FROM, "to": TO, "mode": "regular", "pct": 10}
        body.update(kw)
        out = rep.export_report(self.conn, body)
        self.assertTrue(out["ok"], out)
        return out

    def csv_rows(self, **kw):
        out = self.export(format="csv", **kw)
        return list(csv.reader(io.StringIO(base64.b64decode(out["content_base64"]).decode("utf-8-sig"))))

    def test_every_file_says_what_it_is(self):
        for style in ("detailed", "summary"):
            rows = self.csv_rows(style=style)
            labels = [r[0] for r in rows if r]
            self.assertEqual(rows[0], ["Doctor Margin Report"])
            for want in ("Generated at", "Date range", "Doctor", "Medicine types", "Report mode",
                         "Display style", "Doctor commission", "Purchase cost method", "Totals",
                         "Totals by doctor", "TOTAL"):
                self.assertIn(want, labels, (style, want))
            head = dict(r[:2] for r in rows[1:10])
            self.assertEqual(head["Date range"], "01/10/2026 to 06/10/2026")
            self.assertEqual(head["Doctor"], "All doctors")
            self.assertEqual(head["Doctor commission"], "10%")
            self.assertIn("Scheme-adjusted effective purchase rate", head["Purchase cost method"])
            self.assertTrue(any(r and r[0].startswith("Estimated cost:") for r in rows))
            self.assertTrue(any(r and r[0].startswith("Loss:") for r in rows))
        inj = dict(r[:2] for r in self.csv_rows(mode="injection")[1:10])
        self.assertEqual(inj["Report mode"], "Injection report")
        self.assertIn("100% of positive profit", inj["Doctor commission"])

    def test_the_detailed_file_names_patients_and_bills(self):
        rows = self.csv_rows(style="detailed")
        text = json.dumps(rows)
        self.assertIn("PATIENT ONE", text)
        self.assertIn("SB1", text)
        cols = rows[rows.index(rep.detailed_columns("regular"))]
        self.assertEqual(cols[:4], ["Sale Date", "Customer / Patient", "Doctor", "Bill No"])
        total = next(r for r in rows if r and r[0] == "TOTAL")
        self.assertEqual(total[10:16], ["592.0", "260.0", "332.0", "10.0", "59.2", "272.8"])

    def test_a_summary_never_names_a_patient_or_a_bill(self):
        rows = self.csv_rows(style="summary")
        text = base64.b64decode(self.export(format="csv", style="summary")["content_base64"]).decode("utf-8-sig")
        for word in PATIENTS + ("SB1", "SB2", "SB3", "SB4", "SB5", "SB9", "SB10", "FY2026"):
            self.assertNotIn(word, text)
        self.assertIn(rep.summary_columns("regular"), rows)

        import openpyxl

        out = self.export(format="xlsx", style="summary")
        self.assertTrue(out["filename"].endswith(".xlsx"))
        wb = openpyxl.load_workbook(io.BytesIO(base64.b64decode(out["content_base64"])))
        self.assertEqual(wb.sheetnames, ["Doctor Margin"])  # no extra or hidden sheet
        cells = " ".join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row
                         if c.value is not None)
        for word in PATIENTS + ("SB1", "SB10"):
            self.assertNotIn(word, cells)
        self.assertIn("Doctor Margin Report", cells)
        self.assertIn("Medicine-wise summary", cells)
        self.assertFalse(wb.properties.creator and "PATIENT" in wb.properties.creator)

        def fake_pdf(html, path, **_kw):
            html_path = path.rsplit(".", 1)[0] + ".html"
            with open(html_path, "w", encoding="utf-8") as fh:
                fh.write(html)
            return None, html_path

        with mock.patch("core.document_output.save_html_as_pdf", fake_pdf):
            out = self.export(format="pdf", style="summary")
        page = base64.b64decode(out["content_base64"]).decode("utf-8")
        self.assertIn("Doctor Margin Report", page)
        self.assertIn("Totals by doctor", page)
        for word in PATIENTS + ("SB1", "SB10"):
            self.assertNotIn(word, page)

    def test_the_excel_file_reconciles_with_the_screen(self):
        import openpyxl

        out = self.export(format="xlsx", style="detailed")
        ws = openpyxl.load_workbook(io.BytesIO(base64.b64decode(out["content_base64"]))).active
        total = next(r for r in ws.iter_rows(values_only=True) if r[0] == "TOTAL")
        screen = self.report()["totals"]
        self.assertEqual(total[10:16], (screen["amount"], screen["cost"], screen["profit"], 10,
                                        screen["commission"], screen["remaining"]))

    def test_print_keeps_a_copy_and_prints_it(self):
        folder = tempfile.mkdtemp()

        def fake_pdf(html, path, **_kw):
            with open(path, "wb") as fh:
                fh.write(b"%PDF-1.4 fake")
            return path, path.rsplit(".", 1)[0] + ".html"

        def save_bytes(name, raw, folder_=""):
            p = os.path.join(folder, name)
            with open(p, "wb") as fh:
                fh.write(raw)
            return p

        printed = []
        with mock.patch("core.document_output.save_html_as_pdf", fake_pdf), \
                mock.patch("core.saved_files.save_bytes", save_bytes), \
                mock.patch("core.desktop_export_service.print_pdf_on_printer",
                           lambda p, landscape=False: printed.append((p, landscape)) or ""):
            out = self.export(format="print", print_to="printer", style="summary")
        self.assertTrue(out["printed"])
        self.assertEqual(printed, [(out["path"], True)])
        self.assertTrue(out["path"].startswith(folder) and out["path"].endswith(".pdf"))

    def test_nothing_to_export(self):
        out = rep.export_report(self.conn, {"from": "2026-01-01", "to": "2026-01-02", "format": "csv"})
        self.assertFalse(out["ok"])


class TheCommissionIsRemembered(unittest.TestCase):
    def test_prefs(self):
        from core import desktop_ui_prefs as prefs

        path = os.path.join(tempfile.mkdtemp(), "desktop_ui_prefs.json")
        with mock.patch.object(prefs, "_prefs_path", lambda: path):
            self.assertEqual(prefs.load_desktop_ui_prefs()["doctor_margin_pct"], 10.0)
            prefs.save_desktop_ui_prefs({"doctor_margin_pct": 12.5, "doctor_margin_style": "summary"})
            got = prefs.load_desktop_ui_prefs()
            self.assertEqual((got["doctor_margin_pct"], got["doctor_margin_style"]), (12.5, "summary"))
            prefs.save_desktop_ui_prefs({"doctor_margin_pct": 150, "doctor_margin_style": "odd"})
            got = prefs.load_desktop_ui_prefs()
            self.assertEqual((got["doctor_margin_pct"], got["doctor_margin_style"]), (10.0, "summary"))
            self.assertEqual(rep._saved_pct(), 10.0)


if __name__ == "__main__":
    unittest.main()
