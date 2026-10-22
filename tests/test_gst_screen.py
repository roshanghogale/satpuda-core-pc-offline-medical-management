"""The GST Reports screen (core/gst_screen.py, 22 Oct 2026): every report on a made-up store.

The store is built here with the app's own schema: retail bills at 0 / 5 / 12 / 18 / 28 %,
a bill with a discount and a due, a B2B bill to a customer with a GSTIN, wholesale tax
invoices to a buyer in another state (IGST) and in the shop's own state, returns on a retail
and on a wholesale bill, supplier bills from this state, another state and an unregistered
supplier, and goods sent back. No real shop's data.
"""
import base64
import csv
import io
import os
import sqlite3
import sys
import unittest
from decimal import Decimal
from unittest import mock

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from core import gst_reports as gr  # noqa: E402
from core import gst_screen as gs  # noqa: E402
from core.customer_gst import save_customer_gst  # noqa: E402

SHOP_GSTIN = "27AAPFU0939F1ZV"          # Maharashtra

MARG_HEAD_1 = ["Date", "Particulars", "Address", "Type", "Invoice",
               "0%", "0%", "0%", "0%", "5%", "5%", "2.5%", "2.5%", "12%", "12%", "6%", "6%",
               "18%", "18%", "9%", "9%", "28%", "28%", "14%", "14%",
               "Taxable", "Integrated", "Central", "State", "Total", "Add", "Less", "Invoice"]
MARG_HEAD_2 = ["", "", "", "", "No"] + ["Taxable", "IGST", "CGST", "SGST"] * 5 + \
              ["Value", "Tax Amount", "Tax Amount", "Tax Amount", "Tax Amount", "", "", "Amount"]


def _valid(base14: str) -> str:
    chars = "0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ"
    total = 0
    for i, ch in enumerate(base14):
        v = chars.index(ch) * (2 if i % 2 else 1)
        total += v // 36 + v % 36
    return base14 + chars[(36 - total % 36) % 36]


LOCAL_BUYER = _valid("27AAGCB7383J1Z")
GOA_BUYER = _valid("30AAGCG1234K1Z")
WS_LOCAL = _valid("27AABCW5678L1Z")
LOCAL_SUPPLIER = _valid("27AACCL1111M1Z")
KA_SUPPLIER = _valid("29AACCK2222N1Z")


def D(x):
    return Decimal(str(x))


class _Store(unittest.TestCase):
    FROM, TO = "2026-09-01", "2026-09-30"

    def setUp(self):
        from core.db_setup import initialise

        gs._CACHE.clear()
        self.conn = sqlite3.connect(":memory:")
        initialise(self.conn)
        self.profile = {"name": "TEST MEDICAL", "gstin": SHOP_GSTIN, "gst_enabled": 1}
        for p in (
            mock.patch.object(gr, "_online", return_value=False),
            mock.patch("core.customer_gst._online", return_value=False),
            mock.patch("core.pharmacy_profile_io.load_pharmacy_profile", side_effect=lambda *_a, **_k: self.profile),
        ):
            p.start()
            self.addCleanup(p.stop)
        c = self.conn
        c.execute("INSERT INTO customers (id, name, address) VALUES (1, 'WALK IN ONE', 'MAIN ROAD'),"
                  " (2, 'LOCAL DISTRIBUTOR', 'MARKET YARD'), (3, 'GOA AGENCY', 'PANAJI'), (4, 'TOWN WHOLESALE', 'MIDC')")
        c.execute("INSERT INTO suppliers (id, name, gstin, address) VALUES (1, 'LOCAL PHARMA', ?, 'NAGPUR'),"
                  " (2, 'KARNATAKA PHARMA', ?, 'BENGALURU'), (3, 'NO GSTIN TRADERS', '', 'VILLAGE')",
                  (LOCAL_SUPPLIER, KA_SUPPLIER))
        meds = [(1, "TAB ALPHA", "30049099", 12, "Tablet"), (2, "SYRUP BETA", "3004", 5, "Syrup"),
                (3, "SOAP GAMMA", "3401", 18, "Other"), (4, "SALINE DELTA", "3004", 0, "Injection"),
                (5, "CREAM EPSILON", "3304", 28, "Other")]
        for mid, name, hsn, gst, kind in meds:
            c.execute("INSERT INTO medicines (id, name, hsn_code, gst_percent, type, stock_qty, mrp) "
                      "VALUES (?,?,?,?,?,100,10)", (mid, name, hsn, gst, kind))

        def sale(sid, no, day, cust, lines, discount=0.0, rounding=0.0, due=0.0, wholesale=None):
            total = sum(x[2] for x in lines) - discount + rounding
            mode, gstin, pos = wholesale or ("", "", "")
            c.execute("INSERT INTO sales (id, bill_no, bill_date, customer_id, customer_name, customer_address,"
                      " discount, rounding, total_amount, due_amount, invoice_mode, buyer_gstin, place_of_supply)"
                      " VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
                      (sid, f"SCB{no}/FY2026-27", day, cust,
                       {1: "WALK IN ONE", 2: "LOCAL DISTRIBUTOR", 3: "GOA AGENCY", 4: "TOWN WHOLESALE"}[cust],
                       {1: "MAIN ROAD", 2: "MARKET YARD", 3: "PANAJI", 4: "MIDC"}[cust],
                       discount, rounding, total, due, mode, gstin, pos))
            for line in lines:
                mid, qty, amount, gst = line[:4]
                taxable, gst_amt = (line[4], line[5]) if len(line) > 4 else (None, None)
                c.execute("INSERT INTO sales_items (sale_id, medicine_id, qty, rate, gst_percent, amount, taxable,"
                          " gst_amt) VALUES (?,?,?,?,?,?,?,?)", (sid, mid, qty, amount / qty, gst, amount, taxable,
                                                                 gst_amt))

        sale(1, 1, "2026-09-02", 1, [(1, 10, 112.0, 12), (2, 1, 105.0, 5), (4, 1, 50.0, 0)])
        sale(2, 2, "2026-09-03", 1, [(3, 2, 118.0, 18), (5, 1, 128.0, 28)], discount=6.0, rounding=0.4, due=50.0)
        sale(3, 3, "2026-09-05", 2, [(2, 4, 420.0, 5)])
        sale(4, 4, "2026-09-07", 3, [(1, 100, 1120.0, 12, 1000.0, 120.0), (2, 10, 210.0, 5, 200.0, 10.0)],
             wholesale=("wholesale", GOA_BUYER, "30"))
        sale(5, 5, "2026-09-08", 4, [(3, 10, 590.0, 18, 500.0, 90.0)], wholesale=("wholesale", WS_LOCAL, "27"))
        sale(6, 6, "2026-08-30", 1, [(1, 1, 11.2, 12)])                     # before the period
        c.execute("INSERT INTO sales_returns (id, return_no, sale_id, customer_id, return_date, refund_amount,"
                  " discount) VALUES (1,'SR1',1,1,'2026-09-10',56.0,0), (2,'SR2',4,3,'2026-09-12',112.0,0)")
        c.execute("INSERT INTO sales_return_items (return_id, medicine_id, qty, rate, amount)"
                  " VALUES (1,1,5,11.2,56), (2,1,10,11.2,112)")

        def purchase(pid, day, sup, bill_no, subtotal, gst, lines, known=True):
            c.execute("INSERT INTO purchases (id, purchase_no, supplier_id, bill_number, purchase_date, subtotal,"
                      " total_gst, cgst, sgst, total_amount) VALUES (?,?,?,?,?,?,?,?,?,?)",
                      (pid, f"P{pid}", sup, bill_no, day, subtotal, gst, gst / 2, gst / 2, subtotal + gst))
            for mid, qty, rate, taxable, tax in lines:
                c.execute("INSERT INTO purchase_items (purchase_id, medicine_id, qty, rate, gst_pct, taxable,"
                          " gst_amt, item_amount, hsn_code, type) VALUES (?,?,?,?,?,?,?,?,?,?)",
                          (pid, mid, qty, taxable / qty, rate, taxable if known else None, tax if known else None,
                           taxable + tax, {1: "30049099", 2: "3004", 3: "3401", 4: "3004", 5: "3304"}[mid],
                           {1: "Tablet"}.get(mid, "Other")))

        purchase(1, "2026-09-01", 1, "LP-1", 1100.0, 150.0, [(1, 100, 12, 800.0, 96.0), (3, 10, 18, 300.0, 54.0)])
        purchase(2, "2026-09-04", 2, "KA-1", 500.0, 25.0, [(2, 50, 5, 500.0, 25.0)])
        purchase(3, "2026-09-06", 3, "NG-1", 200.0, 0.0, [(4, 10, 0, 200.0, 0.0)])
        purchase(4, "2026-09-09", 1, "LP-2", 400.0, 48.0, [(1, 40, 12, 250.0, 30.0), (1, 30, 12, 150.0, 18.0)],
                 known=False)
        c.execute("INSERT INTO purchase_returns (id, return_no, purchase_id, supplier_id, return_date, refund_amount)"
                  " VALUES (1,'PR1',1,1,'2026-09-15',89.6)")
        c.execute("INSERT INTO purchase_return_items (return_id, medicine_id, qty, rate, amount)"
                  " VALUES (1,1,10,8,80)")
        c.commit()
        save_customer_gst(c, 2, LOCAL_BUYER, "LOCAL DISTRIBUTOR", since="2026-01-01")

    def rep(self, key, **kw):
        return gs.build(self.conn, key, self.FROM, self.TO, **kw)

    @staticmethod
    def col(sec, name):
        return sec["columns"].index(name)

    def row(self, sec, no):
        return next(r for r in sec["rows"] if r[4] == no)


class TheSaleRegister(_Store):
    def test_one_row_a_bill_and_the_totals_are_the_bills(self):
        sec = self.rep("C")["sections"][0]
        self.assertEqual(["SCB1", "SCB2", "SCB3", "SCB4", "SCB5"], [r[4] for r in sec["rows"]])
        inv = self.col(sec, "Invoice Amount")
        bills = self.conn.execute("SELECT SUM(total_amount) FROM sales WHERE bill_date >= ?", (self.FROM,)).fetchone()[0]
        self.assertAlmostEqual(bills, sec["total"][inv], places=2)
        reg = gr.build(self.conn, self.FROM, self.TO)["sales_register"]
        self.assertAlmostEqual(sum(r["taxable"] for r in reg), sec["total"][self.col(sec, "Taxable Value")], places=2)
        for r in sec["rows"]:
            taxable = sum(D(r[j]) for j, c in enumerate(sec["columns"]) if c.endswith(" Taxable"))
            self.assertEqual(taxable, D(r[self.col(sec, "Taxable Value")]))
            tax = D(r[self.col(sec, "Total Tax Amount")])
            self.assertEqual(D(r[inv]), taxable + tax + D(r[self.col(sec, "Add")]) - D(r[self.col(sec, "Less")]))

    def test_rate_buckets_and_the_printed_discount(self):
        sec = self.rep("C")["sections"][0]
        r = self.row(sec, "SCB2")
        g = gr.printed_bill_gst([(118.0, 18), (128.0, 28)], 6.0)
        self.assertEqual(float(g.lines[0].taxable), r[self.col(sec, "18% Taxable")])
        self.assertEqual(float(g.lines[1].taxable), r[self.col(sec, "28% Taxable")])
        self.assertEqual(0.0, r[self.col(sec, "12% Taxable")])
        self.assertEqual(("CREDIT", "MAIN ROAD"), (r[3], r[2]))
        self.assertEqual(0.4, r[self.col(sec, "Add")])
        r1 = self.row(sec, "SCB1")
        self.assertEqual((50.0, 100.0, 100.0), (r1[self.col(sec, "0% Taxable")], r1[self.col(sec, "5% Taxable")],
                                                r1[self.col(sec, "12% Taxable")]))
        self.assertEqual("CASH", r1[3])

    def test_a_wholesale_buyer_in_another_state_is_igst(self):
        sec = self.rep("C")["sections"][0]
        r = self.row(sec, "SCB4")
        self.assertEqual((1000.0, 120.0, 0.0, 0.0), tuple(r[self.col(sec, c)] for c in
                                                          ("12% Taxable", "12% IGST", "12% CGST 6%", "12% SGST 6%")))
        self.assertEqual((130.0, 0.0, 0.0), (r[self.col(sec, "Integrated Tax Amount")],
                                             r[self.col(sec, "Central Tax Amount")],
                                             r[self.col(sec, "State Tax Amount")]))
        local = self.row(sec, "SCB5")
        self.assertEqual((0.0, 45.0, 45.0), tuple(local[self.col(sec, c)] for c in
                                                  ("18% IGST", "18% CGST 9%", "18% SGST 9%")))


class TheMargLayout(_Store):
    def _csv(self, key):
        res = gs.export(self.conn, {"report": key, "from": self.FROM, "to": self.TO, "format": "csv"})
        self.assertTrue(res["ok"], res)
        return list(csv.reader(io.StringIO(base64.b64decode(res["content_base64"]).decode("utf-8-sig")))), res

    def test_the_sale_csv_is_margs_upload_file(self):
        lines, res = self._csv("C")
        self.assertEqual("GST_Sale_Upload_Excel_File_01-09-2026_30-09-2026.csv", res["filename"])
        self.assertEqual([["TEST MEDICAL"], [SHOP_GSTIN], ["Date : 01-09-2026 - 30-09-2026"], ["Sale GST"], []],
                         lines[:5])
        self.assertEqual(MARG_HEAD_1, lines[5])
        self.assertEqual(MARG_HEAD_2, lines[6])
        self.assertEqual(33, len(lines[7]))
        self.assertEqual("02-09-2026", lines[7][0])
        self.assertEqual("", lines[7][6])                      # no IGST on a local bill: left empty, as Marg
        self.assertEqual([], lines[-2])
        self.assertEqual(["", "", "", "", ""], lines[-1][:5])
        self.assertEqual(33, len(lines[-1]))

    def test_the_purchase_csv_and_the_excel_file(self):
        from openpyxl import load_workbook

        lines, _ = self._csv("D")
        self.assertEqual(["Purchase GST"], lines[3])
        self.assertEqual((MARG_HEAD_1, MARG_HEAD_2), (lines[5], lines[6]))
        res = gs.export(self.conn, {"report": "sale", "from": self.FROM, "to": self.TO, "format": "xlsx"})
        ws = load_workbook(io.BytesIO(base64.b64decode(res["content_base64"]))).active
        self.assertEqual(("TEST MEDICAL", SHOP_GSTIN, "Date : 01-09-2026 - 30-09-2026", "Sale GST"),
                         tuple(ws.cell(row=i, column=1).value for i in range(1, 5)))
        self.assertEqual(MARG_HEAD_1, [ws.cell(row=6, column=j).value for j in range(1, 34)])
        self.assertEqual(MARG_HEAD_2, [ws.cell(row=7, column=j).value or "" for j in range(1, 34)])
        self.assertEqual(50, ws.cell(row=8, column=6).value)              # a number, not text

    def test_returns_registers_use_the_same_layout(self):
        for key, heading in (("E", "Sale Return GST"), ("F", "Purchase Return GST")):
            lines, _ = self._csv(key)
            self.assertEqual([heading], lines[3])
            self.assertEqual((MARG_HEAD_1, MARG_HEAD_2), (lines[5], lines[6]))

    def test_another_rate_gets_its_own_block_after_the_marg_ones(self):
        self.conn.execute("UPDATE sales_items SET gst_percent=40 WHERE sale_id=2 AND medicine_id=5")
        self.conn.commit()
        sec = self.rep("C")["sections"][0]
        self.assertEqual(["40%", "40%", "20%", "20%"], sec["header"][0][25:29])
        self.assertEqual(MARG_HEAD_1[:25], sec["header"][0][:25])


class ThePurchaseSide(_Store):
    def test_purchase_register_rates_igst_and_a_bill_without_line_values(self):
        sec = self.rep("D")["sections"][0]
        ka = self.row(sec, "KA-1")
        self.assertEqual((500.0, 25.0, 0.0), (ka[self.col(sec, "5% Taxable")], ka[self.col(sec, "5% IGST")],
                                              ka[self.col(sec, "5% CGST 2.5%")]))
        lp2 = self.row(sec, "LP-2")                         # one rate on every line: the footer goes there
        self.assertEqual((400.0, 24.0, 24.0), (lp2[self.col(sec, "12% Taxable")], lp2[self.col(sec, "12% CGST 6%")],
                                               lp2[self.col(sec, "12% SGST 6%")]))
        self.assertEqual("URD", self.row(sec, "NG-1")[3])
        itc = gr.build(self.conn, self.FROM, self.TO)["purchase_register"]
        self.assertAlmostEqual(sum(r["total_gst"] for r in itc), sec["total"][self.col(sec, "Total Tax Amount")],
                               places=2)

    def test_gstr2_b2b_leaves_out_unregistered_suppliers(self):
        sec = self.rep("L")["sections"][0]
        self.assertEqual({"LP-1", "KA-1", "LP-2"}, {r[2] for r in sec["rows"]})
        lp1 = [r for r in sec["rows"] if r[2] == "LP-1"]
        self.assertEqual([(12, 800.0, 48.0), (18, 300.0, 27.0)], [(r[8], r[9], r[11]) for r in lp1])
        self.assertEqual(150.0 + 25.0 + 48.0, round(sum(sec["total"][j] for j in (10, 11, 12)), 2))

    def test_debit_notes_and_purchase_hsn_net_of_returns(self):
        sec = self.rep("N")["sections"][0]
        self.assertEqual([("PR1", "LP-1", 12, 80.0, 0.0, 4.8, 4.8)],
                         [(r[2], r[4], r[9], r[10], r[11], r[12], r[13]) for r in sec["rows"]])
        bought = self.rep("M")["sections"][0]
        back = self.rep("O")["sections"][0]
        net = self.rep("P")["sections"][0]
        t = self.col(net, "Taxable Value")
        self.assertAlmostEqual(bought["total"][t] - back["total"][t], net["total"][t], places=2)
        tab = next(r for r in net["rows"] if r[0] == "30049099")
        self.assertEqual(("TBS", 90.0, 720.0), (tab[2], tab[3], tab[6]))
        self.assertIn("1 supplier bill", " ".join(self.rep("M")["notes"]))     # LP-2: no line values

    def test_party_wise(self):
        rows = {r[0]: r for r in self.rep("Q")["sections"][0]["rows"]}
        lp = rows["LOCAL PHARMA"]
        self.assertEqual((2, 1500.0, 198.0, 80.0, 9.6, 1420.0, 188.4), (lp[2], lp[3], lp[7], lp[9], lp[10], lp[11],
                                                                         lp[12]))


class TheGstr1Tables(_Store):
    def test_b2b_has_the_wholesale_and_registered_bills(self):
        sec = self.rep("b2b")["sections"][0]
        by = {(r[2], r[10]): r for r in sec["rows"]}
        self.assertEqual({("SCB3", 5), ("SCB4", 12), ("SCB4", 5), ("SCB5", 18)}, set(by))
        goa = by[("SCB4", 12)]
        self.assertEqual((GOA_BUYER, "30-Goa", 1000.0, 120.0, 0.0), (goa[0], goa[5], goa[11], goa[13], goa[14]))

    def test_b2cs_by_rate_and_place_net_of_returns_and_nil_apart(self):
        rep = self.rep("G")
        b2cs, nil = rep["sections"]
        rates = {r[3]: r for r in b2cs["rows"]}
        self.assertNotIn(0, rates)
        self.assertEqual(100.0 - 50.0, rates[12][4])               # SCB1 12% less SR1 (5 of 10 tablets)
        self.assertEqual("27-Maharashtra", rates[12][1])
        self.assertEqual(50.0, nil["total"][1])

    def test_credit_note_to_a_registered_buyer_in_another_state(self):
        cdnr = self.rep("I")["sections"][0]
        self.assertEqual([("SR2", 100.0, 12.0, 0.0)], [(r[2], r[11], r[14], r[15]) for r in cdnr["rows"]])

    def test_hsn_sale_less_sale_return_is_table_12(self):
        sold = self.rep("H")["sections"][0]
        back = self.rep("J")["sections"][0]
        k = self.rep("K")["sections"]
        t = self.col(sold, "Taxable Value")
        filed = k[0]["total"][t] + k[1]["total"][t]
        self.assertAlmostEqual(sold["total"][t] - back["total"][t], filed, places=2)
        self.assertAlmostEqual(filed, k[2]["total"][t], places=2)
        q = self.col(sold, "Total Quantity")
        tab = {r[0]: r for r in k[2]["rows"]}["30049099"]
        self.assertEqual(110.0 - 15.0, tab[q])

    def test_the_club_and_3b_agree(self):
        club = self.rep("T")["sections"][0]
        net = club["rows"][-1]
        t3b = gr.build(self.conn, self.FROM, self.TO)["gstr3b"]["3.1(a) Outward taxable supplies"]
        self.assertEqual((t3b["taxable"], t3b["igst"], t3b["cgst"], t3b["sgst"]), tuple(net[3:7]))
        self.assertEqual(120.0 + 10.0 - 12.0, t3b["igst"])
        b = self.rep("B")["sections"][0]["rows"][0]
        self.assertEqual(t3b["igst"], b[2])

    def test_portal_json_puts_zero_rated_b2c_in_nil(self):
        js = gr.gstr1_json(gr.load_period(self.conn, self.FROM, self.TO))
        self.assertEqual([{"sply_ty": "INTRAB2C", "expt_amt": 0, "nil_amt": 50.0, "ngsup_amt": 0}], js["nil"]["inv"])
        self.assertFalse(any(r["rt"] == 0 for r in js["b2cs"]))
        self.assertEqual({LOCAL_BUYER, GOA_BUYER, WS_LOCAL}, {x["ctin"] for x in js["b2b"]})


class CompositionAndAnnual(_Store):
    def test_composition_tax_is_one_percent_of_net_taxable_turnover(self):
        rows = self.rep("R")["sections"][0]["rows"]
        self.assertEqual(1, len(rows))
        r = rows[0]
        sales = D(r[2]) - D(r[3])
        self.assertEqual(D(r[6]), sales - (D(r[4]) - D(r[5])))
        self.assertEqual(D(r[10]), (D(r[6]) / 100).quantize(Decimal("0.01")))
        two = self.rep("R", comp_rate=2)["sections"][0]["rows"][0]
        self.assertEqual(2, two[7])

    def test_gstr9_table_4n_is_the_3b_outward_figure(self):
        t4 = self.rep("V")["sections"][0]
        n = t4["rows"][-1]
        t3b = gr.build(self.conn, self.FROM, self.TO)["gstr3b"]["3.1(a) Outward taxable supplies"]
        self.assertEqual((t3b["taxable"], t3b["cgst"], t3b["sgst"], t3b["igst"]), tuple(n[1:5]))

    def test_gstr4_inward_registered_and_not(self):
        secs = self.rep("W")["sections"]
        self.assertEqual({"LOCAL PHARMA", "KARNATAKA PHARMA"}, {r[1] for r in secs[0]["rows"]})
        self.assertEqual(["NO GSTIN TRADERS"], [r[0] for r in secs[1]["rows"]])


class EveryReport(_Store):
    def test_every_report_builds_and_its_totals_are_its_rows(self):
        for item in gs.full_catalogue() + gs.catalogue()["reports"] + [{"key": "check_bills"}]:
            rep = self.rep(item["key"])
            self.assertTrue(rep["sections"], item["key"])
            for sec in rep["sections"]:
                for r in sec["rows"]:
                    self.assertEqual(len(sec["columns"]), len(r), (item["key"], sec["title"]))
                for j in sec["sum_cols"]:
                    if sec["total"]:
                        self.assertAlmostEqual(sum(float(r[j] or 0) for r in sec["rows"]), sec["total"][j], places=2,
                                               msg=(item["key"], sec["title"], sec["columns"][j]))

    def test_every_report_goes_to_csv_excel_and_pdf(self):
        for item in gs.full_catalogue() + gs.catalogue()["reports"]:
            for fmt in ("csv", "xlsx"):
                res = gs.export(self.conn, {"report": item["key"], "from": self.FROM, "to": self.TO, "format": fmt})
                self.assertTrue(res["ok"] and res["content_base64"], (item["key"], fmt))
        pdf = gs.export(self.conn, {"report": "B", "from": self.FROM, "to": self.TO, "format": "pdf"})
        self.assertTrue(pdf["ok"], pdf)
        js = gs.export(self.conn, {"report": "A", "from": self.FROM, "to": self.TO, "format": "json"})
        self.assertTrue(js["filename"].endswith(".json"))

    def test_the_letters_are_margs(self):
        letters = {i["letter"]: i["title"] for i in gs.full_catalogue() if i["letter"]}
        self.assertEqual("ABCDEFGHIJKLMNOPQRSTUVW", "".join(sorted(letters)))
        self.assertEqual("Sale (GSTR-1) Detail", letters["C"])
        self.assertEqual("HSN Sale - Sale Return (GSTR-1) hsn", letters["K"])

    def test_a_bill_of_supply_heading_changes_nothing(self):
        before = self.rep("C")["sections"][0]["rows"]
        gs._CACHE.clear()
        self.profile = {**self.profile, "bill_heading": "Bill of Supply", "invoice_heading": "Bill of Supply"}
        self.assertEqual(before, self.rep("C")["sections"][0]["rows"])

    def test_nothing_is_written(self):
        before = self.conn.total_changes
        for item in gs.full_catalogue() + gs.catalogue(self.conn)["reports"]:
            self.rep(item["key"])
        self.assertEqual(before, self.conn.total_changes)


if __name__ == "__main__":
    unittest.main()


class TheScreenList(_Store):
    """Owner, 22 Oct 2026: Marg's names, sorted so they do not overwhelm -- Monthly filing,
    Registers for your CA, More reports -- and retail / wholesale / composition each their own."""

    REGISTERS = ["Sale (GSTR-1) (Detail)", "Sale Return (GSTR-1) (Detail)", "Purchase (GSTR-2) (Detail)",
                 "Purchase Return (GSTR-2) (Detail)"]

    def lists(self, start=None, end=None, wholesale=False):
        cat = gs.catalogue(self.conn, start or self.FROM, end or self.TO, wholesale)
        by = {}
        for i in cat["reports"]:
            by.setdefault(i["section"], []).append(i["title"])
        return cat, by

    def test_a_retail_month_with_registered_buyers(self):
        cat, by = self.lists()
        self.assertEqual("retail", cat["kind"])
        self.assertEqual(["GSTR-3b", "Sale (GSTR-1) (b2cs)", "HSN Sale - Sale Return (GSTR-1) (hsn)",
                          "Sale (GSTR-1) (b2b)", "Sale Return (GSTR-1) (cdnr)"], by["Monthly filing"])
        self.assertEqual(self.REGISTERS, by["Registers for your CA"])
        self.assertEqual(["HSN Sale (GSTR-1) (hsn)", "HSN Sale Return (GSTR-1) (hsn)", "Purchase (GSTR-2) (b2b)",
                          "HSN Purchase (GSTR-2) (hsn)", "Purchase Return (GSTR-2) (dbnr)",
                          "HSN Purchase Rtn (GSTR-2) (hsn)", "HSN Purchase - Purchase Return (GSTR-2) (hsn)",
                          "Purchase (GSTR-2) (Party Wise)", "Sale GSTR-1 (club)", "Purchase GSTR-2 (club)",
                          "Annual GSTR"], by["More reports"])
        items = {i["key"]: i for i in cat["reports"]}
        self.assertEqual((5, 2, 3, 1), (items["sale"]["count"], items["sale_return"]["count"],
                                        items["b2b"]["count"], items["cdnr"]["count"]))
        self.assertTrue(items["b2cs"]["json"] and items["sale"]["json"])
        self.assertFalse(items["purchase"]["json"] or items["gstr3b"]["json"])

    def test_a_retail_month_without_registered_buyers_folds_b2b_and_cdnr(self):
        cat, by = self.lists("2026-08-01", "2026-08-31")
        self.assertEqual(["GSTR-3b", "Sale (GSTR-1) (b2cs)", "HSN Sale - Sale Return (GSTR-1) (hsn)"],
                         by["Monthly filing"])
        self.assertEqual(["Sale (GSTR-1) (b2b)", "Sale Return (GSTR-1) (cdnr)"], by["More reports"][:2])

    def test_a_wholesale_store(self):
        cat, by = self.lists(wholesale=True)
        self.assertEqual("wholesale", cat["kind"])
        self.assertEqual(["GSTR-3b", "Sale (GSTR-1) (b2b)", "Sale Return (GSTR-1) (cdnr)",
                          "HSN Sale - Sale Return (GSTR-1) (hsn)", "Sale (GSTR-1) (b2cs)"], by["Monthly filing"])
        self.assertEqual(self.REGISTERS, by["Registers for your CA"])
        self.assertEqual(["Purchase (GSTR-2) (b2b)", "Purchase (GSTR-2) (Party Wise)",
                          "Purchase Return (GSTR-2) (dbnr)"], by["More reports"][:3])
        # only registered buyers in these two days: b2cs folds away, still reachable
        _, only_b2b = self.lists("2026-09-07", "2026-09-08", wholesale=True)
        self.assertNotIn("Sale (GSTR-1) (b2cs)", only_b2b["Monthly filing"])
        self.assertEqual("Sale (GSTR-1) (b2cs)", only_b2b["More reports"][0])

    def test_a_composition_dealer(self):
        from core.bill_heading import save_bill_heading

        save_bill_heading(self.conn, "bill_of_supply", "composition")
        for ws in (False, True):
            cat, by = self.lists(wholesale=ws)
            self.assertEqual("composition", cat["kind"])
            self.assertEqual(["Sale - Sale Return Composition GST", "Purchase Composition GST"], by["Monthly filing"])
            self.assertEqual(self.REGISTERS, by["Registers for your CA"])
            self.assertEqual(["Annual GSTR4"], by["More reports"])
            self.assertFalse(any(i["json"] for i in cat["reports"]))
        save_bill_heading(self.conn, "bill_of_supply", "exempt")
        self.assertEqual("retail", self.lists()[0]["kind"])

    def test_every_listed_report_builds_under_its_marg_name(self):
        for ws in (False, True):
            for i in self.lists(wholesale=ws)[0]["reports"]:
                self.assertEqual(i["title"], self.rep(i["key"])["title"])
        self.assertEqual(self.rep("C")["sections"], self.rep("sale")["sections"])

    def test_gstr1_summary_has_its_sections_and_matches_the_tables(self):
        rep = self.rep("gstr1")
        titles = [s["title"] for s in rep["sections"]]
        self.assertEqual(["B2B (buyers with GSTIN)", "B2CS (other sales, by rate)", "Nil rated (0%)", "HSN (B2B)",
                          "HSN (B2C)", "Credit notes (buyers with GSTIN)", "Documents issued"], titles)
        b2b = rep["sections"][0]
        self.assertEqual(self.rep("b2b")["sections"][0]["total"], b2b["total"])

    def test_hsn_summary_is_sales_less_returns(self):
        net = self.rep("hsn")["sections"][0]
        k = self.rep("K")["sections"]
        t = self.col(net, "Taxable Value")
        self.assertAlmostEqual(k[0]["total"][t] + k[1]["total"][t], net["total"][t], places=2)
