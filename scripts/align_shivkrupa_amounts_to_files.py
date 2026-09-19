"""
Force Shivkrupa sales / purchase / return / payment amounts to match eVital export files.

Source of truth:
  - Sales bill total + paid  -> vm9xzvqmyn.csv (aligned 1:1 by import order)
  - Purchase bill total+paid -> xpjfe4h0im.csv (matched by bill_number)
  - Sales returns            -> Sales_Return_Register.xlsx
  - Purchase returns         -> Purchase_Return_Register.xlsx (non-Draft)

Then recalculates customer + supplier dues.
"""
from __future__ import annotations

import argparse
import csv
import io
import os
import shutil
import sqlite3
import sys
from collections import defaultdict
from contextlib import redirect_stdout
from datetime import datetime
from pathlib import Path

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

import core.online_guard as _og  # noqa: E402

_og.ensure_can_mutate = lambda: None  # type: ignore

from core.customer_service import (  # noqa: E402
    get_or_create_customer,
    recalculate_customer_due,
)
from core.purchase_service import recalculate_supplier_due  # noqa: E402

try:
    import openpyxl
except ImportError as exc:  # pragma: no cover
    raise SystemExit("openpyxl required") from exc

DB_PATH = os.path.join(
    ROOT, "config", "stores", "Store_Shivkrupa_Medical_General_Store", "veterinary.db"
)
SALES_CSV = r"c:\Users\win10\Downloads\vm9xzvqmyn.csv"
PURCH_CSV = r"c:\Users\win10\Downloads\xpjfe4h0im.csv"
SR_XLSX = r"c:\Users\win10\Downloads\Sales_Return_Register.xlsx"
PR_XLSX = r"c:\Users\win10\Downloads\Purchase_Return_Register.xlsx"
REPORT = os.path.join(
    ROOT,
    "config",
    "stores",
    "Store_Shivkrupa_Medical_General_Store",
    "align_amounts_to_files_report.txt",
)


def _log(msg: str) -> None:
    print(msg, flush=True)


def _f(v) -> float:
    try:
        return round(float(v or 0), 2)
    except Exception:
        return 0.0


def _parse_dt(raw: str) -> datetime:
    raw = (raw or "").strip()
    for fmt in (
        "%Y-%m-%d %H:%M:%S.%f",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d",
        "%d-%m-%y",
        "%d-%m-%Y",
    ):
        try:
            val = raw if "%f" in fmt else raw.split(".")[0]
            return datetime.strptime(val, fmt)
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(raw[:19])
    except ValueError:
        return datetime.min


def _parse_date(raw: str) -> str:
    return _parse_dt(raw).date().isoformat()


def _split_modes(total_paid: float, method: str, old_cash: float, old_online: float) -> tuple[float, float]:
    method = (method or "").strip().lower()
    if total_paid <= 0:
        return 0.0, 0.0
    if "upi" in method or "online" in method or "card" in method or "neft" in method:
        return 0.0, total_paid
    if "credit" in method or "due" in method:
        return 0.0, 0.0
    # Keep prior split ratio when possible
    old = old_cash + old_online
    if old > 0.01:
        cash = round(total_paid * (old_cash / old), 2)
        return cash, round(total_paid - cash, 2)
    return total_paid, 0.0


def align_sales(conn: sqlite3.Connection) -> dict:
    groups: dict[str, list[dict]] = defaultdict(list)
    with open(SALES_CSV, encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            groups[str(row.get("order_number") or "").strip()].append(row)
    keys = sorted(
        groups, key=lambda k: (_parse_dt(groups[k][0].get("created_date", "")), k)
    )
    sales = conn.execute(
        """
        SELECT id, total_amount, amount_paid, cash_paid, online_paid, due_amount
        FROM sales
        WHERE COALESCE(deleted,0)=0 AND COALESCE(is_autosave,0)=0
        ORDER BY id
        """
    ).fetchall()
    if len(keys) != len(sales):
        raise RuntimeError(f"Sales count mismatch CSV={len(keys)} DB={len(sales)}")

    updated = 0
    samples: list[str] = []
    for i, (sid, tot, paid, cash, online, due) in enumerate(sales):
        head = groups[keys[i]][0]
        csv_total = _f(head.get("total"))
        csv_paid = _f(head.get("paid_amount"))
        method = head.get("payment_method") or ""
        # Credit / unpaid in source
        if csv_paid <= 0 and csv_total > 0 and "credit" in method.lower():
            csv_paid = 0.0
        new_due = round(max(0.0, csv_total - csv_paid), 2)
        new_credit = round(max(0.0, csv_paid - csv_total), 2)
        new_cash, new_online = _split_modes(
            csv_paid, method, _f(cash), _f(online)
        )
        if (
            abs(_f(tot) - csv_total) <= 0.009
            and abs(_f(paid) - csv_paid) <= 0.009
            and abs(_f(due) - new_due) <= 0.009
        ):
            continue
        conn.execute(
            """
            UPDATE sales SET
              total_amount=?,
              amount_paid=?, cash_paid=?, online_paid=?,
              due_amount=?, credit_amount=?,
              total_due=?,
              account_cleared=CASE WHEN ?<=0.009 THEN 1 ELSE 0 END,
              bill_cleared=CASE WHEN ?<=0.009 THEN 1 ELSE 0 END
            WHERE id=?
            """,
            (
                csv_total,
                csv_paid,
                new_cash,
                new_online,
                new_due,
                new_credit,
                new_due,
                new_due,
                new_due,
                sid,
            ),
        )
        updated += 1
        if len(samples) < 20:
            samples.append(
                f"sale {sid}: total {_f(tot)}->{csv_total} paid {_f(paid)}->{csv_paid} "
                f"due {_f(due)}->{new_due} [{method}]"
            )
    return {"updated": updated, "total": len(sales), "samples": samples}


def align_purchases(conn: sqlite3.Connection) -> dict:
    by_bill: dict[str, dict] = {}
    with open(PURCH_CSV, encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            bn = (row.get("bill_no") or "").strip()
            if bn and bn not in by_bill:
                by_bill[bn] = row

    rows = conn.execute(
        """
        SELECT id, bill_number, final_amount, total_amount,
               amount_paid_at_entry, amount_paid, need_to_pay, due_amount,
               cash_paid_at_entry, online_paid_at_entry
        FROM purchases
        WHERE COALESCE(deleted,0)=0 AND COALESCE(is_autosave,0)=0
        """
    ).fetchall()
    updated = unmatched = 0
    samples: list[str] = []
    for (
        pid,
        bill_number,
        final_amt,
        total_amt,
        entry_paid,
        amount_paid,
        need_to_pay,
        due_amount,
        cash_entry,
        online_entry,
    ) in rows:
        src = by_bill.get(str(bill_number or "").strip())
        if not src:
            unmatched += 1
            continue
        csv_total = _f(src.get("total") if src.get("total") not in (None, "") else src.get("amount"))
        csv_paid = _f(src.get("paid_amount"))
        method = src.get("payment_method") or ""
        new_due = round(max(0.0, csv_total - csv_paid), 2)
        new_credit = round(max(0.0, csv_paid - csv_total), 2)
        new_cash, new_online = _split_modes(
            csv_paid, method, _f(cash_entry), _f(online_entry)
        )
        fa = _f(final_amt if final_amt is not None else total_amt)
        if (
            abs(fa - csv_total) <= 0.009
            and abs(_f(entry_paid) - csv_paid) <= 0.009
            and abs(_f(due_amount) - new_due) <= 0.009
        ):
            continue
        conn.execute(
            """
            UPDATE purchases SET
              total_amount=?, final_amount=?, need_to_pay=?,
              amount_paid_at_entry=?, amount_paid=?,
              cash_paid_at_entry=?, online_paid_at_entry=?,
              due_amount=?, due=?, total_due=?,
              credit_amount=?, current_credit=?,
              account_cleared=CASE WHEN ?<=0.009 THEN 1 ELSE 0 END,
              bill_cleared=CASE WHEN ?<=0.009 THEN 1 ELSE 0 END
            WHERE id=?
            """,
            (
                csv_total,
                csv_total,
                csv_total,
                csv_paid,
                csv_paid,
                new_cash,
                new_online,
                new_due,
                new_due,
                new_due,
                new_credit,
                new_credit,
                new_due,
                new_due,
                pid,
            ),
        )
        updated += 1
        if len(samples) < 20:
            samples.append(
                f"purch {pid} bill={bill_number}: total {fa}->{csv_total} "
                f"paid {_f(entry_paid)}->{csv_paid} due {_f(due_amount)}->{new_due}"
            )
    return {
        "updated": updated,
        "matched": len(rows) - unmatched,
        "unmatched": unmatched,
        "samples": samples,
    }


def align_sales_returns(conn: sqlite3.Connection) -> dict:
    wb = openpyxl.load_workbook(SR_XLSX, read_only=True, data_only=True)
    ws = wb.active
    headers = None
    groups: dict[str, dict] = {}
    for row in ws.iter_rows(values_only=True):
        if row and row[0] == "Return Bill No.":
            headers = list(row)
            continue
        if not headers or not row or row[0] is None:
            continue
        key = str(row[0]).strip()
        if not key or key.lower() == "total":
            continue
        amt = _f(row[14])
        if key not in groups:
            groups[key] = {
                "amount": 0.0,
                "date": _parse_date(str(row[1] or "")),
                "customer": str(row[2] or "").strip() or "COUNTER SALE",
                "mode": str(row[4] or "").strip(),
            }
        groups[key]["amount"] = round(groups[key]["amount"] + amt, 2)
    wb.close()

    # Existing returns ordered by date/id — map by refund+date first
    existing = conn.execute(
        """
        SELECT id, return_no, sale_id, customer_id, return_date, refund_amount
        FROM sales_returns
        WHERE COALESCE(deleted,0)=0
        ORDER BY return_date, id
        """
    ).fetchall()
    # Known mapping from prior import (file # -> existing rows excluding summary #3)
    # Match by (date, amount) then create missing.
    used: set[int] = set()
    updated = created = 0
    notes: list[str] = []

    for key in sorted(groups, key=lambda k: int(k) if k.isdigit() else 9999):
        g = groups[key]
        target = g["amount"]
        # Skip empty
        if target <= 0:
            continue
        match = None
        for row in existing:
            rid, rno, sale_id, cust_id, rdate, refund = row
            if rid in used:
                continue
            if abs(_f(refund) - target) <= 0.02 and str(rdate) == g["date"]:
                match = row
                break
        if match is None:
            for row in existing:
                rid, rno, sale_id, cust_id, rdate, refund = row
                if rid in used:
                    continue
                if abs(_f(refund) - target) <= 0.02:
                    match = row
                    break
        if match:
            used.add(int(match[0]))
            if abs(_f(match[5]) - target) > 0.009 or str(match[4]) != g["date"]:
                conn.execute(
                    "UPDATE sales_returns SET refund_amount=?, return_date=? WHERE id=?",
                    (target, g["date"], match[0]),
                )
                updated += 1
                notes.append(f"updated {match[1]} -> {target} on {g['date']}")
            else:
                notes.append(f"ok {match[1]} = {target}")
            continue

        # Create missing return (e.g. summary return #3)
        cust_id = get_or_create_customer(conn, g["customer"], "", "")
        # Try find a sale for that customer on/before date to attach
        sale_row = conn.execute(
            """
            SELECT id FROM sales
            WHERE customer_id=? AND COALESCE(deleted,0)=0
              AND bill_date<=?
            ORDER BY bill_date DESC, id DESC LIMIT 1
            """,
            (cust_id, g["date"]),
        ).fetchone()
        sale_id = int(sale_row[0]) if sale_row else None
        # Pick a free return_no
        return_no = f"SR{key}" if str(key).isdigit() else "SRX"
        exists = conn.execute(
            "SELECT 1 FROM sales_returns WHERE return_no=? LIMIT 1", (return_no,)
        ).fetchone()
        if exists:
            nxt = conn.execute("SELECT COALESCE(MAX(id),0)+1 FROM sales_returns").fetchone()[0]
            return_no = f"SR{nxt}"
        conn.execute(
            """
            INSERT INTO sales_returns
              (return_no, sale_id, customer_id, return_date, refund_amount, discount, reason,
               deleted, version)
            VALUES (?,?,?,?,?,0,?,0,1)
            """,
            (
                return_no,
                sale_id,
                cust_id,
                g["date"],
                target,
                f"Imported from Sales Return Register #{key}",
            ),
        )
        created += 1
        notes.append(
            f"created {return_no} customer={g['customer']} amt={target} date={g['date']}"
        )

    return {
        "file_returns": len(groups),
        "updated": updated,
        "created": created,
        "notes": notes,
        "file_total": round(sum(g["amount"] for g in groups.values()), 2),
    }


def align_purchase_returns(conn: sqlite3.Connection) -> dict:
    wb = openpyxl.load_workbook(PR_XLSX, read_only=True, data_only=True)
    ws = wb.active
    headers = None
    groups: dict[str, dict] = {}
    for row in ws.iter_rows(values_only=True):
        if row and row[0] == "Return ID / Voucher No.":
            headers = list(row)
            continue
        if not headers or not row or not row[0]:
            continue
        key = str(row[0]).strip()
        if not key or key.lower() == "total":
            continue
        status = str(row[4] or "").strip()
        if status.lower() == "draft":
            continue
        amt = _f(row[17])
        if key not in groups:
            groups[key] = {
                "amount": 0.0,
                "date": _parse_date(str(row[1] or "")),
                "supplier": str(row[2] or "").strip(),
                "status": status,
            }
        groups[key]["amount"] = round(groups[key]["amount"] + amt, 2)
    wb.close()

    existing = conn.execute(
        """
        SELECT id, return_no, refund_amount, return_date
        FROM purchase_returns WHERE COALESCE(deleted,0)=0
        ORDER BY id
        """
    ).fetchall()
    updated = 0
    notes: list[str] = []
    used: set[int] = set()
    for key, g in groups.items():
        target = g["amount"]
        match = None
        for row in existing:
            if row[0] in used:
                continue
            if abs(_f(row[2]) - target) <= 0.02:
                match = row
                break
        if not match:
            notes.append(f"MISSING purchase return {key} amt={target} {g['supplier']}")
            continue
        used.add(int(match[0]))
        if abs(_f(match[2]) - target) > 0.009:
            conn.execute(
                "UPDATE purchase_returns SET refund_amount=?, return_date=? WHERE id=?",
                (target, g["date"], match[0]),
            )
            updated += 1
            notes.append(f"updated {match[1]} -> {target}")
        else:
            notes.append(f"ok {match[1]} = {target}")
    return {
        "file_returns": len(groups),
        "updated": updated,
        "notes": notes,
        "file_total": round(sum(g["amount"] for g in groups.values()), 2),
    }


def verify(conn: sqlite3.Connection) -> list[str]:
    lines: list[str] = []
    # sales
    groups: dict[str, list[dict]] = defaultdict(list)
    with open(SALES_CSV, encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            groups[str(row.get("order_number") or "").strip()].append(row)
    keys = sorted(
        groups, key=lambda k: (_parse_dt(groups[k][0].get("created_date", "")), k)
    )
    sales = conn.execute(
        "SELECT id,total_amount,amount_paid FROM sales "
        "WHERE COALESCE(deleted,0)=0 ORDER BY id"
    ).fetchall()
    mism = 0
    for i, (sid, tot, paid) in enumerate(sales):
        head = groups[keys[i]][0]
        if abs(_f(tot) - _f(head.get("total"))) > 0.01 or abs(
            _f(paid) - _f(head.get("paid_amount"))
        ) > 0.01:
            mism += 1
    lines.append(f"sales amount mismatches after align: {mism}/{len(sales)}")

    by_bill: dict[str, dict] = {}
    with open(PURCH_CSV, encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            bn = (row.get("bill_no") or "").strip()
            if bn and bn not in by_bill:
                by_bill[bn] = row
    pm = 0
    matched = 0
    for pid, bn, fa, ape in conn.execute(
        "SELECT id, bill_number, COALESCE(final_amount,total_amount), "
        "COALESCE(amount_paid_at_entry,0) FROM purchases "
        "WHERE COALESCE(deleted,0)=0"
    ):
        src = by_bill.get(str(bn or "").strip())
        if not src:
            continue
        matched += 1
        if abs(_f(fa) - _f(src.get("total") or src.get("amount"))) > 0.01 or abs(
            _f(ape) - _f(src.get("paid_amount"))
        ) > 0.01:
            pm += 1
    lines.append(f"purchase amount mismatches after align: {pm}/{matched}")

    sr_sum = _f(
        conn.execute(
            "SELECT COALESCE(SUM(refund_amount),0) FROM sales_returns "
            "WHERE COALESCE(deleted,0)=0"
        ).fetchone()[0]
    )
    pr_sum = _f(
        conn.execute(
            "SELECT COALESCE(SUM(refund_amount),0) FROM purchase_returns "
            "WHERE COALESCE(deleted,0)=0"
        ).fetchone()[0]
    )
    lines.append(f"sales_returns sum={sr_sum} count={conn.execute('SELECT COUNT(*) FROM sales_returns WHERE COALESCE(deleted,0)=0').fetchone()[0]}")
    lines.append(f"purchase_returns sum={pr_sum} count={conn.execute('SELECT COUNT(*) FROM purchase_returns WHERE COALESCE(deleted,0)=0').fetchone()[0]}")
    return lines


def recalculate_all(conn: sqlite3.Connection) -> dict:
    cust_ids = [
        r[0]
        for r in conn.execute(
            "SELECT id FROM customers WHERE COALESCE(deleted,0)=0"
        )
    ]
    supp_ids = [
        r[0]
        for r in conn.execute(
            "SELECT id FROM suppliers WHERE COALESCE(deleted,0)=0"
        )
    ]
    with redirect_stdout(io.StringIO()):
        for cid in cust_ids:
            recalculate_customer_due(conn, cid, commit=False, sync=False)
        for sid in supp_ids:
            recalculate_supplier_due(conn, sid, commit=False)
    return {"customers": len(cust_ids), "suppliers": len(supp_ids)}


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    if not os.path.isfile(DB_PATH):
        raise SystemExit(f"DB not found: {DB_PATH}")

    backup = ""
    if not args.dry_run:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup = DB_PATH + f".pre_align_amounts_{stamp}"
        shutil.copy2(DB_PATH, backup)
        _log(f"Backup: {backup}")

    conn = sqlite3.connect(DB_PATH)
    report: list[str] = []
    try:
        _log("Aligning sales amounts to CSV…")
        s = align_sales(conn)
        report.append(f"SALES: {s}")
        _log(f"  sales updated={s['updated']}/{s['total']}")

        _log("Aligning purchase amounts to CSV…")
        p = align_purchases(conn)
        report.append(f"PURCHASES: {p}")
        _log(f"  purchases updated={p['updated']} unmatched={p['unmatched']}")

        _log("Aligning sales returns…")
        sr = align_sales_returns(conn)
        report.append(f"SALES_RETURNS: {sr}")
        _log(f"  sales returns updated={sr['updated']} created={sr['created']} file_total={sr['file_total']}")

        _log("Aligning purchase returns…")
        pr = align_purchase_returns(conn)
        report.append(f"PURCHASE_RETURNS: {pr}")
        _log(f"  purchase returns updated={pr['updated']} file_total={pr['file_total']}")

        _log("Recalculating dues…")
        dues = recalculate_all(conn)
        report.append(f"DUES: {dues}")

        vlines = verify(conn)
        report.extend(vlines)
        for line in vlines:
            _log("  " + line)

        cd = conn.execute(
            "SELECT COUNT(*), ROUND(SUM(total_due),2), ROUND(SUM(total_credit),2) "
            "FROM customers WHERE COALESCE(total_due,0)>0.01 OR COALESCE(total_credit,0)>0.01"
        ).fetchone()
        sd = conn.execute(
            "SELECT COUNT(*), ROUND(SUM(total_due),2), ROUND(SUM(total_credit),2) "
            "FROM suppliers WHERE COALESCE(total_due,0)>0.01 OR COALESCE(total_credit,0)>0.01"
        ).fetchone()
        report.append(f"customers due/credit: {cd}")
        report.append(f"suppliers due/credit: {sd}")
        _log(f"  customers due/credit {cd}")
        _log(f"  suppliers due/credit {sd}")

        with open(REPORT, "w", encoding="utf-8") as f:
            f.write("\n".join(report) + "\n")
        _log(f"Report: {REPORT}")

        if args.dry_run:
            conn.rollback()
            _log("DRY RUN — rolled back")
        else:
            conn.commit()
            _log("Done — amounts now match export files.")
    finally:
        conn.close()


if __name__ == "__main__":
    main()
