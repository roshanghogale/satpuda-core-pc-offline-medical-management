"""Audit Shivkrupa paid amounts / returns vs eVital exports."""
from __future__ import annotations

import csv
import sqlite3
from collections import defaultdict
from pathlib import Path

import openpyxl

DB = Path(
    r"D:\Satpuda Core Server Update\mac2\config\stores"
    r"\Store_Shivkrupa_Medical_General_Store\veterinary.db"
)
SALES_CSV = Path(r"c:\Users\win10\Downloads\vm9xzvqmyn.csv")
PURCH_CSV = Path(r"c:\Users\win10\Downloads\xpjfe4h0im.csv")
PR_XLSX = Path(r"c:\Users\win10\Downloads\Purchase_Return_Register.xlsx")
SR_XLSX = Path(r"c:\Users\win10\Downloads\Sales_Return_Register.xlsx")


def main() -> None:
    conn = sqlite3.connect(str(DB))
    c = conn.cursor()

    print(
        "sales",
        c.execute("SELECT COUNT(*) FROM sales").fetchone()[0],
        "active",
        c.execute(
            "SELECT COUNT(*) FROM sales WHERE COALESCE(deleted,0)=0 "
            "AND COALESCE(is_autosave,0)=0"
        ).fetchone()[0],
        "deleted",
        c.execute(
            "SELECT COUNT(*) FROM sales WHERE COALESCE(deleted,0)!=0"
        ).fetchone()[0],
        "autosave",
        c.execute(
            "SELECT COUNT(*) FROM sales WHERE COALESCE(is_autosave,0)!=0"
        ).fetchone()[0],
    )
    print(
        "purch",
        c.execute("SELECT COUNT(*) FROM purchases").fetchone()[0],
        "deleted",
        c.execute(
            "SELECT COUNT(*) FROM purchases WHERE COALESCE(deleted,0)!=0"
        ).fetchone()[0],
        "autosave",
        c.execute(
            "SELECT COUNT(*) FROM purchases WHERE COALESCE(is_autosave,0)!=0"
        ).fetchone()[0],
    )

    # Sales CSV: first row per bill_no
    by_bill: dict[str, dict] = {}
    with open(SALES_CSV, encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            bn = (row.get("bill_no") or "").strip()
            if not bn or bn in by_bill:
                continue
            by_bill[bn] = {
                "paid": float(row.get("paid_amount") or 0),
                "total": float(row.get("total") or 0),
                "method": (row.get("payment_method") or "").strip(),
                "patient": (row.get("patient_name") or "").strip(),
            }
    print("csv unique sales bills", len(by_bill))

    mism = []
    unmatched = 0
    for _sid, bill_no, total, paid, due, cust in c.execute(
        """
        SELECT s.id, s.bill_no, COALESCE(s.total_amount,0),
               COALESCE(s.amount_paid,0), COALESCE(s.due_amount,0),
               COALESCE(c.name,'')
        FROM sales s LEFT JOIN customers c ON c.id=s.customer_id
        WHERE COALESCE(s.deleted,0)=0 AND COALESCE(s.is_autosave,0)=0
        """
    ):
        key = str(bill_no or "").strip()
        # DB bill_no may be encoded (FY serial); CSV bill_no is plain number
        src = by_bill.get(key)
        if not src:
            digits = "".join(ch for ch in key if ch.isdigit())
            src = by_bill.get(digits)
        if not src:
            unmatched += 1
            continue
        db_paid = float(paid)
        csv_paid = float(src["paid"])
        db_total = float(total)
        csv_total = float(src["total"])
        if abs(db_paid - csv_paid) > 0.5 or abs(db_total - csv_total) > 0.5:
            mism.append(
                (key, db_total, csv_total, db_paid, csv_paid, float(due), cust[:24])
            )
    print("sales unmatched", unmatched, "mismatches", len(mism))
    for m in mism[:30]:
        print("  SALE", m)

    by_pb: dict[str, dict] = {}
    with open(PURCH_CSV, encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            bn = (row.get("bill_no") or "").strip()
            if not bn or bn in by_pb:
                continue
            by_pb[bn] = {
                "paid": float(row.get("paid_amount") or 0),
                "total": float(row.get("total") or row.get("amount") or 0),
                "method": (row.get("payment_method") or "").strip(),
                "dist": (row.get("distributor_name") or "").strip(),
            }
    print("csv unique purch bills", len(by_pb))

    cols = [r[1] for r in c.execute("PRAGMA table_info(purchases)")]
    inv_col = "invoice_no" if "invoice_no" in cols else "supplier_invoice"
    if inv_col not in cols:
        inv_col = "purchase_no"

    pm = []
    matched = 0
    for row in c.execute(
        f"""
        SELECT p.id, COALESCE(p.{inv_col},''), COALESCE(p.purchase_no,''),
               COALESCE(p.final_amount, p.total_amount, 0),
               COALESCE(p.amount_paid_at_entry,0), COALESCE(p.amount_paid,0),
               COALESCE(p.due_amount,0), COALESCE(s.name,'')
        FROM purchases p LEFT JOIN suppliers s ON s.id=p.supplier_id
        WHERE COALESCE(p.deleted,0)=0 AND COALESCE(p.is_autosave,0)=0
        """
    ):
        pid, inv, pno, final_amt, entry_paid, amt_paid, due, sup = row
        key = str(inv or "").strip() or str(pno or "").strip()
        src = by_pb.get(key)
        if not src:
            digits = "".join(ch for ch in key if ch.isdigit())
            src = by_pb.get(digits)
        if not src:
            continue
        matched += 1
        if abs(float(entry_paid) - src["paid"]) > 0.5 or abs(
            float(final_amt) - src["total"]
        ) > 0.5:
            pm.append(
                (
                    key,
                    float(final_amt),
                    src["total"],
                    float(entry_paid),
                    src["paid"],
                    float(amt_paid),
                    float(due),
                    (sup or "")[:24],
                )
            )
    print("purch matched", matched, "mismatches", len(pm))
    for m in pm[:30]:
        print("  PURCH", m)

    # Returns registers
    wb = openpyxl.load_workbook(str(PR_XLSX), read_only=True, data_only=True)
    ws = wb.active
    headers = None
    pr: dict[str, float] = defaultdict(float)
    pr_status: dict[str, str] = {}
    for row in ws.iter_rows(values_only=True):
        if row and row[0] == "Return ID / Voucher No.":
            headers = row
            continue
        if not headers or not row or not row[0]:
            continue
        status = str(row[4] or "").strip()
        pr_status[str(row[0])] = status
        if status == "Draft":
            continue
        pr[str(row[0])] += float(row[17] or 0)
    wb.close()
    print("settled purch returns register", len(pr), round(sum(pr.values()), 2))
    for k, v in pr.items():
        print(" ", k, v, pr_status.get(k))

    wb = openpyxl.load_workbook(str(SR_XLSX), read_only=True, data_only=True)
    ws = wb.active
    headers = None
    sr: dict[str, float] = defaultdict(float)
    for row in ws.iter_rows(values_only=True):
        if row and row[0] == "Return Bill No.":
            headers = row
            continue
        if not headers or not row or row[0] is None:
            continue
        sr[str(row[0])] += float(row[14] or 0)
    wb.close()
    print("sales returns register", len(sr), round(sum(sr.values()), 2))
    for k, v in sorted(sr.items(), key=lambda x: int(x[0]) if x[0].isdigit() else 0):
        print(" ", k, round(v, 2))

    print(
        "db sales return sum",
        c.execute("SELECT SUM(refund_amount) FROM sales_returns").fetchone()[0],
        "count",
        c.execute("SELECT COUNT(*) FROM sales_returns").fetchone()[0],
    )
    print(
        "db purch return sum",
        c.execute("SELECT SUM(refund_amount) FROM purchase_returns").fetchone()[0],
        "count",
        c.execute("SELECT COUNT(*) FROM purchase_returns").fetchone()[0],
    )

    # Customers with credit (possible over-pay / return double count)
    cred = c.execute(
        "SELECT COUNT(*), ROUND(SUM(total_credit),2) FROM customers "
        "WHERE COALESCE(total_credit,0)>0.01"
    ).fetchone()
    duec = c.execute(
        "SELECT COUNT(*), ROUND(SUM(total_due),2) FROM customers "
        "WHERE COALESCE(total_due,0)>0.01"
    ).fetchone()
    print("customers with credit", cred, "with due", duec)
    dues = c.execute(
        "SELECT COUNT(*), ROUND(SUM(total_due),2) FROM suppliers "
        "WHERE COALESCE(total_due,0)>0.01"
    ).fetchone()
    print("suppliers with due", dues)

    # Sample: sales where paid > total (overpaid on bill)
    over = c.execute(
        """
        SELECT COUNT(*), ROUND(SUM(amount_paid-total_amount),2)
        FROM sales
        WHERE COALESCE(deleted,0)=0 AND COALESCE(is_autosave,0)=0
          AND COALESCE(amount_paid,0) > COALESCE(total_amount,0)+0.5
        """
    ).fetchone()
    under = c.execute(
        """
        SELECT COUNT(*), ROUND(SUM(total_amount-amount_paid),2)
        FROM sales
        WHERE COALESCE(deleted,0)=0 AND COALESCE(is_autosave,0)=0
          AND COALESCE(due_amount,0) > 0.5
        """
    ).fetchone()
    print("sales overpaid bills", over, "with due_amount>0.5", under)

    under_p = c.execute(
        """
        SELECT COUNT(*), ROUND(SUM(COALESCE(final_amount,total_amount)-COALESCE(amount_paid_at_entry,0)),2)
        FROM purchases
        WHERE COALESCE(deleted,0)=0 AND COALESCE(is_autosave,0)=0
          AND COALESCE(due_amount,0) > 0.5
        """
    ).fetchone()
    print("purch with due_amount>0.5", under_p)

    conn.close()


if __name__ == "__main__":
    main()
