"""Named export report builders for the Tauri desktop (mirrors Tk export menus)."""
from __future__ import annotations

from datetime import datetime, timedelta
from typing import Any, Optional


def _safe_float(v: Any, default: float = 0.0) -> float:
    try:
        return float(v if v is not None else default)
    except Exception:
        return default


def _is_online() -> bool:
    try:
        from core.sync_prefs import is_online_mode

        return bool(is_online_mode())
    except Exception:
        return False


def _flatten_export_cell(v: Any) -> Any:
    if isinstance(v, dict):
        return (
            v.get("text")
            or v.get("label")
            or v.get("status")
            or v.get("value")
            or ""
        )
    return v


def _flatten_export_row(row: Any) -> list[Any]:
    return [_flatten_export_cell(c) for c in (list(row) if row is not None else [])]


def _resolve_export_dates(fd: str, td: str) -> tuple[str, str]:
    fd = (fd or "").strip()
    td = (td or "").strip()
    if fd or td:
        return fd, td
    try:
        from core.history_prefs import resolve_history_dates

        a, b, _ = resolve_history_dates("", "")
        return a, b
    except Exception:
        return fd, td


def _expiry_sort_key(raw: Any) -> str:
    """Normalize expiry to YYYY-MM-DD for comparisons (ISO or MM/YY)."""
    s = str(raw or "").strip()
    if not s:
        return ""
    if len(s) >= 10 and s[4] == "-" and s[7] == "-":
        return s[:10]
    parts = s.replace(".", "/").replace("-", "/").split("/")
    if len(parts) == 2:
        mm, yy = parts[0].strip(), parts[1].strip()
        if mm.isdigit() and yy.isdigit():
            year = int(yy)
            if year < 100:
                year += 2000 if year < 80 else 1900
            month = max(1, min(12, int(mm)))
            return f"{year:04d}-{month:02d}-01"
    if len(parts) == 3 and all(p.isdigit() for p in parts):
        d, m, y = parts
        year = int(y)
        if year < 100:
            year += 2000
        return f"{year:04d}-{int(m):02d}-{int(d):02d}"
    return s


def _schedule_matches(value: Any, mode: str, picked: list[str]) -> bool:
    sch = str(value or "").strip()
    mode = (mode or "all").strip().lower()
    if mode in ("", "all"):
        return True
    if mode in ("non_scheduled", "nonscheduled", "non-scheduled"):
        return not sch
    if not picked:
        return True
    want = {p.strip().upper() for p in picked if str(p).strip()}
    return sch.upper() in want


def _online_sale_docs(fd: str, td: str) -> list[dict[str, Any]]:
    from core import store_query_client as sq
    from core.online_mutation_queue import merge_server_rows, overlay_sales_dicts
    from core.server_crud import get_doc

    raw = list(
        (sq.list_sales(from_date=fd or "", to_date=td or "", limit=5000) or {}).get(
            "rows"
        )
        or []
    )
    merged = merge_server_rows(raw, overlay_sales_dicts(), collection="sales")
    out: list[dict[str, Any]] = []
    for r in merged:
        if not isinstance(r, dict) or r.get("deleted") or int(r.get("is_autosave") or 0):
            continue
        items = r.get("items") or r.get("medicines")
        if isinstance(items, list) and items:
            doc = dict(r)
            doc["items"] = items
            out.append(doc)
            continue
        try:
            sid = int(r.get("id") or r.get("local_id") or 0)
        except (TypeError, ValueError):
            sid = 0
        if sid:
            try:
                doc = get_doc("sales", sid) or dict(r)
            except Exception:
                doc = dict(r)
            if not isinstance(doc, dict):
                doc = dict(r)
            items = doc.get("items") or doc.get("medicines")
            if not isinstance(items, list) or not items:
                try:
                    extra = sq.get_sale(sid) or {}
                    if isinstance(extra, dict):
                        items = extra.get("items") or extra.get("medicines") or extra.get("lines")
                        if items:
                            doc = dict(doc)
                            doc["items"] = items
                            if extra.get("customer_name") and not doc.get("customer_name"):
                                doc["customer_name"] = extra.get("customer_name")
                except Exception:
                    pass
            out.append(doc)
        else:
            out.append(dict(r))
    return out


def _sale_item_rows(doc: dict[str, Any]) -> list[dict[str, Any]]:
    items = doc.get("items") or doc.get("medicines") or []
    if not isinstance(items, list):
        return []
    out: list[dict[str, Any]] = []
    for it in items:
        if not isinstance(it, dict):
            continue
        mid = 0
        try:
            mid = int(it.get("medicine_id") or it.get("id") or 0)
        except (TypeError, ValueError):
            mid = 0
        mp: dict[str, Any] = {}
        if mid:
            try:
                from core.online_catalog import medicine_by_id

                mp = medicine_by_id(mid) or {}
            except Exception:
                mp = {}
        out.append(
            {
                "name": it.get("medicine_name")
                or it.get("name")
                or mp.get("name")
                or "",
                "batch": it.get("batch_no")
                or it.get("batch")
                or mp.get("batch_no")
                or "",
                "content": it.get("content_drug")
                or it.get("content")
                or mp.get("content_drug")
                or "",
                "schedule": it.get("schedule") or mp.get("schedule") or "",
                "expiry": it.get("expiry_date")
                or it.get("expiry")
                or mp.get("expiry_date")
                or "",
                "qty": _safe_float(it.get("qty") or it.get("quantity")),
                "rate": _safe_float(it.get("rate")),
                "amount": _safe_float(
                    it.get("amount")
                    if it.get("amount") is not None
                    else it.get("item_amount")
                ),
            }
        )
    return out


def _online_inventory_meds() -> list[dict[str, Any]]:
    meds: list[dict[str, Any]] = []
    try:
        from core.online_catalog import medicines as oc_medicines

        meds = [
            m
            for m in oc_medicines()
            if isinstance(m, dict) and not m.get("is_hidden") and not m.get("deleted")
        ]
    except Exception:
        meds = []
    if meds:
        return meds
    try:
        from core import store_query_client as sq

        data = sq.list_inventory(limit=5000) or {}
        rows = data.get("rows") or data.get("medicines") or data.get("items") or []
        return [m for m in rows if isinstance(m, dict) and not m.get("is_hidden")]
    except Exception:
        return []


def _filter_export_cols(
    page_key: str, report_key: str, columns: list[str], rows: list[list[Any]]
) -> tuple[list[str], list[list[Any]]]:
    try:
        from core.column_config import get_export_column_visibility

        vis = get_export_column_visibility(page_key, report_key)
        idxs = [i for i, c in enumerate(columns) if vis.get(c, True)]
        if not idxs:
            idxs = list(range(len(columns)))
        return (
            [columns[i] for i in idxs],
            [[r[i] if i < len(r) else "" for i in idxs] for r in rows],
        )
    except Exception:
        return columns, rows


def _strip_fy_from_bill_columns(
    columns: list[str], rows: list[list[Any]]
) -> list[list[Any]]:
    """Show short bill numbers (SCB1) — strip /FY… suffix on export."""
    idxs = [i for i, c in enumerate(columns) if c in ("Bill No", "Bill No.")]
    if not idxs:
        return rows
    try:
        from core.fy_serial import display_sales_bill_no
    except Exception:
        def display_sales_bill_no(v):  # type: ignore
            s = str(v or "")
            return s.split("/FY", 1)[0] if "/FY" in s else s

    out: list[list[Any]] = []
    for row in rows:
        nr = list(row)
        for i in idxs:
            if i < len(nr):
                nr[i] = display_sales_bill_no(str(nr[i] or ""))
        out.append(nr)
    return out


def export_options(page: str) -> list[dict[str, str]]:
    """Report picker options matching Tk export menus."""
    if page == "inventory":
        return [
            {"key": "current_view", "label": "Current View (filtered)"},
            {"key": "stock_statement", "label": "Stock Statement (all)"},
            {"key": "near_expiry", "label": "Near Expiry Report"},
            {"key": "expired_stock", "label": "Expired Stock Report"},
            {"key": "schedule_wise_stock", "label": "Schedule-wise Stock"},
        ]
    if page == "sales_history":
        return [
            {"key": "current_view", "label": "Current View (with filters)"},
            {"key": "sales_register", "label": "Sales Register (all bills)"},
            {"key": "monthly_summary", "label": "Monthly Summary"},
            {"key": "daily_summary", "label": "Daily Sales Summary"},
            {"key": "customer_due", "label": "Customer Due Report"},
            {"key": "doctor_wise", "label": "Doctor-wise Sales"},
            {"key": "payment_mode", "label": "Payment Mode Report"},
            {"key": "schedule_report", "label": "Schedule Report"},
        ]
    if page == "purchase_history":
        return [
            {"key": "current_view", "label": "Current View"},
            {"key": "purchase_register", "label": "Purchase Register"},
            {"key": "monthly_summary", "label": "Monthly Summary"},
            {"key": "supplier_due", "label": "Supplier Due Report"},
            {"key": "gst_purchase", "label": "GST Purchase Report"},
        ]
    return [{"key": "current_view", "label": "Current View"}]


def run_export(
    conn,
    page: str,
    report: str,
    *,
    from_date: str = "",
    to_date: str = "",
    current_columns: Optional[list[str]] = None,
    current_rows: Optional[list[list[Any]]] = None,
    schedule: Optional[dict[str, Any]] = None,
) -> dict[str, Any]:
    page = (page or "").strip()
    report = (report or "").strip() or "current_view"
    fd, td = _resolve_export_dates(from_date, to_date)
    if page == "inventory":
        fd, td = (from_date or "").strip(), (to_date or "").strip()

    if report == "current_view":
        cols = list(current_columns or [])
        rows = [_flatten_export_row(r) for r in (current_rows or [])]
        if cols and any(c.lower() == "status" for c in cols):
            # Status cells from the live table are objects; flatten already handled.
            pass
        return {
            "columns": cols,
            "rows": rows,
            "filename": f"{page}_current_view",
            "title": "Current View",
        }

    if page == "inventory":
        return _inventory_export(conn, report)
    if page == "sales_history":
        return _sales_export(conn, report, fd, td, schedule=schedule)
    if page == "purchase_history":
        return _purchase_export(conn, report, fd, td)
    return {"error": f"Unknown page {page}", "columns": [], "rows": []}


def list_schedules_for_export() -> dict[str, Any]:
    """Schedules + saved layout for the Schedule Report dialog."""
    try:
        from core.layout_config import get_configured_schedules
        from core.export_prefs import (
            SCHEDULE_LAYOUT_LANDSCAPE,
            SCHEDULE_LAYOUT_PORTRAIT,
            SCHEDULE_LAYOUT_STYLED,
            SCHEDULE_LAYOUT_LABELS,
            SCHEDULE_DM_STYLE_CLASSIC,
            SCHEDULE_DM_STYLE_SIGN,
            SCHEDULE_DM_STYLE_LABELS,
            load_schedule_report_layout,
            load_schedule_dm_prefs,
        )

        layout = load_schedule_report_layout()
        dm = load_schedule_dm_prefs()
        valid = (
            SCHEDULE_LAYOUT_PORTRAIT,
            SCHEDULE_LAYOUT_LANDSCAPE,
            SCHEDULE_LAYOUT_STYLED,
        )
        return {
            "schedules": list(get_configured_schedules()),
            "layout": layout if layout in valid else SCHEDULE_LAYOUT_PORTRAIT,
            "layouts": [
                {"key": k, "label": SCHEDULE_LAYOUT_LABELS[k]} for k in valid
            ],
            "dm_style": dm.get("style") or SCHEDULE_DM_STYLE_CLASSIC,
            "dm_borders": bool(dm.get("borders", True)),
            "dm_styles": [
                {
                    "key": SCHEDULE_DM_STYLE_CLASSIC,
                    "label": SCHEDULE_DM_STYLE_LABELS[SCHEDULE_DM_STYLE_CLASSIC],
                },
                {
                    "key": SCHEDULE_DM_STYLE_SIGN,
                    "label": SCHEDULE_DM_STYLE_LABELS[SCHEDULE_DM_STYLE_SIGN],
                },
            ],
        }
    except Exception as exc:
        return {
            "schedules": [],
            "layout": "portrait",
            "layouts": [],
            "dm_style": "classic",
            "dm_borders": True,
            "dm_styles": [],
            "error": str(exc),
        }


def _inventory_export(conn, report: str) -> dict[str, Any]:
    today = datetime.now().strftime("%Y-%m-%d")
    if _is_online():
        try:
            meds = _online_inventory_meds()
            if report == "stock_statement":
                cols = [
                    "Name", "Type", "Batch", "Expiry", "Stock", "Unit",
                    "MRP", "Rate", "Manufacturer", "Schedule",
                ]
                rows = [[
                    m.get("name") or "",
                    m.get("type") or "",
                    m.get("batch_no") or "",
                    m.get("expiry_date") or "",
                    _safe_float(m.get("stock_qty")),
                    m.get("unit") or "",
                    _safe_float(m.get("mrp")),
                    _safe_float(m.get("rate")),
                    m.get("manufacturer") or "",
                    m.get("schedule") or "",
                ] for m in meds]
                cols, rows = _filter_export_cols("inventory", "stock_statement", cols, rows)
                return {"columns": cols, "rows": rows, "filename": "stock_statement", "title": "Stock Statement"}
            if report == "near_expiry":
                threshold = (datetime.now() + timedelta(days=90)).strftime("%Y-%m-%d")
                cols = ["Name", "Type", "Batch", "Expiry", "Stock", "Manufacturer"]
                rows = []
                for m in meds:
                    exp = _expiry_sort_key(m.get("expiry_date") or m.get("expiry"))
                    if exp and today < exp <= threshold:
                        rows.append([
                            m.get("name") or "", m.get("type") or "", m.get("batch_no") or "",
                            m.get("expiry_date") or m.get("expiry") or exp,
                            _safe_float(m.get("stock_qty")), m.get("manufacturer") or "",
                        ])
                cols, rows = _filter_export_cols("inventory", "near_expiry", cols, rows)
                return {"columns": cols, "rows": rows, "filename": "near_expiry", "title": "Near Expiry Report"}
            if report == "expired_stock":
                cols = ["Name", "Type", "Batch", "Expiry", "Stock", "Manufacturer"]
                rows = []
                for m in meds:
                    exp = _expiry_sort_key(m.get("expiry_date") or m.get("expiry"))
                    if exp and exp <= today:
                        rows.append([
                            m.get("name") or "", m.get("type") or "", m.get("batch_no") or "",
                            m.get("expiry_date") or m.get("expiry") or exp,
                            _safe_float(m.get("stock_qty")), m.get("manufacturer") or "",
                        ])
                cols, rows = _filter_export_cols("inventory", "expired_stock", cols, rows)
                return {"columns": cols, "rows": rows, "filename": "expired_stock", "title": "Expired Stock Report"}
            if report == "schedule_wise_stock":
                cols = ["Schedule", "Name", "Batch", "Expiry", "Stock", "MRP"]
                rows = [[
                    m.get("schedule") or "Non-Scheduled",
                    m.get("name") or "",
                    m.get("batch_no") or "",
                    m.get("expiry_date") or "",
                    _safe_float(m.get("stock_qty")),
                    _safe_float(m.get("mrp")),
                ] for m in meds]
                cols, rows = _filter_export_cols(
                    "inventory", "schedule_wise_stock", cols, rows
                )
                return {
                    "columns": cols, "rows": rows,
                    "filename": "schedule_wise_stock", "title": "Schedule-wise Stock",
                }
        except Exception as exc:
            return {"error": str(exc), "columns": [], "rows": []}
    if report == "stock_statement":
        cols = [
            "Name",
            "Type",
            "Batch",
            "Expiry",
            "Stock",
            "Unit",
            "MRP",
            "Rate",
            "Manufacturer",
            "Schedule",
        ]
        raw = conn.execute(
            """
            SELECT name, COALESCE(type,''), COALESCE(batch_no,''), COALESCE(expiry_date,''),
                   COALESCE(stock_qty,0), COALESCE(unit,''), COALESCE(mrp,0), COALESCE(rate,0),
                   COALESCE(manufacturer,''), COALESCE(schedule,'')
            FROM medicines
            WHERE COALESCE(deleted,0)=0 AND COALESCE(is_hidden,0)=0
            ORDER BY name COLLATE NOCASE
            """
        ).fetchall()
        rows = [list(r) for r in raw]
        cols, rows = _filter_export_cols("inventory", "stock_statement", cols, rows)
        return {
            "columns": cols,
            "rows": rows,
            "filename": "stock_statement",
            "title": "Stock Statement",
        }

    if report == "near_expiry":
        threshold = (datetime.now() + timedelta(days=90)).strftime("%Y-%m-%d")
        cols = ["Name", "Type", "Batch", "Expiry", "Stock", "Manufacturer"]
        raw = conn.execute(
            """
            SELECT name, COALESCE(type,''), COALESCE(batch_no,''), COALESCE(expiry_date,''),
                   COALESCE(stock_qty,0), COALESCE(manufacturer,'')
            FROM medicines
            WHERE COALESCE(deleted,0)=0 AND COALESCE(is_hidden,0)=0
              AND expiry_date>=? AND expiry_date<=? AND COALESCE(stock_qty,0)>0
            ORDER BY expiry_date
            """,
            (today, threshold),
        ).fetchall()
        rows = [list(r) for r in raw]
        cols, rows = _filter_export_cols("inventory", "near_expiry", cols, rows)
        return {
            "columns": cols,
            "rows": rows,
            "filename": "near_expiry",
            "title": "Near Expiry",
        }

    if report == "expired_stock":
        cols = ["Name", "Type", "Batch", "Expiry", "Stock", "Manufacturer"]
        raw = conn.execute(
            """
            SELECT name, COALESCE(type,''), COALESCE(batch_no,''), COALESCE(expiry_date,''),
                   COALESCE(stock_qty,0), COALESCE(manufacturer,'')
            FROM medicines
            WHERE COALESCE(deleted,0)=0 AND COALESCE(is_hidden,0)=0
              AND expiry_date<? AND COALESCE(stock_qty,0)>0
            ORDER BY expiry_date
            """,
            (today,),
        ).fetchall()
        rows = [list(r) for r in raw]
        cols, rows = _filter_export_cols("inventory", "expired_stock", cols, rows)
        return {
            "columns": cols,
            "rows": rows,
            "filename": "expired_stock",
            "title": "Expired Stock",
        }

    if report == "schedule_wise_stock":
        cols = ["Schedule", "Name", "Batch", "Expiry", "Stock", "MRP"]
        raw = conn.execute(
            """
            SELECT COALESCE(NULLIF(TRIM(schedule),''),'—'), name, COALESCE(batch_no,''),
                   COALESCE(expiry_date,''), COALESCE(stock_qty,0), COALESCE(mrp,0)
            FROM medicines
            WHERE COALESCE(deleted,0)=0 AND COALESCE(is_hidden,0)=0
            ORDER BY schedule, name
            """
        ).fetchall()
        rows = [list(r) for r in raw]
        cols, rows = _filter_export_cols(
            "inventory", "schedule_wise_stock", cols, rows
        )
        return {
            "columns": cols,
            "rows": rows,
            "filename": "schedule_wise_stock",
            "title": "Schedule-wise Stock",
        }

    return {"error": f"Unknown inventory report {report}", "columns": [], "rows": []}


def _sales_export(
    conn,
    report: str,
    fd: str,
    td: str,
    *,
    schedule: Optional[dict[str, Any]] = None,
) -> dict[str, Any]:
    if _is_online() and report != "schedule_report":
        try:
            from core.fy_serial import display_sales_bill_no
            from core.online_mutation_queue import merge_server_rows, overlay_sales_dicts
            from core import store_query_client as sq

            listed = list(
                (sq.list_sales(from_date=fd or "", to_date=td or "", limit=5000) or {}).get(
                    "rows"
                )
                or []
            )
            raw = merge_server_rows(listed, overlay_sales_dicts(), collection="sales")
            raw = [
                r
                for r in raw
                if isinstance(r, dict)
                and not r.get("deleted")
                and not int(r.get("is_autosave") or 0)
            ]
            if report == "sales_register":
                cols = [
                    "Bill No", "Date", "Customer", "Phone", "Total Amount", "Discount",
                    "Cash Paid", "Online Paid", "Amount Paid", "Previous Due",
                    "Due Amount", "Total Due", "Doctor",
                ]
                rows = [[
                    r.get("bill_no") or "",
                    r.get("bill_date") or "",
                    r.get("customer_name") or "",
                    r.get("customer_phone") or r.get("phone") or "",
                    _safe_float(r.get("total_amount")),
                    _safe_float(r.get("discount")),
                    _safe_float(r.get("cash_paid")),
                    _safe_float(r.get("online_paid")),
                    _safe_float(r.get("amount_paid")),
                    _safe_float(r.get("previous_due")),
                    _safe_float(r.get("due_amount")),
                    _safe_float(r.get("total_due")),
                    r.get("doctor_name") or "",
                ] for r in raw]
                cols, rows = _filter_export_cols("sales_history", "sales_register", cols, rows)
                rows = _strip_fy_from_bill_columns(cols, rows)
                return {"columns": cols, "rows": rows, "filename": "sales_register", "title": "Sales Register"}
            if report == "monthly_summary":
                buckets: dict[str, list] = {}
                for r in raw:
                    ym = str(r.get("bill_date") or "")[:7]
                    if not ym:
                        continue
                    b = buckets.setdefault(ym, [0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0])
                    b[0] += 1
                    b[1] += _safe_float(r.get("total_amount"))
                    b[2] += _safe_float(r.get("discount"))
                    b[3] += _safe_float(r.get("cash_paid"))
                    b[4] += _safe_float(r.get("online_paid"))
                    b[5] += _safe_float(r.get("amount_paid"))
                    b[6] += _safe_float(r.get("due_amount"))
                cols = [
                    "Month", "Bills", "Total Sales", "Discount",
                    "Cash Paid", "Online Paid", "Amount Paid", "Due Amount",
                ]
                rows = [[k, *v] for k, v in sorted(buckets.items(), reverse=True)]
                cols, rows = _filter_export_cols("sales_history", "monthly_summary", cols, rows)
                return {
                    "columns": cols, "rows": rows,
                    "filename": "sales_monthly_summary", "title": "Monthly Summary",
                }
            if report == "daily_summary":
                buckets = {}
                for r in raw:
                    d = str(r.get("bill_date") or "")[:10]
                    if not d:
                        continue
                    b = buckets.setdefault(d, [0, 0.0, 0.0, 0.0, 0.0, 0.0])
                    b[0] += 1
                    b[1] += _safe_float(r.get("total_amount"))
                    b[2] += _safe_float(r.get("cash_paid"))
                    b[3] += _safe_float(r.get("online_paid"))
                    b[4] += _safe_float(r.get("amount_paid"))
                    b[5] += _safe_float(r.get("due_amount"))
                cols = [
                    "Date", "Bills", "Total Amount", "Cash Paid",
                    "Online Paid", "Amount Paid", "Due Amount",
                ]
                rows = [[k, *v] for k, v in sorted(buckets.items(), reverse=True)]
                cols, rows = _filter_export_cols("sales_history", "daily_summary", cols, rows)
                return {
                    "columns": cols, "rows": rows,
                    "filename": "sales_daily_summary", "title": "Daily Summary",
                }
            if report == "customer_due":
                cols = ["Customer", "Phone", "Date", "Bill No", "Total Amount", "Due Amount"]
                rows = []
                fifo = {}
                try:
                    from core.due_fifo import fifo_sales_remaining_by_bill
                    from core.online_mutation_queue import overlay_customer_payment_dicts

                    pays = merge_server_rows(
                        list(
                            (sq.list_customer_payments(limit=5000) or {}).get("rows")
                            or []
                        ),
                        overlay_customer_payment_dicts(),
                        collection="customer_payments",
                    )
                    rets = list(
                        (sq.list_sales_returns(limit=5000) or {}).get("rows") or []
                    )
                    fifo = fifo_sales_remaining_by_bill(raw, pays, rets)
                except Exception:
                    fifo = {}
                for r in raw:
                    try:
                        sid = int(r.get("id") or r.get("local_id") or 0)
                    except (TypeError, ValueError):
                        sid = 0
                    hit = fifo.get(sid) if sid else None
                    if hit is not None:
                        due = _safe_float(hit[0])
                        cleared = bool(hit[1])
                    else:
                        due = _safe_float(r.get("total_due") or r.get("due_amount"))
                        cleared = bool(r.get("account_cleared"))
                    if due <= 0.01 or cleared:
                        continue
                    rows.append([
                        r.get("customer_name") or "",
                        r.get("customer_phone") or r.get("phone") or "",
                        r.get("bill_date") or "",
                        display_sales_bill_no(str(r.get("bill_no") or "")),
                        _safe_float(r.get("total_amount")),
                        due,
                    ])
                cols, rows = _filter_export_cols("sales_history", "customer_due", cols, rows)
                return {
                    "columns": cols, "rows": rows,
                    "filename": "customer_due", "title": "Customer Due Report",
                }
            if report == "payment_mode":
                cols = [
                    "Date", "Bill No", "Customer", "Total Amount",
                    "Cash Paid", "Online Paid", "Amount Paid", "Due Amount",
                ]
                rows = [[
                    r.get("bill_date") or "",
                    r.get("bill_no") or "",
                    r.get("customer_name") or "",
                    _safe_float(r.get("total_amount")),
                    _safe_float(r.get("cash_paid")),
                    _safe_float(r.get("online_paid")),
                    _safe_float(r.get("amount_paid")),
                    _safe_float(r.get("due_amount")),
                ] for r in raw]
                cols, rows = _filter_export_cols("sales_history", "payment_mode", cols, rows)
                rows = _strip_fy_from_bill_columns(cols, rows)
                return {
                    "columns": cols, "rows": rows,
                    "filename": "payment_mode", "title": "Payment Mode Report",
                }
            if report == "doctor_wise":
                cols = [
                    "Doctor", "Date", "Customer", "Bill No",
                    "Medicine", "Schedule", "Qty", "Rate", "Amount",
                ]
                rows = []
                for doc in _online_sale_docs(fd, td):
                    doctor = str(doc.get("doctor_name") or "").strip()
                    if not doctor:
                        continue
                    bno = display_sales_bill_no(str(doc.get("bill_no") or ""))
                    bdate = doc.get("bill_date") or ""
                    cust = doc.get("customer_name") or ""
                    lines = _sale_item_rows(doc)
                    if not lines:
                        rows.append([
                            doctor, bdate, cust, bno, "", "", "", "",
                            _safe_float(doc.get("total_amount")),
                        ])
                        continue
                    for it in lines:
                        rows.append([
                            doctor,
                            bdate,
                            cust,
                            bno,
                            it.get("name") or "",
                            it.get("schedule") or "",
                            it.get("qty") or 0,
                            it.get("rate") or 0,
                            it.get("amount") or 0,
                        ])
                cols, rows = _filter_export_cols("sales_history", "doctor_wise", cols, rows)
                rows = _strip_fy_from_bill_columns(cols, rows)
                return {
                    "columns": cols, "rows": rows,
                    "filename": "doctor_wise_sales", "title": "Doctor-wise Sales",
                }
        except Exception as exc:
            return {"error": str(exc), "columns": [], "rows": []}

    where = ["COALESCE(s.deleted,0)=0", "COALESCE(s.is_autosave,0)=0"]
    params: list[Any] = []
    if fd:
        where.append("s.bill_date>=?")
        params.append(fd)
    if td:
        where.append("s.bill_date<=?")
        params.append(td)
    wsql = " AND ".join(where)

    if report == "sales_register":
        cols = [
            "Bill No",
            "Date",
            "Customer",
            "Phone",
            "Total Amount",
            "Discount",
            "Cash Paid",
            "Online Paid",
            "Amount Paid",
            "Previous Due",
            "Due Amount",
            "Total Due",
            "Doctor",
        ]
        raw = conn.execute(
            f"""
            SELECT s.bill_no, s.bill_date, COALESCE(c.name,''), COALESCE(c.phone,''),
                   COALESCE(s.total_amount,0), COALESCE(s.discount,0),
                   COALESCE(s.cash_paid,0), COALESCE(s.online_paid,0),
                   COALESCE(s.amount_paid,0), COALESCE(s.previous_due,0),
                   COALESCE(s.due_amount,0), COALESCE(s.total_due,0),
                   COALESCE(s.doctor_name,'')
            FROM sales s LEFT JOIN customers c ON s.customer_id=c.id
            WHERE {wsql}
            ORDER BY s.bill_date DESC
            """,
            params,
        ).fetchall()
        rows = [list(r) for r in raw]
        cols, rows = _filter_export_cols("sales_history", "sales_register", cols, rows)
        rows = _strip_fy_from_bill_columns(cols, rows)
        return {
            "columns": cols,
            "rows": rows,
            "filename": "sales_register",
            "title": "Sales Register",
        }

    if report == "monthly_summary":
        cols = [
            "Month",
            "Bills",
            "Total Sales",
            "Discount",
            "Cash Paid",
            "Online Paid",
            "Amount Paid",
            "Due Amount",
        ]
        raw = conn.execute(
            f"""
            SELECT strftime('%Y-%m', s.bill_date), COUNT(*),
                   SUM(COALESCE(s.total_amount,0)), SUM(COALESCE(s.discount,0)),
                   SUM(COALESCE(s.cash_paid,0)), SUM(COALESCE(s.online_paid,0)),
                   SUM(COALESCE(s.amount_paid,0)), SUM(COALESCE(s.due_amount,0))
            FROM sales s WHERE {wsql}
            GROUP BY strftime('%Y-%m', s.bill_date) ORDER BY 1 DESC
            """,
            params,
        ).fetchall()
        rows = [list(r) for r in raw]
        cols, rows = _filter_export_cols(
            "sales_history", "monthly_summary", cols, rows
        )
        return {
            "columns": cols,
            "rows": rows,
            "filename": "sales_monthly_summary",
            "title": "Monthly Summary",
        }

    if report == "daily_summary":
        cols = [
            "Date",
            "Bills",
            "Total Amount",
            "Cash Paid",
            "Online Paid",
            "Amount Paid",
            "Due Amount",
        ]
        raw = conn.execute(
            f"""
            SELECT s.bill_date, COUNT(*),
                   SUM(COALESCE(s.total_amount,0)), SUM(COALESCE(s.cash_paid,0)),
                   SUM(COALESCE(s.online_paid,0)), SUM(COALESCE(s.amount_paid,0)),
                   SUM(COALESCE(s.due_amount,0))
            FROM sales s WHERE {wsql}
            GROUP BY s.bill_date ORDER BY 1 DESC
            """,
            params,
        ).fetchall()
        rows = [list(r) for r in raw]
        cols, rows = _filter_export_cols("sales_history", "daily_summary", cols, rows)
        return {
            "columns": cols,
            "rows": rows,
            "filename": "sales_daily_summary",
            "title": "Daily Sales Summary",
        }

    if report == "customer_due":
        cols = ["Customer", "Phone", "Date", "Bill No", "Total Amount", "Due Amount"]
        raw = conn.execute(
            """
            SELECT COALESCE(c.name,''), COALESCE(c.phone,''), s.bill_date, s.bill_no,
                   COALESCE(s.total_amount,0), COALESCE(s.due_amount,0)
            FROM sales s LEFT JOIN customers c ON s.customer_id=c.id
            WHERE COALESCE(s.deleted,0)=0 AND COALESCE(s.is_autosave,0)=0
              AND COALESCE(s.due_amount,0)>0
            ORDER BY s.bill_date DESC
            """
        ).fetchall()
        rows = [list(r) for r in raw]
        cols, rows = _filter_export_cols("sales_history", "customer_due", cols, rows)
        rows = _strip_fy_from_bill_columns(cols, rows)
        return {
            "columns": cols,
            "rows": rows,
            "filename": "customer_due",
            "title": "Customer Due",
        }

    if report == "doctor_wise":
        cols = [
            "Doctor",
            "Date",
            "Customer",
            "Bill No",
            "Medicine",
            "Schedule",
            "Qty",
            "Rate",
            "Amount",
        ]
        item_where = list(where) + [
            "s.doctor_name IS NOT NULL",
            "TRIM(s.doctor_name)!=''",
        ]
        raw = conn.execute(
            f"""
            SELECT COALESCE(s.doctor_name,''), s.bill_date, COALESCE(c.name,''), s.bill_no,
                   COALESCE(m.name,''), COALESCE(m.schedule,''),
                   COALESCE(si.qty,0), COALESCE(si.rate,0), COALESCE(si.amount,0)
            FROM sales s
            LEFT JOIN customers c ON s.customer_id=c.id
            JOIN sales_items si ON s.id=si.sale_id
            JOIN medicines m ON si.medicine_id=m.id
            WHERE {" AND ".join(item_where)}
            ORDER BY s.doctor_name, s.bill_date
            """,
            params,
        ).fetchall()
        rows = [list(r) for r in raw]
        cols, rows = _filter_export_cols("sales_history", "doctor_wise", cols, rows)
        rows = _strip_fy_from_bill_columns(cols, rows)
        return {
            "columns": cols,
            "rows": rows,
            "filename": "doctor_wise_sales",
            "title": "Doctor-wise Sales",
        }

    if report == "payment_mode":
        cols = [
            "Date",
            "Bill No",
            "Customer",
            "Total Amount",
            "Cash Paid",
            "Online Paid",
            "Amount Paid",
            "Due Amount",
        ]
        raw = conn.execute(
            f"""
            SELECT s.bill_date, s.bill_no, COALESCE(c.name,''),
                   COALESCE(s.total_amount,0), COALESCE(s.cash_paid,0),
                   COALESCE(s.online_paid,0), COALESCE(s.amount_paid,0),
                   COALESCE(s.due_amount,0)
            FROM sales s LEFT JOIN customers c ON s.customer_id=c.id
            WHERE {wsql}
            ORDER BY s.bill_date DESC
            """,
            params,
        ).fetchall()
        rows = [list(r) for r in raw]
        cols, rows = _filter_export_cols("sales_history", "payment_mode", cols, rows)
        rows = _strip_fy_from_bill_columns(cols, rows)
        return {
            "columns": cols,
            "rows": rows,
            "filename": "payment_mode_report",
            "title": "Payment Mode Report",
        }

    if report == "schedule_report":
        return _schedule_report_export(conn, fd, td, schedule or {})

    return {"error": f"Unknown sales report {report}", "columns": [], "rows": []}


def _schedule_sql_filter(choice: dict[str, Any]) -> tuple[str, list[Any]]:
    mode = (choice.get("mode") or "all").strip()
    if mode == "all":
        return "", []
    if mode == "non_scheduled":
        return " AND (m.schedule IS NULL OR TRIM(m.schedule)='')", []
    names = [str(s).strip() for s in (choice.get("schedules") or []) if str(s).strip()]
    if not names:
        return "", []
    placeholders = ",".join("?" * len(names))
    return f" AND m.schedule IN ({placeholders})", names


def _fmt_schedule_expiry(raw: Any) -> str:
    if not raw:
        return ""
    try:
        p = str(raw).split("-")
        return f"{p[2]}/{p[1]}/{p[0][2:]}" if len(p) == 3 else str(raw)
    except Exception:
        return str(raw)


def _schedule_report_export(
    conn, fd: str, td: str, choice: dict[str, Any]
) -> dict[str, Any]:
    """Line-item sales × medicines filtered by schedule (mirrors Tk)."""
    try:
        from core.export_prefs import (
            SCHEDULE_LAYOUT_PORTRAIT,
            save_schedule_report_layout,
        )
    except Exception:
        SCHEDULE_LAYOUT_PORTRAIT = "portrait"

        def save_schedule_report_layout(_layout: str) -> None:
            return None

    sch_fd = str(choice.get("from_date") or fd or "").strip()
    sch_td = str(choice.get("to_date") or td or "").strip()
    mode = str(choice.get("mode") or "all").strip()
    picked = [str(s).strip() for s in (choice.get("schedules") or []) if str(s).strip()]
    if mode == "selected" and not picked:
        return {"error": "Select at least one schedule", "columns": [], "rows": []}
    if mode == "all":
        label = "All Schedules"
    elif mode == "non_scheduled":
        label = "Non-Scheduled"
    else:
        label = ", ".join(picked) if picked else "Selected"
    page_layout = str(choice.get("page_layout") or SCHEDULE_LAYOUT_PORTRAIT).strip()
    try:
        save_schedule_report_layout(page_layout)
    except Exception:
        pass
    try:
        from core.export_prefs import save_schedule_dm_prefs
        if "dm_style" in choice or "dm_borders" in choice:
            save_schedule_dm_prefs(
                style=choice.get("dm_style"),
                borders=choice.get("dm_borders") if "dm_borders" in choice else None,
            )
    except Exception:
        pass

    if _is_online():
        try:
            from core.fy_serial import display_sales_bill_no

            cols = [
                "Date",
                "Bill No",
                "Customer",
                "Doctor",
                "Medicine",
                "Batch",
                "Content/Drug",
                "Schedule",
                "Expiry",
                "Qty",
                "Rate",
                "Amount",
            ]
            rows: list[list[Any]] = []
            for doc in _online_sale_docs(sch_fd, sch_td):
                bdate = str(doc.get("bill_date") or "")[:10]
                bno = display_sales_bill_no(str(doc.get("bill_no") or ""))
                cust = doc.get("customer_name") or ""
                doctor = doc.get("doctor_name") or "—"
                for it in _sale_item_rows(doc):
                    if not _schedule_matches(it.get("schedule"), mode, picked):
                        continue
                    sch = str(it.get("schedule") or "").strip() or "—"
                    rows.append(
                        [
                            bdate,
                            bno,
                            cust,
                            doctor or "—",
                            it.get("name") or "",
                            it.get("batch") or "—",
                            it.get("content") or "—",
                            sch,
                            _fmt_schedule_expiry(it.get("expiry")),
                            int(it.get("qty") or 0),
                            f"{_safe_float(it.get('rate')):.2f}",
                            f"{_safe_float(it.get('amount')):.2f}",
                        ]
                    )
            rows.sort(key=lambda r: (str(r[7] or ""), str(r[0] or ""), str(r[2] or "")))
            pl = str(page_layout or "").strip().lower()
            if rows and not pl.startswith("land"):
                drop = {"Rate", "Amount"}
                keep = [i for i, c in enumerate(cols) if c not in drop]
                cols = [cols[i] for i in keep]
                rows = [[row[i] for i in keep] for row in rows]
            cols, rows = _filter_export_cols(
                "sales_history", "schedule_report", cols, rows
            )
            rows = _strip_fy_from_bill_columns(cols, rows)
            safe = "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in label)[:40]
            dm_style = str(choice.get("dm_style") or "classic")
            dm_borders = bool(choice.get("dm_borders", True))
            return {
                "columns": cols,
                "rows": rows,
                "filename": f"schedule_report_{safe or 'all'}",
                "title": f"{label} Schedule Report",
                "meta": {
                    "schedule_label": label,
                    "from": sch_fd,
                    "to": sch_td,
                    "page_layout": page_layout,
                    "dm_style": dm_style,
                    "dm_borders": dm_borders,
                    "date_range": f"{sch_fd or '…'} → {sch_td or '…'}",
                },
            }
        except Exception as exc:
            print(f"[export] online schedule report: {exc}")

    where = ["1=1"]
    params: list[Any] = []
    if sch_fd:
        where.append("s.bill_date>=?")
        params.append(sch_fd)
    if sch_td:
        where.append("s.bill_date<=?")
        params.append(sch_td)
    sch_sql, sch_params = _schedule_sql_filter(
        {"mode": mode, "schedules": picked}
    )
    where_sql = " AND ".join(where) + sch_sql
    params.extend(sch_params)

    raw = conn.execute(
        f"""
        SELECT s.bill_date, s.bill_no, COALESCE(c.name,''), COALESCE(s.doctor_name,''),
               COALESCE(m.name,''), COALESCE(m.batch_no,''), COALESCE(m.content_drug,''),
               COALESCE(m.schedule,''), m.expiry_date,
               COALESCE(si.qty,0), COALESCE(si.rate,0), COALESCE(si.amount,0)
        FROM sales s
        LEFT JOIN customers c ON s.customer_id=c.id
        JOIN sales_items si ON s.id=si.sale_id
        JOIN medicines m ON si.medicine_id=m.id
        WHERE COALESCE(s.deleted,0)=0 AND COALESCE(s.is_autosave,0)=0
          AND {where_sql}
        ORDER BY m.schedule, s.bill_date, c.name
        """,
        params,
    ).fetchall()

    cols = [
        "Date",
        "Bill No",
        "Customer",
        "Doctor",
        "Medicine",
        "Batch",
        "Content/Drug",
        "Schedule",
        "Expiry",
        "Qty",
        "Rate",
        "Amount",
    ]
    rows: list[list[Any]] = []
    for r in raw:
        rows.append(
            [
                r[0],
                r[1],
                r[2],
                r[3] or "—",
                r[4],
                r[5] or "—",
                r[6] or "—",
                r[7] or "—",
                _fmt_schedule_expiry(r[8]),
                int(r[9] or 0),
                f"{_safe_float(r[10]):.2f}",
                f"{_safe_float(r[11]):.2f}",
            ]
        )

    # Portrait / styled: drop Rate/Amount like classic vertical layout.
    pl = str(page_layout or "").strip().lower()
    if rows and not pl.startswith("land"):
        drop = {"Rate", "Amount"}
        keep = [i for i, c in enumerate(cols) if c not in drop]
        cols = [cols[i] for i in keep]
        rows = [[row[i] for i in keep] for row in rows]

    cols, rows = _filter_export_cols("sales_history", "schedule_report", cols, rows)
    rows = _strip_fy_from_bill_columns(cols, rows)
    safe = "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in label)[:40]
    dm_style = str(choice.get("dm_style") or "classic")
    dm_borders = bool(choice.get("dm_borders", True))
    return {
        "columns": cols,
        "rows": rows,
        "filename": f"schedule_report_{safe or 'all'}",
        "title": f"{label} Schedule Report",
        "meta": {
            "schedule_label": label,
            "from": sch_fd,
            "to": sch_td,
            "page_layout": page_layout,
            "dm_style": dm_style,
            "dm_borders": dm_borders,
            "date_range": f"{sch_fd or '…'} → {sch_td or '…'}",
        },
    }


def _purchase_export(conn, report: str, fd: str, td: str) -> dict[str, Any]:
    if _is_online():
        try:
            from core import store_query_client as sq
            from core.due_fifo import purchase_entry_paid
            from core.fy_serial import display_purchase_no
            from core.online_catalog import suppliers as oc_suppliers
            from core.online_mutation_queue import (
                merge_server_rows,
                overlay_purchase_dicts,
            )

            listed = list(
                (sq.list_purchases(from_date=fd or "", to_date=td or "", limit=5000) or {}).get(
                    "rows"
                )
                or []
            )
            raw = merge_server_rows(
                listed, overlay_purchase_dicts(), collection="purchases"
            )
            raw = [
                r
                for r in raw
                if isinstance(r, dict)
                and not r.get("deleted")
                and not int(r.get("is_autosave") or 0)
            ]
            if report == "purchase_register":
                cols = [
                    "Bill No", "Date", "Supplier", "Phone", "Final Amount",
                    "Paid at Entry", "Cash Paid", "Online Paid", "Returns",
                ]
                rows = []
                for r in raw:
                    pno = str(r.get("purchase_no") or "")
                    bno = str(r.get("bill_number") or "") or display_purchase_no(pno)
                    rows.append([
                        bno,
                        r.get("purchase_date") or "",
                        r.get("supplier_name") or "",
                        r.get("supplier_phone") or r.get("phone") or "",
                        _safe_float(r.get("final_amount") or r.get("total_amount")),
                        purchase_entry_paid(r),
                        _safe_float(r.get("cash_paid_at_entry")),
                        _safe_float(r.get("online_paid_at_entry")),
                        _safe_float(
                            r.get("returns_amount")
                            or r.get("refund_amount")
                            or r.get("returns")
                        ),
                    ])
                cols, rows = _filter_export_cols(
                    "purchase_history", "purchase_register", cols, rows
                )
                return {
                    "columns": cols, "rows": rows,
                    "filename": "purchase_register", "title": "Purchase Register",
                }
            if report == "monthly_summary":
                buckets: dict[str, list] = {}
                for r in raw:
                    ym = str(r.get("purchase_date") or "")[:7]
                    if not ym:
                        continue
                    b = buckets.setdefault(ym, [0, 0.0, 0.0])
                    b[0] += 1
                    b[1] += _safe_float(r.get("final_amount") or r.get("total_amount"))
                    b[2] += purchase_entry_paid(r)
                cols = ["Month", "Purchases", "Final Amount", "Paid at Entry"]
                rows = [[k, *v] for k, v in sorted(buckets.items(), reverse=True)]
                cols, rows = _filter_export_cols(
                    "purchase_history", "monthly_summary", cols, rows
                )
                return {
                    "columns": cols, "rows": rows,
                    "filename": "purchase_monthly_summary", "title": "Monthly Summary",
                }
            if report == "supplier_due":
                cols = ["Name", "Phone", "Total Due", "Credit"]
                rows = []
                dues: dict[int, float] = {}
                try:
                    from core.due_fifo import (
                        fifo_paid_via_by_bill,
                        purchase_remaining_and_via,
                    )
                    from core.online_mutation_queue import overlay_supplier_payment_dicts

                    pays = merge_server_rows(
                        list(
                            (sq.list_supplier_payments(limit=5000) or {}).get("rows")
                            or []
                        ),
                        overlay_supplier_payment_dicts(),
                        collection="supplier_payments",
                    )
                    fifo = fifo_paid_via_by_bill(raw, pays)
                    for r in raw:
                        try:
                            sid = int(r.get("supplier_id") or 0)
                            pid = int(r.get("id") or r.get("local_id") or 0)
                        except (TypeError, ValueError):
                            continue
                        if sid <= 0:
                            continue
                        rem, _via = purchase_remaining_and_via(r, fifo.get(pid, 0.0))
                        dues[sid] = round(dues.get(sid, 0.0) + rem, 2)
                except Exception:
                    dues = {}
                for s in oc_suppliers():
                    if not isinstance(s, dict):
                        continue
                    try:
                        sid = int(s.get("id") or s.get("local_id") or 0)
                    except (TypeError, ValueError):
                        sid = 0
                    due = dues.get(sid)
                    if due is None:
                        due = _safe_float(s.get("total_due"))
                    if due <= 0.01:
                        continue
                    rows.append([
                        s.get("name") or "",
                        s.get("phone") or "",
                        due,
                        0.0 if due > 0.01 else _safe_float(
                            s.get("total_credit") or s.get("credit")
                        ),
                    ])
                cols, rows = _filter_export_cols("suppliers", "supplier_due", cols, rows)
                return {
                    "columns": cols, "rows": rows,
                    "filename": "supplier_due", "title": "Supplier Due Report",
                }
            if report == "gst_purchase":
                from core.server_crud import get_doc

                cols = [
                    "Bill No",
                    "Date",
                    "Supplier",
                    "Medicine",
                    "HSN",
                    "GST%",
                    "Qty",
                    "Rate",
                    "Amount",
                ]
                rows = []
                for r in raw:
                    pid = int(r.get("id") or 0)
                    if pid <= 0:
                        continue
                    try:
                        doc = get_doc("purchases", pid) or {}
                    except Exception:
                        doc = {}
                    if not (doc.get("items") or []):
                        try:
                            extra = sq.get_purchase(pid) or {}
                            if isinstance(extra, dict):
                                doc = extra or doc
                        except Exception:
                            pass
                    pno = str(doc.get("purchase_no") or r.get("purchase_no") or "")
                    bno = (
                        str(doc.get("bill_number") or r.get("bill_number") or "").strip()
                        or display_purchase_no(pno)
                    )
                    pdate = str(
                        doc.get("purchase_date") or r.get("purchase_date") or ""
                    ).strip().split(" ")[0]
                    supplier = (
                        str(doc.get("supplier_name") or r.get("supplier_name") or "")
                    )
                    for it in doc.get("items") or doc.get("medicines") or []:
                        if not isinstance(it, dict):
                            continue
                        rows.append([
                            bno,
                            pdate,
                            supplier,
                            it.get("name") or it.get("medicine_name") or "",
                            it.get("hsn_code") or it.get("hsn") or "",
                            _safe_float(
                                it.get("gst_pct")
                                if it.get("gst_pct") is not None
                                else it.get("gst_percent")
                            ),
                            _safe_float(it.get("qty") or it.get("quantity")),
                            _safe_float(it.get("rate")),
                            _safe_float(
                                it.get("item_amount")
                                if it.get("item_amount") is not None
                                else it.get("amount")
                            ),
                        ])
                cols, rows = _filter_export_cols(
                    "purchase_history", "gst_purchase", cols, rows
                )
                return {
                    "columns": cols,
                    "rows": rows,
                    "filename": "gst_purchase",
                    "title": "GST Purchase",
                }
        except Exception as exc:
            return {"error": str(exc), "columns": [], "rows": []}

    where = ["COALESCE(p.deleted,0)=0", "COALESCE(p.is_autosave,0)=0"]
    params: list[Any] = []
    try:
        cols_p = {r[1] for r in conn.execute("PRAGMA table_info(purchases)")}
        if "is_autosave" in cols_p:
            where.append("COALESCE(p.is_autosave,0)=0")
    except Exception:
        pass
    if fd:
        where.append("p.purchase_date>=?")
        params.append(fd)
    if td:
        where.append("p.purchase_date<=?")
        params.append(td)
    wsql = " AND ".join(where)

    if report == "purchase_register":
        cols = [
            "Bill No",
            "Date",
            "Supplier",
            "Phone",
            "Final Amount",
            "Paid at Entry",
            "Cash Paid",
            "Online Paid",
            "Returns",
        ]
        raw = conn.execute(
            f"""
            SELECT COALESCE(p.bill_number, p.purchase_no), p.purchase_date,
                   COALESCE(s.name,''), COALESCE(s.phone,''),
                   COALESCE(p.final_amount, p.total_amount, 0),
                   COALESCE(p.amount_paid_at_entry, p.amount_paid, 0),
                   COALESCE(p.cash_paid_at_entry, 0),
                   COALESCE(p.online_paid_at_entry, 0),
                   COALESCE((SELECT SUM(pr.refund_amount) FROM purchase_returns pr
                             WHERE pr.purchase_id=p.id), 0)
            FROM purchases p LEFT JOIN suppliers s ON p.supplier_id=s.id
            WHERE {wsql}
            ORDER BY p.purchase_date DESC
            """,
            params,
        ).fetchall()
        rows = [list(r) for r in raw]
        cols, rows = _filter_export_cols(
            "purchase_history", "purchase_register", cols, rows
        )
        return {
            "columns": cols,
            "rows": rows,
            "filename": "purchase_register",
            "title": "Purchase Register",
        }

    if report == "monthly_summary":
        cols = ["Month", "Purchases", "Final Amount", "Paid at Entry"]
        raw = conn.execute(
            f"""
            SELECT strftime('%Y-%m', p.purchase_date), COUNT(*),
                   SUM(COALESCE(p.final_amount, p.total_amount, 0)),
                   SUM(COALESCE(p.amount_paid_at_entry, p.amount_paid, 0))
            FROM purchases p WHERE {wsql}
            GROUP BY strftime('%Y-%m', p.purchase_date) ORDER BY 1 DESC
            """,
            params,
        ).fetchall()
        rows = [list(r) for r in raw]
        cols, rows = _filter_export_cols(
            "purchase_history", "monthly_summary", cols, rows
        )
        return {
            "columns": cols,
            "rows": rows,
            "filename": "purchase_monthly_summary",
            "title": "Monthly Summary",
        }

    if report == "supplier_due":
        cols = ["Name", "Phone", "Total Due", "Credit"]
        raw = conn.execute(
            """
            SELECT name, COALESCE(phone,''), COALESCE(total_due,0), COALESCE(credit,0)
            FROM suppliers
            WHERE COALESCE(deleted,0)=0 AND COALESCE(total_due,0)>0
            ORDER BY name
            """
        ).fetchall()
        rows = [list(r) for r in raw]
        cols, rows = _filter_export_cols("suppliers", "supplier_due", cols, rows)
        return {
            "columns": cols,
            "rows": rows,
            "filename": "supplier_due",
            "title": "Supplier Due",
        }

    if report == "gst_purchase":
        cols = [
            "Bill No",
            "Date",
            "Supplier",
            "Medicine",
            "HSN",
            "GST%",
            "Qty",
            "Rate",
            "Amount",
        ]
        # purchase_items has no medicine_name, hsn or gst_percent (db_setup: hsn_code,
        # gst_pct, item_amount), so this query failed on every offline shop with "no
        # such column: pi.medicine_name". The name comes from medicines; the rest are the
        # same fields the Online branch above reads off the purchase document.
        item_where = ""
        try:
            if "deleted" in {r[1] for r in conn.execute("PRAGMA table_info(purchase_items)")}:
                item_where = " AND COALESCE(pi.deleted,0)=0"
        except Exception:
            pass
        raw = conn.execute(
            f"""
            SELECT COALESCE(p.bill_number, p.purchase_no), p.purchase_date,
                   COALESCE(s.name,''), COALESCE(m.name,''),
                   COALESCE(pi.hsn_code,''), COALESCE(pi.gst_pct,0),
                   COALESCE(pi.qty,0), COALESCE(pi.rate,0),
                   COALESCE(pi.item_amount, pi.amount, 0)
            FROM purchase_items pi
            JOIN purchases p ON p.id=pi.purchase_id
            LEFT JOIN suppliers s ON p.supplier_id=s.id
            LEFT JOIN medicines m ON pi.medicine_id=m.id
            WHERE {wsql}{item_where}
            ORDER BY p.purchase_date DESC
            """,
            params,
        ).fetchall()
        rows = [list(r) for r in raw]
        cols, rows = _filter_export_cols("purchase_history", "gst_purchase", cols, rows)
        return {
            "columns": cols,
            "rows": rows,
            "filename": "gst_purchase",
            "title": "GST Purchase",
        }

    return {"error": f"Unknown purchase report {report}", "columns": [], "rows": []}


def _schedule_styled_pdf_bytes(
    data: dict[str, Any],
    *,
    do_print: bool = False,
) -> dict[str, Any]:
    """Build Classic-style schedule HTML/PDF; optionally RAW-print on dot matrix."""
    import base64
    import os

    from core.document_output import (
        build_schedule_report_html,
        prepare_schedule_report_display,
        save_schedule_report_document,
        schedule_report_plain_from_display,
        schedule_report_title,
    )

    cols = list(data.get("columns") or [])
    rows = [list(r) for r in (data.get("rows") or [])]
    meta = data.get("meta") if isinstance(data.get("meta"), dict) else {}
    page_layout = str(meta.get("page_layout") or "portrait")
    dm_style = str(meta.get("dm_style") or "classic")
    dm_borders = bool(meta.get("dm_borders", True))
    sch_label = str(meta.get("schedule_label") or data.get("title") or "Schedule")
    date_range = str(meta.get("date_range") or "")
    title = schedule_report_title(sch_label)

    # Apply Classic/Sign reshape on vertical layouts (styled + portrait) for PDF + DM.
    hdrs, shaped = cols, rows
    if not str(page_layout).lower().startswith("land"):
        try:
            from core.dot_matrix_print import _reshape_schedule_for_dm_style

            hdrs, shaped = _reshape_schedule_for_dm_style(cols, rows, dm_style)
        except Exception:
            hdrs, shaped = cols, rows

    qty_idx = hdrs.index("Qty") if "Qty" in hdrs else None
    amt_idx = hdrs.index("Amount") if "Amount" in hdrs else None
    total_qty = None
    total_amount = None
    try:
        if qty_idx is not None:
            total_qty = sum(int(float(r[qty_idx] or 0)) for r in shaped)
        if amt_idx is not None:
            total_amount = sum(float(r[amt_idx] or 0) for r in shaped)
    except Exception:
        pass

    html = build_schedule_report_html(
        sch_label=sch_label,
        date_range=date_range,
        headers=hdrs,
        table_rows=shaped,
        total_qty=total_qty,
        total_amount=total_amount,
        qty_idx=qty_idx,
        amt_idx=amt_idx,
        page_layout=page_layout,
        report_title=title,
    )
    suffix = "L" if str(page_layout).startswith("land") else (
        "S" if str(page_layout).startswith("styl") else "P"
    )
    base_name = f"{data.get('filename') or 'schedule_report'}_{suffix}"
    pdf_path, html_path = save_schedule_report_document(
        html, base_name, page_layout=page_layout,
    )
    target = pdf_path or html_path
    if not target or not os.path.isfile(target):
        return {
            "ok": False,
            "error": "Could not save schedule PDF. Install Edge or Chrome.",
        }

    printed = False
    print_error = ""
    if do_print:
        try:
            from core.printer_manager import PrinterManager

            if PrinterManager.is_dot_matrix_mode():
                display = prepare_schedule_report_display(
                    hdrs,
                    shaped,
                    page_layout=page_layout,
                    total_qty=total_qty,
                    total_amount=total_amount,
                    qty_idx=qty_idx,
                    amt_idx=amt_idx,
                )
                plain = schedule_report_plain_from_display(display)
                from core.dot_matrix_print import print_schedule_report_dot_matrix

                print_schedule_report_dot_matrix(
                    title,
                    f"Period: {date_range} | Records: {len(shaped)}",
                    plain,
                    page_layout=page_layout,
                    dm_style=dm_style,
                    borders=dm_borders,
                )
                printed = True
            elif pdf_path and os.name == "nt":
                os.startfile(pdf_path, "print")
                printed = True
            else:
                # Open file as last resort (browser/PDF viewer print).
                if target:
                    os.startfile(target) if os.name == "nt" else None
                    printed = True
        except Exception as exc:
            print_error = str(exc)

    with open(target, "rb") as fh:
        payload = fh.read()
    mime = "application/pdf" if target.lower().endswith(".pdf") else "text/html"
    ext = "pdf" if mime.endswith("pdf") else "html"
    return {
        "ok": True,
        "filename": f"{base_name}.{ext}",
        "mime": mime,
        "format": ext,
        "content_base64": base64.b64encode(payload).decode("ascii"),
        "row_count": len(shaped),
        "pdf_path": pdf_path or "",
        "html_path": html_path or "",
        "path": target,
        "printed": printed,
        "print_error": print_error,
        "page_layout": page_layout,
        "dm_style": dm_style,
    }


def export_to_file(
    conn,
    page: str,
    report: str,
    fmt: str,
    *,
    from_date: str = "",
    to_date: str = "",
    current_columns: Optional[list[str]] = None,
    current_rows: Optional[list[list[Any]]] = None,
    schedule: Optional[dict[str, Any]] = None,
    do_print: bool = False,
) -> dict[str, Any]:
    """Run export and return downloadable file bytes (PDF / XLSX / CSV)."""
    import base64
    import csv
    import io
    import os
    import tempfile

    data = run_export(
        conn,
        page,
        report,
        from_date=from_date,
        to_date=to_date,
        current_columns=current_columns,
        current_rows=current_rows,
        schedule=schedule,
    )
    if data.get("error"):
        return {"ok": False, "error": data["error"]}
    cols = list(data.get("columns") or [])
    rows = [_flatten_export_row(r) for r in (data.get("rows") or [])]
    if not rows:
        return {"ok": False, "error": "No records to export."}

    title = str(data.get("title") or data.get("filename") or "Export")
    base_name = str(data.get("filename") or f"{page}_export")
    format_key = (fmt or "csv").strip().lower()
    if format_key == "excel":
        format_key = "xlsx"

    # Schedule report PDF/print → Classic styled HTML path (vertical + DM styles).
    if report == "schedule_report" and (format_key == "pdf" or do_print):
        return _schedule_styled_pdf_bytes(data, do_print=do_print)

    if format_key == "csv":
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow(cols)
        writer.writerows(rows)
        payload = buf.getvalue().encode("utf-8-sig")
        ext = "csv"
        mime = "text/csv;charset=utf-8"
    elif format_key == "xlsx":
        try:
            import openpyxl
        except ImportError:
            return {"ok": False, "error": "openpyxl not installed for Excel export."}
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = title[:31]
        for ci, h in enumerate(cols, 1):
            ws.cell(row=1, column=ci, value=h)
        for ri, row in enumerate(rows, 2):
            for ci, val in enumerate(row, 1):
                ws.cell(row=ri, column=ci, value=val)
        bio = io.BytesIO()
        wb.save(bio)
        payload = bio.getvalue()
        ext = "xlsx"
        mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    elif format_key == "pdf":
        from core.export_manager import _save_pdf_to_path

        fd, path = tempfile.mkstemp(suffix=".pdf", prefix="desktop_export_")
        os.close(fd)
        saved = _save_pdf_to_path(path, title, cols, rows)
        with open(saved, "rb") as fh:
            payload = fh.read()
        for p in (saved, path):
            try:
                if p and os.path.isfile(p):
                    os.remove(p)
            except OSError:
                pass
        html_path = saved.replace(".pdf", ".html") if saved else ""
        if html_path and os.path.isfile(html_path):
            try:
                os.remove(html_path)
            except OSError:
                pass
        ext = "pdf"
        mime = "application/pdf"
    else:
        return {"ok": False, "error": f"Unknown export format: {fmt}"}

    return {
        "ok": True,
        "filename": f"{base_name}.{ext}",
        "mime": mime,
        "format": format_key,
        "content_base64": base64.b64encode(payload).decode("ascii"),
        "row_count": len(rows),
    }
