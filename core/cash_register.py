"""Cash & Expenses: the day's cash register (owner, 25 Oct 2026).

"Enter the cash which is present in the store when I open it, then all the sales bills and
sales returns with the cash get noted into this, enter the expense like tea or medicines from
another medical store, and at the end of the day see the list and a report whether the actual
cash in my store is proper or not."

One register per date. What the shop TYPES is stored: the opening cash, the manual lines (an
expense, or cash the owner added) and the counted cash at close -- in ``cash_days`` and
``cash_expenses``, two synced tables (offline-first triggers, outbox, the server's collections
of the same names), so the phone and the web can show the register later. What the books
already hold is NOT stored again: the cash part of every bill, receipt, return, purchase and
supplier payment of the date is read from those tables each time, so the register always
shows the same figures as the bills (a deleted bill leaves it; an edited one follows).

    expected cash = opening
                    + cash on bills + cash receipts + purchase returns refunded in cash + cash added
                    - cash refunded on sales returns - cash paid on purchases at entry
                    - cash paid to suppliers - expenses

Routes (core/desktop_api.py hands /api/cash/* here):
    GET  /api/cash/day?date=          the register of one date
    GET  /api/cash/range?from=&to=    a month: one line per day, expenses by category
    POST /api/cash/opening            {date, opening_cash}
    POST /api/cash/expense            {date, id?, kind: out|in, category, description, amount, at}
    POST /api/cash/expense/delete     {id}
    POST /api/cash/close              {date, counted_cash, note} / {date, reopen: true}
    POST /api/cash/export             {date} or {from, to}; format xlsx|pdf, print_to printer
"""
from __future__ import annotations

import base64
import re
import sqlite3
from datetime import date, datetime, timedelta, timezone
from typing import Any, Optional

COLLECTIONS = ("cash_days", "cash_expenses")

# Default categories offered in the box (the shop may type any other).
CATEGORIES = ("Tea & snacks", "Medicines bought", "Transport", "Stationery", "Salary / advance",
              "Electricity / rent", "Owner took cash", "Other")

TABLES = (
    """CREATE TABLE IF NOT EXISTS cash_days (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        day DATE NOT NULL,
        opening_cash REAL DEFAULT 0,
        counted_cash REAL,
        closed_at TEXT,
        note TEXT,
        created_at TEXT,
        updated_at TEXT,
        version INTEGER DEFAULT 1,
        device_id TEXT,
        deleted INTEGER DEFAULT 0
    )""",
    "CREATE INDEX IF NOT EXISTS idx_cash_days_day ON cash_days (day)",
    """CREATE TABLE IF NOT EXISTS cash_expenses (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        day DATE NOT NULL,
        at TEXT,
        kind TEXT DEFAULT 'out',
        category TEXT,
        description TEXT,
        amount REAL DEFAULT 0,
        created_at TEXT,
        updated_at TEXT,
        version INTEGER DEFAULT 1,
        device_id TEXT,
        deleted INTEGER DEFAULT 0
    )""",
    "CREATE INDEX IF NOT EXISTS idx_cash_expenses_day ON cash_expenses (day)",
)

_DAY_COLS = ("day", "opening_cash", "counted_cash", "closed_at", "note", "created_at", "updated_at",
             "version", "device_id", "deleted")
_EXP_COLS = ("day", "at", "kind", "category", "description", "amount", "created_at", "updated_at",
             "version", "device_id", "deleted")
_DOC_COLS = {"cash_days": _DAY_COLS, "cash_expenses": _EXP_COLS}

_ISO_DAY = re.compile(r"^\d{4}-\d{2}-\d{2}$")


# ── helpers ───────────────────────────────────────────────────────────────────

def r2(v) -> float:
    try:
        return round(float(v or 0) + 0.0, 2)
    except (TypeError, ValueError):
        return 0.0


def _now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="milliseconds").replace("+00:00", "Z")


def today() -> str:
    return date.today().isoformat()


def clean_day(raw) -> str:
    s = str(raw or "").strip()[:10]
    if not s:
        return today()
    if not _ISO_DAY.match(s):
        raise ValueError("Date must be YYYY-MM-DD")
    date.fromisoformat(s)                      # a real date
    return s


def _dmy(day: str) -> str:
    return f"{day[8:10]}-{day[5:7]}-{day[0:4]}" if day and len(day) >= 10 else str(day or "")


def _local_hm(ts) -> str:
    """'HH:MM' local time of a stored stamp (SQLite CURRENT_TIMESTAMP is UTC; our own stamps
    end in Z). A stamp with no zone and a 'T' is taken as local (the Android apps')."""
    s = str(ts or "").strip()
    if not s:
        return ""
    try:
        if s.endswith("Z"):
            dt = datetime.fromisoformat(s[:-1]).replace(tzinfo=timezone.utc)
        elif "+" in s[10:] or (len(s) > 19 and s[19] == "-"):
            dt = datetime.fromisoformat(s)
        elif "T" in s:
            return s[11:16]
        else:
            dt = datetime.fromisoformat(s).replace(tzinfo=timezone.utc)
        return dt.astimezone().strftime("%H:%M")
    except ValueError:
        return s[11:16] if len(s) >= 16 else ""


def ensure_tables(conn: sqlite3.Connection) -> None:
    for sql in TABLES:
        conn.execute(sql)


def _cols(conn: sqlite3.Connection, table: str) -> set:
    try:
        return {str(r[1]) for r in conn.execute(f"PRAGMA table_info({table})").fetchall()}
    except sqlite3.Error:
        return set()


def _alive(cols: set, alias: str, *flags: str) -> str:
    """'COALESCE(a.deleted,0)=0 AND COALESCE(a.is_autosave,0)=0' for the flags the table has."""
    parts = [f"COALESCE({alias}.{c},0)=0" for c in flags if c in cols]
    return " AND ".join(parts) or "1=1"


def _f(v) -> float:
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        return 0.0


def _rows(conn: sqlite3.Connection, sql: str, params=()) -> list:
    cur = conn.execute(sql, params)
    names = [d[0] for d in cur.description]
    return [dict(zip(names, r)) for r in cur.fetchall()]


# ── the typed part: cash_days / cash_expenses ─────────────────────────────────

def day_row(conn: sqlite3.Connection, day: str) -> Optional[dict]:
    """The register row of a date. Two devices that both opened the day before syncing hold
    two rows; the closed one, else the lowest id, is the one every device shows."""
    ensure_tables(conn)
    rows = _rows(conn, "SELECT * FROM cash_days WHERE day=? AND COALESCE(deleted,0)=0 "
                       "ORDER BY (counted_cash IS NULL), id LIMIT 1", (day,))
    return rows[0] if rows else None


def _stamp_update(conn: sqlite3.Connection, table: str, row_id: int, sets: dict) -> None:
    sets = dict(sets)
    sets["updated_at"] = _now()
    cols = ", ".join(f"{k}=?" for k in sets)
    conn.execute(f"UPDATE {table} SET {cols}, version=COALESCE(version,1)+1 WHERE id=?",
                 (*sets.values(), int(row_id)))


def set_opening(conn: sqlite3.Connection, day: str, opening) -> dict:
    day = clean_day(day)
    opening = r2(opening)
    if opening < 0:
        raise ValueError("Opening cash cannot be negative")
    row = day_row(conn, day)
    if row:
        _stamp_update(conn, "cash_days", row["id"], {"opening_cash": opening})
    else:
        now = _now()
        conn.execute(
            "INSERT INTO cash_days (day, opening_cash, created_at, updated_at, version, deleted) "
            "VALUES (?,?,?,?,1,0)", (day, opening, now, now))
    conn.commit()
    return day_row(conn, day) or {}


def close_day(conn: sqlite3.Connection, day: str, counted, note: str = "") -> dict:
    """Save the cash counted at close. When nobody typed an opening the row keeps it NULL
    (shown as 0), so the screen can still ask for it."""
    day = clean_day(day)
    counted = r2(counted)
    if counted < 0:
        raise ValueError("Counted cash cannot be negative")
    row = day_row(conn, day)
    now = _now()
    if row:
        _stamp_update(conn, "cash_days", row["id"],
                      {"counted_cash": counted, "closed_at": now, "note": str(note or "")[:300]})
    else:
        conn.execute(
            "INSERT INTO cash_days (day, opening_cash, counted_cash, closed_at, note, created_at, updated_at, "
            "version, deleted) VALUES (?,NULL,?,?,?,?,?,1,0)",
            (day, counted, now, str(note or "")[:300], now, now))
    conn.commit()
    return day_row(conn, day) or {}


def reopen_day(conn: sqlite3.Connection, day: str) -> dict:
    day = clean_day(day)
    row = day_row(conn, day)
    if row:
        _stamp_update(conn, "cash_days", row["id"], {"counted_cash": None, "closed_at": None})
        conn.commit()
    return day_row(conn, day) or {}


def save_expense(conn: sqlite3.Connection, body: dict) -> dict:
    """Add or edit a manual line. kind 'out' = an expense (tea, medicines bought from another
    store...), 'in' = cash the owner put in the drawer."""
    ensure_tables(conn)
    day = clean_day(body.get("date") or body.get("day"))
    kind = "in" if str(body.get("kind") or "out").strip().lower() in ("in", "cash_in", "added") else "out"
    amount = r2(body.get("amount"))
    if amount <= 0:
        raise ValueError("Enter the amount")
    desc = " ".join(str(body.get("description") or "").split())[:200]
    category = " ".join(str(body.get("category") or "").split())[:60]
    if not desc and not category:
        raise ValueError("Say what the cash was for (e.g. Tea)")
    if not desc:
        desc = category
    at = str(body.get("at") or "").strip()[:5]
    if not re.match(r"^\d{2}:\d{2}$", at):
        at = datetime.now().strftime("%H:%M")
    row_id = int(body.get("id") or 0)
    if row_id:
        if not conn.execute("SELECT 1 FROM cash_expenses WHERE id=?", (row_id,)).fetchone():
            raise ValueError("That line was removed already")
        _stamp_update(conn, "cash_expenses", row_id,
                      {"day": day, "at": at, "kind": kind, "category": category, "description": desc,
                       "amount": amount})
    else:
        now = _now()
        cur = conn.execute(
            "INSERT INTO cash_expenses (day, at, kind, category, description, amount, created_at, updated_at, "
            "version, deleted) VALUES (?,?,?,?,?,?,?,?,1,0)",
            (day, at, kind, category, desc, amount, now, now))
        row_id = int(cur.lastrowid)
    conn.commit()
    rows = _rows(conn, "SELECT * FROM cash_expenses WHERE id=?", (row_id,))
    return rows[0] if rows else {"id": row_id}


def delete_expense(conn: sqlite3.Connection, row_id) -> bool:
    ensure_tables(conn)
    cur = conn.execute("DELETE FROM cash_expenses WHERE id=?", (int(row_id or 0),))
    conn.commit()
    return bool(cur.rowcount)


def manual_lines(conn: sqlite3.Connection, day: str) -> list:
    ensure_tables(conn)
    out = []
    for r in _rows(conn, "SELECT * FROM cash_expenses WHERE day=? AND COALESCE(deleted,0)=0 "
                         "ORDER BY at, id", (day,)):
        out.append({"id": int(r["id"]), "time": str(r.get("at") or ""), "kind": r.get("kind") or "out",
                    "category": r.get("category") or "", "description": r.get("description") or "",
                    "amount": r2(r.get("amount"))})
    return out


# ── the books: what moved cash on the date ────────────────────────────────────

def _cash_of(cash, online, whole) -> float:
    """The cash part of a payment stored as a cash / online split, or as one amount with the
    mode (older rows, the phone): same rule as the server's purchaseEntryPaid."""
    c, o = _f(cash), _f(online)
    if c or o:
        return c
    return _f(whole)


def cash_sales(conn: sqlite3.Connection, day: str) -> list:
    cols = _cols(conn, "sales")
    if not cols:
        return []
    cust = "COALESCE(NULLIF(s.customer_name,''), c.name, 'Walk-in')" if "customer_name" in cols else "COALESCE(c.name,'Walk-in')"
    rows = _rows(conn, f"""
        SELECT s.id, s.bill_no, s.created_at, s.total_amount, s.cash_paid, s.online_paid, s.amount_paid,
               {cust} AS party
        FROM sales s LEFT JOIN customers c ON c.id = s.customer_id
        WHERE s.bill_date=? AND {_alive(cols, 's', 'deleted', 'is_autosave')}
        ORDER BY s.created_at, s.id""", (day,))
    from core.fy_serial import display_sales_bill_no

    out = []
    for r in rows:
        cash = r2(_cash_of(r.get("cash_paid"), r.get("online_paid"), r.get("amount_paid")))
        if not cash:
            continue
        out.append({"id": int(r["id"]), "time": _local_hm(r.get("created_at")),
                    "ref": display_sales_bill_no(r.get("bill_no")), "party": r.get("party") or "Walk-in",
                    "detail": f"Bill {r2(r.get('total_amount')):.2f}", "amount": cash})
    return out


def _payment_cash(r: dict) -> float:
    c, o = _f(r.get("cash_amount")), _f(r.get("online_amount"))
    if c or o:
        return c
    mode = str(r.get("payment_mode") or "cash").strip().lower()
    return _f(r.get("amount")) if mode in ("cash", "") else 0.0


def _customer_payments(conn: sqlite3.Connection, day: str) -> list:
    cols = _cols(conn, "customer_payments")
    if not cols:
        return []
    return _rows(conn, f"""
        SELECT p.*, c.name AS party FROM customer_payments p LEFT JOIN customers c ON c.id = p.customer_id
        WHERE p.payment_date=? AND {_alive(cols, 'p', 'deleted')} ORDER BY p.created_at, p.id""", (day,))


def cash_receipts(conn: sqlite3.Connection, day: str) -> list:
    """Dues received in cash (Payments, Sales History 'Pay'). Refunds (negative rows) are the
    returns section's."""
    out = []
    for r in _customer_payments(conn, day):
        if _f(r.get("amount")) <= 0:
            continue
        cash = r2(_payment_cash(r))
        if cash <= 0:
            continue
        out.append({"id": int(r["id"]), "time": _local_hm(r.get("created_at")),
                    "ref": str(r.get("reference_no") or r.get("receipt_no") or ""),
                    "party": r.get("party") or "", "detail": str(r.get("note") or "Payment received"),
                    "amount": cash})
    return out


def cash_sales_returns(conn: sqlite3.Connection, day: str) -> list:
    """Cash handed back on sales returns: the refund rows (negative cash receipts, which carry
    what was really handed back once dues were cleared), and a walk-in return settled in cash
    (no customer, so no receipt row: the refund itself)."""
    out = []
    seen_refs = set()
    for r in _customer_payments(conn, day):
        if _f(r.get("amount")) >= 0:
            continue
        cash = r2(-_payment_cash(r))
        if cash <= 0:
            continue
        ref = str(r.get("reference_no") or "")
        seen_refs.add(ref)
        out.append({"id": int(r["id"]), "time": _local_hm(r.get("created_at")), "ref": ref,
                    "party": r.get("party") or "", "detail": str(r.get("note") or "Refund"), "amount": cash})
    cols = _cols(conn, "sales_returns")
    if cols and "settle_mode" in cols:
        for r in _rows(conn, f"""
                SELECT r.id, r.return_no, r.created_at, r.refund_amount, r.customer_id
                FROM sales_returns r
                WHERE r.return_date=? AND LOWER(COALESCE(r.settle_mode,''))='cash' AND COALESCE(r.customer_id, 0) = 0
                  AND {_alive(cols, 'r', 'deleted')} ORDER BY r.created_at, r.id""", (day,)):
            if str(r.get("return_no") or "") in seen_refs:
                continue
            amt = r2(r.get("refund_amount"))
            if amt <= 0:
                continue
            out.append({"id": int(r["id"]), "time": _local_hm(r.get("created_at")),
                        "ref": str(r.get("return_no") or ""), "party": "Walk-in", "detail": "Return refund",
                        "amount": amt})
    return out


def cash_purchases(conn: sqlite3.Connection, day: str) -> list:
    cols = _cols(conn, "purchases")
    if not cols:
        return []
    rows = _rows(conn, f"""
        SELECT p.id, p.purchase_no, p.bill_number, p.created_at, p.final_amount, p.total_amount,
               p.cash_paid_at_entry, p.online_paid_at_entry, p.amount_paid_at_entry, s.name AS party
        FROM purchases p LEFT JOIN suppliers s ON s.id = p.supplier_id
        WHERE p.purchase_date=? AND {_alive(cols, 'p', 'deleted', 'is_autosave')}
        ORDER BY p.created_at, p.id""", (day,))
    out = []
    for r in rows:
        cash = r2(_cash_of(r.get("cash_paid_at_entry"), r.get("online_paid_at_entry"), r.get("amount_paid_at_entry")))
        if cash <= 0:
            continue
        total = _f(r.get("final_amount")) or _f(r.get("total_amount"))
        out.append({"id": int(r["id"]), "time": _local_hm(r.get("created_at")),
                    "ref": str(r.get("bill_number") or r.get("purchase_no") or ""),
                    "party": r.get("party") or "", "detail": f"Purchase {r2(total):.2f}", "amount": cash})
    return out


def _supplier_payments(conn: sqlite3.Connection, day: str) -> list:
    cols = _cols(conn, "supplier_payments")
    if not cols:
        return []
    return _rows(conn, f"""
        SELECT p.*, s.name AS party FROM supplier_payments p LEFT JOIN suppliers s ON s.id = p.supplier_id
        WHERE p.payment_date=? AND LOWER(COALESCE(p.mode,'cash'))='cash' AND {_alive(cols, 'p', 'deleted')}
        ORDER BY p.created_at, p.id""", (day,))


def cash_supplier_payments(conn: sqlite3.Connection, day: str) -> list:
    out = []
    for r in _supplier_payments(conn, day):
        amt = r2(r.get("amount"))
        if amt <= 0:
            continue
        out.append({"id": int(r["id"]), "time": _local_hm(r.get("created_at")),
                    "ref": str(r.get("reference") or r.get("payment_no") or ""), "party": r.get("party") or "",
                    "detail": "Paid to supplier", "amount": amt})
    return out


def cash_purchase_returns(conn: sqlite3.Connection, day: str) -> list:
    """A supplier handing cash back for a purchase return (core/return_refunds.py writes it as
    a negative cash supplier payment)."""
    out = []
    for r in _supplier_payments(conn, day):
        amt = r2(-_f(r.get("amount")))
        if amt <= 0:
            continue
        out.append({"id": int(r["id"]), "time": _local_hm(r.get("created_at")),
                    "ref": str(r.get("reference") or r.get("payment_no") or ""), "party": r.get("party") or "",
                    "detail": "Return refund from supplier", "amount": amt})
    return out


# ── the sums ──────────────────────────────────────────────────────────────────

def expected_cash(opening, *, cash_sales=0.0, cash_receipts=0.0, cash_purchase_returns=0.0, cash_added=0.0,
                  cash_sales_returns=0.0, cash_purchases=0.0, cash_supplier_payments=0.0, expenses=0.0) -> float:
    """What should be in the drawer. Pure arithmetic, so the test can pin it."""
    cash_in = _f(cash_sales) + _f(cash_receipts) + _f(cash_purchase_returns) + _f(cash_added)
    cash_out = _f(cash_sales_returns) + _f(cash_purchases) + _f(cash_supplier_payments) + _f(expenses)
    return r2(_f(opening) + cash_in - cash_out)


def _total(lines: list) -> float:
    return r2(sum(_f(x.get("amount")) for x in lines))


def day_report(conn: sqlite3.Connection, day: str) -> dict:
    day = clean_day(day)
    row = day_row(conn, day)
    manual = manual_lines(conn, day)
    expenses = [m for m in manual if m["kind"] == "out"]
    added = [m for m in manual if m["kind"] == "in"]
    sections = [
        {"key": "sales", "title": "Cash on sales bills", "flow": "in", "lines": cash_sales(conn, day)},
        {"key": "receipts", "title": "Dues received in cash", "flow": "in", "lines": cash_receipts(conn, day)},
        {"key": "sales_returns", "title": "Cash refunded on sales returns", "flow": "out",
         "lines": cash_sales_returns(conn, day)},
        {"key": "purchases", "title": "Cash paid on purchases (at entry)", "flow": "out",
         "lines": cash_purchases(conn, day)},
        {"key": "supplier_payments", "title": "Cash paid to suppliers", "flow": "out",
         "lines": cash_supplier_payments(conn, day)},
        {"key": "purchase_returns", "title": "Cash back from suppliers (purchase returns)", "flow": "in",
         "lines": cash_purchase_returns(conn, day)},
    ]
    for s in sections:
        s["total"] = _total(s["lines"])
    by = {s["key"]: s["total"] for s in sections}
    opening = r2(row.get("opening_cash")) if row and row.get("opening_cash") is not None else 0.0
    exp_total, add_total = _total(expenses), _total(added)
    expected = expected_cash(opening, cash_sales=by["sales"], cash_receipts=by["receipts"],
                             cash_purchase_returns=by["purchase_returns"], cash_added=add_total,
                             cash_sales_returns=by["sales_returns"], cash_purchases=by["purchases"],
                             cash_supplier_payments=by["supplier_payments"], expenses=exp_total)
    counted = r2(row["counted_cash"]) if row and row.get("counted_cash") is not None else None
    return {
        "ok": True,
        "date": day,
        "date_dmy": _dmy(day),
        # set: somebody typed the opening (a day closed without one keeps it NULL, shown as 0).
        "opening": {"value": opening, "set": bool(row and row.get("opening_cash") is not None),
                    "row_id": int(row["id"]) if row else None},
        "sections": sections,
        "expenses": expenses,
        "cash_added": added,
        "totals": {
            "cash_in": r2(by["sales"] + by["receipts"] + by["purchase_returns"] + add_total),
            "cash_out": r2(by["sales_returns"] + by["purchases"] + by["supplier_payments"]),
            "expenses": exp_total,
            "cash_added": add_total,
            "expected": expected,
        },
        "counted": counted,
        "closed_at": (row or {}).get("closed_at"),
        "note": (row or {}).get("note") or "",
        "difference": r2(counted - expected) if counted is not None else None,
        "categories": list(CATEGORIES),
    }


def _days_between(start: str, end: str) -> list:
    a, b = date.fromisoformat(start), date.fromisoformat(end)
    if b < a:
        a, b = b, a
    if (b - a).days > 92:
        raise ValueError("Pick up to three months at a time")
    return [(a + timedelta(days=i)).isoformat() for i in range((b - a).days + 1)]


def range_report(conn: sqlite3.Connection, start: str, end: str) -> dict:
    """One line per day (days with nothing are left out), expenses by category, every expense."""
    start, end = clean_day(start), clean_day(end)
    days, cats, lines = [], {}, []
    for d in _days_between(start, end):
        rep = day_report(conn, d)
        t = rep["totals"]
        has = any(s["lines"] for s in rep["sections"]) or rep["expenses"] or rep["cash_added"] or rep["opening"]["set"] \
            or rep["counted"] is not None
        if not has:
            continue
        days.append({"date": d, "date_dmy": rep["date_dmy"], "opening": rep["opening"]["value"],
                     "cash_in": t["cash_in"], "cash_out": t["cash_out"], "expenses": t["expenses"],
                     "expected": t["expected"], "counted": rep["counted"], "difference": rep["difference"],
                     "closed": bool(rep["closed_at"])})
        for e in rep["expenses"]:
            key = e["category"] or "Other"
            cats[key] = r2(cats.get(key, 0) + e["amount"])
            lines.append({**e, "date": d, "date_dmy": rep["date_dmy"]})
    tot = {
        "opening": r2(sum(x["opening"] for x in days)),
        "cash_in": r2(sum(x["cash_in"] for x in days)),
        "cash_out": r2(sum(x["cash_out"] for x in days)),
        "expenses": r2(sum(x["expenses"] for x in days)),
        # Only the days that were counted; None when no day was (the screen shows a dash, not
        # "Tallies").
        "difference": (r2(sum(x["difference"] for x in days if x["difference"] is not None))
                       if any(x["difference"] is not None for x in days) else None),
        "days_closed": sum(1 for x in days if x["closed"]),
    }
    return {"ok": True, "from": start, "to": end, "from_dmy": _dmy(start), "to_dmy": _dmy(end), "days": days,
            "categories": [{"category": k, "amount": v} for k, v in sorted(cats.items(), key=lambda kv: -kv[1])],
            "expenses": lines, "totals": tot}


# ── export: Excel / PDF / Print ───────────────────────────────────────────────

def _day_sections(rep: dict) -> list:
    cols = ["Time", "Bill / Ref", "Party", "Detail", "Amount"]
    out = [{"title": "Opening cash", "columns": ["What", "Amount"],
            "rows": [["Cash in the drawer when the shop opened", rep["opening"]["value"]]]}]
    for s in rep["sections"]:
        rows = [[l["time"], l["ref"], l["party"], l["detail"], l["amount"]] for l in s["lines"]]
        if rows:
            rows.append(["", "", "", "Total", s["total"]])
            out.append({"title": f"{s['title']} ({'cash in' if s['flow'] == 'in' else 'cash out'})",
                        "columns": cols, "rows": rows})
    for key, title in (("expenses", "Expenses (cash out)"), ("cash_added", "Cash added (cash in)")):
        rows = [[e["time"], e["category"], e["description"], e["amount"]] for e in rep[key]]
        if rows:
            rows.append(["", "", "Total", _total(rep[key])])
            out.append({"title": title, "columns": ["Time", "Category", "Description", "Amount"], "rows": rows})
    t = rep["totals"]
    summary = [["Opening cash", rep["opening"]["value"]], ["+ Cash in (bills, dues, suppliers' refunds, cash added)", t["cash_in"]],
               ["- Cash out (returns, purchases, suppliers)", t["cash_out"]], ["- Expenses", t["expenses"]],
               ["= Expected cash in drawer", t["expected"]]]
    if rep["counted"] is not None:
        summary += [["Counted cash", rep["counted"]], ["Difference (counted - expected)", rep["difference"]]]
        if rep.get("note"):
            summary.append(["Note", rep["note"]])
    else:
        summary.append(["Counted cash", "not counted yet"])
    out.append({"title": "Summary", "columns": ["What", "Amount"], "rows": summary})
    return out


def _range_sections(rep: dict) -> list:
    rows = [[d["date_dmy"], d["opening"], d["cash_in"], d["cash_out"], d["expenses"], d["expected"],
             d["counted"] if d["counted"] is not None else "", d["difference"] if d["difference"] is not None else "",
             "Closed" if d["closed"] else "Open"] for d in rep["days"]]
    t = rep["totals"]
    rows.append(["Total", t["opening"], t["cash_in"], t["cash_out"], t["expenses"], "", "",
                 t["difference"] if t["difference"] is not None else "", ""])
    out = [{"title": "Day by day", "columns": ["Date", "Opening", "Cash in", "Cash out", "Expenses", "Expected",
                                              "Counted", "Difference", "Status"], "rows": rows}]
    cats = [[c["category"], c["amount"]] for c in rep["categories"]]
    if cats:
        cats.append(["Total", t["expenses"]])
        out.append({"title": "Expenses by category", "columns": ["Category", "Amount"], "rows": cats})
    lines = [[e["date_dmy"], e["time"], e["category"], e["description"], e["amount"]] for e in rep["expenses"]]
    if lines:
        out.append({"title": "Every expense", "columns": ["Date", "Time", "Category", "Description", "Amount"],
                    "rows": lines})
    return out


def _xlsx_bytes(title: str, sections: list) -> bytes:
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Cash"
    bold = Font(bold=True)
    widths: dict = {}
    n = 1
    ws.cell(row=n, column=1, value=title).font = Font(bold=True, size=13)
    n += 2
    for s in sections:
        ws.cell(row=n, column=1, value=s["title"]).font = bold
        n += 1
        for j, c in enumerate(s["columns"], start=1):
            ws.cell(row=n, column=j, value=c).font = bold
            widths[j] = max(widths.get(j, 8), len(str(c)) + 2)
        n += 1
        for r in s["rows"]:
            is_total = str(r[0]) == "Total" or (len(r) > 1 and str(r[-2]) == "Total")
            for j, v in enumerate(r, start=1):
                cell = ws.cell(row=n, column=j, value=None if v == "" else v)
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    cell.number_format = "0.00"
                if is_total:
                    cell.font = bold
                if v not in ("", None):
                    widths[j] = min(max(widths.get(j, 8), len(str(v)) + 2), 48)
            n += 1
        n += 1
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _shop_name(conn: sqlite3.Connection) -> str:
    try:
        row = conn.execute("SELECT name FROM pharmacy_profile ORDER BY id LIMIT 1").fetchone()
        return str(row[0] or "") if row else ""
    except sqlite3.Error:
        return ""


def export(conn: sqlite3.Connection, body: dict) -> dict:
    fmt = str(body.get("format") or "pdf").strip().lower()
    print_to = str(body.get("print_to") or "").strip().lower()
    if body.get("from") or body.get("to"):
        rep = range_report(conn, body.get("from") or body.get("to"), body.get("to") or body.get("from"))
        sections = _range_sections(rep)
        stem = f"Cash_Register_{rep['from_dmy']}_{rep['to_dmy']}"
        title = f"Cash & Expenses {rep['from_dmy']} - {rep['to_dmy']}"
    else:
        rep = day_report(conn, body.get("date"))
        sections = _day_sections(rep)
        stem = f"Cash_Register_{rep['date_dmy']}"
        title = f"Cash & Expenses {rep['date_dmy']}"
    shop = _shop_name(conn)
    if shop:
        title = f"{title} -- {shop}"
    if fmt in ("xlsx", "excel") and not print_to:
        data = _xlsx_bytes(title, sections)
        return {"ok": True, "filename": f"{stem}.xlsx", "format": "xlsx",
                "mime": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "content_base64": base64.b64encode(data).decode("ascii")}
    from core.desktop_export_service import export_alert_sections

    res = export_alert_sections({"title": title, "sections": sections, "format": "pdf" if fmt not in ("csv",) else fmt,
                                 "print_to": print_to, "page_layout": body.get("page_layout") or "portrait",
                                 "paper": body.get("paper")})
    if res.get("ok") and res.get("filename"):
        res["filename"] = f"{stem}.{res.get('format') or 'pdf'}"
    return res


# ── sync: the document the server keeps, and a pulled one written here ────────

def build_doc(conn: sqlite3.Connection, collection: str, local_id: int) -> Optional[dict]:
    """The record as the server's collection takes it (server-live extraUpserts.js), or None
    when it is gone from this PC."""
    if collection not in COLLECTIONS:
        return None
    ensure_tables(conn)
    rows = _rows(conn, f"SELECT * FROM {collection} WHERE id=?", (int(local_id),))
    if not rows:
        return None
    r = rows[0]
    doc: dict[str, Any] = {"id": int(r["id"])}
    for c in _DOC_COLS[collection]:
        doc[c] = r.get(c)
    doc["version"] = int(r.get("version") or 1)
    doc["deleted"] = bool(r.get("deleted"))
    doc["device_id"] = str(r.get("device_id") or "")
    if not doc.get("updated_at"):
        doc["updated_at"] = r.get("created_at") or _now()
    return doc


def apply_pulled(conn: sqlite3.Connection, collection: str, local_id: int, doc, *, delete: bool = False) -> str:
    """Write another device's record (call with capture paused). A delete, or deleted: true,
    removes the row."""
    if collection not in COLLECTIONS:
        return "ignored"
    ensure_tables(conn)
    lid = int(local_id or (doc or {}).get("id") or 0)
    if lid <= 0:
        return "ignored"
    if delete or not isinstance(doc, dict) or doc.get("deleted") in (True, 1, "1", "true"):
        conn.execute(f"DELETE FROM {collection} WHERE id=?", (lid,))
        return "soft_deleted"
    cols = list(_DOC_COLS[collection])
    vals = []
    for c in cols:
        v = doc.get(c)
        if c == "deleted":
            v = 1 if v in (True, 1, "1", "true") else 0
        elif c == "version":
            v = int(v or 1)
        elif c in ("opening_cash", "counted_cash", "amount"):
            v = None if v is None else r2(v)
        elif c == "day":
            v = str(v or "")[:10]
        elif v is not None and not isinstance(v, (int, float, str)):
            v = str(v)
        vals.append(v)
    sets = ", ".join(f"{c}=excluded.{c}" for c in cols)
    conn.execute(
        f"INSERT INTO {collection} (id, {', '.join(cols)}) VALUES (?{', ?' * len(cols)}) "
        f"ON CONFLICT(id) DO UPDATE SET {sets}",
        (lid, *vals))
    return "applied"


# ── routes ────────────────────────────────────────────────────────────────────

def _q(qs: dict, key: str) -> str:
    v = qs.get(key)
    if isinstance(v, (list, tuple)):
        v = v[0] if v else ""
    return str(v or "").strip()


def handle_get(path: str, qs: dict, conn) -> Optional[tuple]:
    if not path.startswith("/api/cash/"):
        return None
    if conn is None:
        return 503, {"ok": False, "error": "Database not open"}
    try:
        if path == "/api/cash/day":
            return 200, day_report(conn, _q(qs, "date"))
        if path == "/api/cash/range":
            return 200, range_report(conn, _q(qs, "from"), _q(qs, "to"))
    except Exception as exc:
        return 400, {"ok": False, "error": str(exc)}
    return None


def _after_write() -> None:
    """Send it now (the worker's next cycle is a few seconds away anyway)."""
    try:
        from core.offline_first import worker

        worker.kick()
    except Exception:
        pass


ONLINE_REFUSAL = ("Cash & Expenses is kept on this PC's own copy, and Online mode has none -- "
                  "the line would be lost on restart. Switch to Offline-first to use it.")


def _online_shell() -> bool:
    try:
        from core.sync_prefs import is_online_mode

        return bool(is_online_mode())
    except Exception:
        return False


def handle_post(path: str, body: dict, conn) -> Optional[tuple]:
    if not path.startswith("/api/cash/"):
        return None
    if conn is None:
        return 503, {"ok": False, "error": "Database not open"}
    body = body if isinstance(body, dict) else {}
    if path != "/api/cash/export" and _online_shell():
        # Online the engine's connection is a throwaway :memory: database (desktop_api
        # _open_memory_shell): an opening, an expense or a close typed there "saved",
        # vanished on the next restart and never reached the other PCs.
        return 409, {"ok": False, "error": ONLINE_REFUSAL}
    try:
        if path == "/api/cash/opening":
            set_opening(conn, body.get("date"), body.get("opening_cash"))
            _after_write()
            return 200, day_report(conn, body.get("date"))
        if path == "/api/cash/expense":
            saved = save_expense(conn, body)
            _after_write()
            out = day_report(conn, saved.get("day") or body.get("date"))
            out["saved_id"] = saved.get("id")
            return 200, out
        if path == "/api/cash/expense/delete":
            row = conn.execute("SELECT day FROM cash_expenses WHERE id=?", (int(body.get("id") or 0),)).fetchone()
            delete_expense(conn, body.get("id"))
            _after_write()
            return 200, day_report(conn, (row[0] if row else None) or body.get("date"))
        if path == "/api/cash/close":
            if body.get("reopen"):
                reopen_day(conn, body.get("date"))
            else:
                close_day(conn, body.get("date"), body.get("counted_cash"), str(body.get("note") or ""))
            _after_write()
            return 200, day_report(conn, body.get("date"))
        if path == "/api/cash/export":
            out = export(conn, body)
            return (200 if out.get("ok") else 400), out
    except Exception as exc:
        return 400, {"ok": False, "error": str(exc)}
    return None
