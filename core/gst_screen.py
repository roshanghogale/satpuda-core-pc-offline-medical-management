"""The GST Reports screen (owner, 22 Oct 2026): every GST report a shop is asked for, listed
the way Marg lists them (A .. W) plus GSTR-1 B2B / B2CL / nil / documents, each one a table
in the government's column order with its totals -- and the Marg "GST upload" registers
(Sale / Purchase / Sale Return / Purchase Return GST) as Excel and CSV in Marg's own layout.

Every figure comes from core.gst_reports: a sale's tax is the printed bill's own
(core.bill_gst, the bill discount spread over the lines), a credit note's is the refund's,
a supplier bill's is its printed footer. Nothing here works a figure out a second way, so
a register, the GSTR-1 tables and GSTR-3B cannot disagree. A bill's heading ("Tax Invoice"
or "Bill of Supply") plays no part: the lines decide the tax.

Nothing here writes to the store.
"""
from __future__ import annotations

import time
from collections import defaultdict
from decimal import ROUND_HALF_UP, Decimal
from typing import Callable, Optional

from core import gst_reports as gr
from core.bill_gst import split_tax
from core.customer_gst import state_of

ZERO = gr.ZERO
CENT = Decimal("0.01")
MARG_RATES = tuple(gr._rate(r) for r in (0, 5, 12, 18, 28))


def _f(v) -> float:
    return gr._f(v if isinstance(v, Decimal) else Decimal(str(v or 0)))


def _num(v) -> Decimal:
    try:
        return Decimal(str(v)) if v not in ("", None) else ZERO
    except Exception:
        return ZERO


def _pct(r) -> str:
    return f"{gr._rate_out(gr._rate(r))}%"


# ── a table on the screen / in a file ─────────────────────────────────────────────────────

def section(title: str, columns: list, rows: list, sums=(), nums=(), header=None, marg: bool = False,
            total_label: str = "Total") -> dict:
    """A titled table. ``sums`` are the columns added up in the totals row, ``nums`` the other
    numeric columns (right-aligned, never summed: a rate, a count of bills is summed)."""
    idx = {c: j for j, c in enumerate(columns)}
    s_idx = [idx[c] if isinstance(c, str) else c for c in sums]
    n_idx = sorted(set(s_idx) | {idx[c] if isinstance(c, str) else c for c in nums})
    total = None
    if s_idx and rows:
        total = [""] * len(columns)
        for j in s_idx:
            total[j] = _f(sum((_num(r[j]) for r in rows), ZERO))
        if not marg and 0 not in s_idx:
            total[0] = total_label
    return {"title": title, "columns": list(columns), "rows": rows, "total": total, "sum_cols": s_idx,
            "num_cols": n_idx, "header": header, "marg": marg}


# ── the period, read once per screen visit ───────────────────────────────────────────────

_CACHE: dict = {}
_CACHE_SECONDS = 45


def period(conn, start: str, end: str) -> gr.Period:
    """The period's bills, notes and supplier bills. Switching reports on the screen reads
    the store once: the same dates within a few seconds and no write since reuse it."""
    key = (id(conn), gr._iso(start), gr._iso(end))
    try:
        changes = conn.total_changes
    except Exception:
        changes = -1
    hit = _CACHE.get(key)
    if hit and hit[1] == changes and time.monotonic() - hit[0] < _CACHE_SECONDS:
        return hit[2]
    p = gr.load_period(conn, start, end)
    _CACHE.clear()
    _CACHE[key] = (time.monotonic(), changes, p)
    return p


class Ctx:
    """One report's inputs: the period, and the GSTR-1 tables worked out once."""

    def __init__(self, conn, p: gr.Period, comp_rate: float = 1.0):
        self.conn = conn
        self.p = p
        self.shop_state = p.shop.get("state", "")
        self.comp_rate = Decimal(str(comp_rate or 1))
        self._g1 = None
        self._3b = None

    @property
    def g1(self) -> dict:
        if self._g1 is None:
            self._g1 = gr.gstr1(self.p)
        return self._g1

    @property
    def t3b(self) -> dict:
        if self._3b is None:
            self._3b = gr.gstr3b(self.p, self.g1)
        return self._3b

    def legacy(self, title: str) -> dict:
        """A table exactly as the old GST window laid it out (gst_reports.sections)."""
        rep = {"gstr1": self.g1, "gstr3b": self.t3b, "sales_register": [], "purchase_register": [],
               "purchase_rates": [], "checks": [], "tally": []}
        return next(s for s in gr.sections(rep) if s["title"] == title)


# ── the Marg registers: one row per bill, a column block per GST rate ──────────────────

def _entry(when: str, party: str, address: str, kind: str, no: str, buckets: dict, invoice: Decimal,
           totals: Optional[tuple] = None) -> dict:
    """buckets: rate -> [taxable, igst, cgst, sgst]. totals: (taxable, igst, cgst, sgst) when the
    bill's figures are known only as a whole (its rate split is not)."""
    return {"date": when, "party": party, "address": address, "type": kind, "no": no, "buckets": buckets,
            "invoice": invoice, "totals": totals}


def _register(title: str, entries: list) -> dict:
    rates = sorted(set(MARG_RATES) | {r for e in entries for r in e["buckets"]})
    h1, h2, cols = ["Date", "Particulars", "Address", "Type", "Invoice"], ["", "", "", "", "No"], \
        ["Date", "Particulars", "Address", "Type", "Invoice No"]
    for r in rates:
        lab, half = _pct(r), _pct(r / 2)
        h1 += [lab, lab, half, half]
        h2 += ["Taxable", "IGST", "CGST", "SGST"]
        cols += [f"{lab} Taxable", f"{lab} IGST", f"{lab} CGST {half}", f"{lab} SGST {half}"]
    h1 += ["Taxable", "Integrated", "Central", "State", "Total", "Add", "Less", "Invoice"]
    h2 += ["Value", "Tax Amount", "Tax Amount", "Tax Amount", "Tax Amount", "", "", "Amount"]
    cols += ["Taxable Value", "Integrated Tax Amount", "Central Tax Amount", "State Tax Amount",
             "Total Tax Amount", "Add", "Less", "Invoice Amount"]
    rows = []
    for e in entries:
        row = [gr._dmy(e["date"]), e["party"], e["address"], e["type"], e["no"]]
        tot = [ZERO, ZERO, ZERO, ZERO]
        for r in rates:
            b = e["buckets"].get(r) or [ZERO, ZERO, ZERO, ZERO]
            row += [_f(x) for x in b]
            tot = [a + x for a, x in zip(tot, b)]
        if e["totals"] is not None:
            tot = list(e["totals"])
        tax = tot[1] + tot[2] + tot[3]
        diff = e["invoice"] - tot[0] - tax
        row += [_f(tot[0]), _f(tot[1]), _f(tot[2]), _f(tot[3]), _f(tax),
                _f(diff if diff > 0 else ZERO), _f(-diff if diff < 0 else ZERO), _f(e["invoice"])]
        rows.append(row)
    return section(title, cols, rows, sums=list(range(5, len(cols))), header=[h1, h2], marg=True)


def _sale_buckets(c: Ctx, b: gr.Bill) -> dict:
    g = gr.bill_tax(b)
    by: dict = defaultdict(lambda: [ZERO, ZERO])
    for lg in g.lines:
        by[lg.rate][0] += lg.taxable
        by[lg.rate][1] += lg.tax
    inter = gr._inter(b.pos, c.shop_state)
    return {r: [t, *gr._split(x, inter)] for r, (t, x) in by.items()}


def _note_buckets(c: Ctx, n: gr.Note) -> dict:
    g = gr.note_tax(n)
    by: dict = defaultdict(lambda: [ZERO, ZERO])
    for lg in g.lines:
        by[lg.rate][0] += lg.taxable
        by[lg.rate][1] += lg.tax
    inter = gr._inter(n.bill.pos, c.shop_state)
    return {r: [t, *gr._split(x, inter)] for r, (t, x) in by.items()}


def _supplier_inter(c: Ctx, gstin: str) -> bool:
    return gr._inter(state_of(gstin), c.shop_state)


def purchase_buckets(c: Ctx, pu: gr.Purchase) -> tuple:
    """(rate -> [taxable, igst, cgst, sgst], known). The supplier's own footer is the bill's
    figure: lines give the split by rate, a paisa or a rupee of difference from the footer goes
    to the biggest rate. A bill whose lines lack taxable values has a split only when all its
    lines share one rate (or it carries no tax at all); otherwise it is not guessed."""
    inter = _supplier_inter(c, pu.gstin)
    by: dict = defaultdict(lambda: [ZERO, ZERO])
    if pu.lines_known:
        for ln in pu.lines:
            by[ln.rate][0] += ln.taxable
            by[ln.rate][1] += ln.tax
        dt = pu.taxable - sum((v[0] for v in by.values()), ZERO)
        dx = pu.tax - sum((v[1] for v in by.values()), ZERO)
        if by and abs(dt) <= Decimal("1.00") and abs(dx) <= Decimal("1.00"):
            big = max(by, key=lambda r: by[r][0])
            by[big][0] += dt
            by[big][1] += dx
        known = True
    else:
        rates = {ln.rate for ln in pu.lines}
        if len(rates) == 1:
            by[next(iter(rates))] = [pu.taxable, pu.tax]
            known = True
        elif pu.tax == 0 and pu.taxable > 0:
            by[gr._rate(0)] = [pu.taxable, ZERO]
            known = True
        else:
            known = False
    return {r: [t, *gr._split(x, inter)] for r, (t, x) in by.items()}, known


def _purchase_totals(c: Ctx, pu: gr.Purchase) -> tuple:
    if _supplier_inter(c, pu.gstin):
        return pu.taxable, pu.tax, ZERO, ZERO
    return pu.taxable, ZERO, pu.cgst, pu.sgst


def _return_buckets(c: Ctx, r: gr.PurchaseReturn) -> dict:
    gstin = r.purchase.gstin if r.purchase else ""
    inter = _supplier_inter(c, gstin)
    by: dict = defaultdict(lambda: [ZERO, ZERO])
    for ln in r.lines:
        by[ln.rate][0] += ln.taxable
        by[ln.rate][1] += ln.tax
    return {k: [t, *gr._split(x, inter)] for k, (t, x) in by.items()}


def _kind(credit: bool) -> str:
    return "CREDIT" if credit else "CASH"


def sale_register(c: Ctx) -> dict:
    entries = [_entry(b.date, b.customer, b.address, _kind(b.credit), b.no, _sale_buckets(c, b), b.total)
               for b in c.p.bills if b.lines]
    notes = []
    skipped = [b for b in c.p.bills if not b.lines and b.total != 0]
    if skipped:
        notes.append(f"{len(skipped)} bill(s) have an amount but no medicine lines (Rs "
                     f"{_f(sum((b.total for b in skipped), ZERO))}): their GST rate is not known, so they are "
                     "not in this register -- see Checks.")
    return {"sections": [_register("Sale GST", entries)], "notes": notes}


def purchase_register_marg(c: Ctx) -> dict:
    entries, unsplit = [], 0
    for pu in c.p.purchases:
        buckets, known = purchase_buckets(c, pu)
        if not known:
            unsplit += 1
        entries.append(_entry(pu.date, pu.supplier, pu.address, "GST" if pu.gstin else "URD",
                              pu.bill_number or pu.no, buckets, pu.value,
                              None if known else _purchase_totals(c, pu)))
    notes = []
    if unsplit:
        notes.append(f"{unsplit} supplier bill(s) were saved without line-wise taxable values: their totals "
                     "are in the Taxable Value / Tax columns, the rate columns are empty for them.")
    return {"sections": [_register("Purchase GST", entries)], "notes": notes}


def sale_return_register(c: Ctx) -> dict:
    entries = [_entry(n.date, n.bill.customer, n.bill.address, _kind(n.bill.credit), n.no, _note_buckets(c, n),
                      n.refund) for n in c.p.notes if not n.problem and n.bill and n.lines]
    bad = [n for n in c.p.notes if n.problem or not n.bill or not n.lines]
    notes = []
    if bad:
        notes.append(f"{len(bad)} return(s) could not be matched to their bill lines and are not here -- see Checks.")
    return {"sections": [_register("Sale Return GST", entries)], "notes": notes}


def purchase_return_register(c: Ctx) -> dict:
    entries, unknown = [], 0
    for r in c.p.purchase_returns:
        pu = r.purchase
        buckets = _return_buckets(c, r)
        totals = None
        if not r.known:
            unknown += 1
            inter = _supplier_inter(c, pu.gstin if pu else "")
            totals = (r.taxable, *gr._split(r.tax, inter))
        entries.append(_entry(r.date, pu.supplier if pu else "", pu.address if pu else "",
                              "GST" if pu and pu.gstin else "URD", r.no, buckets, r.value, totals))
    notes = []
    if unknown:
        notes.append(f"{unknown} return(s) come from purchase lines without taxable values: only what is known "
                     "is shown (see Checks).")
    return {"sections": [_register("Purchase Return GST", entries)], "notes": notes}


# ── HSN ───────────────────────────────────────────────────────────────────────────────────

HSN_COLS = ["HSN", "Description", "UQC", "Total Quantity", "Total Value", "Rate", "Taxable Value",
            "Integrated Tax Amount", "Central Tax Amount", "State/UT Tax Amount", "Cess Amount"]
HSN_SUMS = ["Total Quantity", "Total Value", "Taxable Value", "Integrated Tax Amount", "Central Tax Amount",
            "State/UT Tax Amount", "Cess Amount"]


class _Hsn:
    def __init__(self):
        self.rows: dict = defaultdict(lambda: [ZERO, ZERO, ZERO, ZERO, ZERO, ""])   # qty, value, taxable, igst, tax, desc

    def add(self, hsn: str, rate: Decimal, kind: str, name: str, qty, value, taxable, tax, inter: bool, sign=1):
        row = self.rows[(hsn or "", rate, gr._uqc(kind))]
        row[0] += sign * Decimal(str(qty or 0))
        row[1] += sign * value
        row[2] += sign * taxable
        if inter:
            row[3] += sign * tax
        else:
            row[4] += sign * tax
        row[5] = row[5] or gr.HSN_DESC.get((hsn or "")[:4], "") or name

    def table(self, title: str) -> dict:
        out = []
        for (code, rate, uqc), r in sorted(self.rows.items(), key=lambda kv: (kv[0][0], kv[0][1], kv[0][2])):
            c, s = split_tax(r[4])
            out.append([code, r[5], uqc, float(r[0]), _f(r[1]), gr._rate_out(rate), _f(r[2]), _f(r[3]), _f(c), _f(s),
                        0.0])
        return section(title, HSN_COLS, out, sums=HSN_SUMS, nums=["Rate"])


def hsn_sales(c: Ctx, sales: bool, returns: bool, title: str) -> dict:
    """HSN of sales (H), of sale returns (J); the net (K) is the portal's Table 12 as filed."""
    h = _Hsn()
    if sales:
        for b in c.p.bills:
            g = gr.bill_tax(b)
            inter = gr._inter(b.pos, c.shop_state)
            for ln, lg in zip(b.lines, g.lines):
                h.add(ln.hsn, lg.rate, ln.kind, ln.name, ln.qty, lg.net, lg.taxable, lg.tax, inter)
    if returns:
        sign = -1 if sales else 1
        for n in c.p.notes:
            if n.problem or not n.bill or not n.lines:
                continue
            g = gr.note_tax(n)
            inter = gr._inter(n.bill.pos, c.shop_state)
            for (ln, qty, _), lg in zip(n.lines, g.lines):
                h.add(ln.hsn, lg.rate, ln.kind, ln.name, qty, lg.net, lg.taxable, lg.tax, inter, sign)
    return h.table(title)


def hsn_purchases(c: Ctx, purchases: bool, returns: bool, title: str) -> tuple:
    h = _Hsn()
    unknown = 0
    if purchases:
        for pu in c.p.purchases:
            if not pu.lines_known:
                unknown += 1
                continue
            inter = _supplier_inter(c, pu.gstin)
            for ln in pu.lines:
                h.add(ln.hsn, ln.rate, ln.kind, "", ln.qty, ln.taxable + ln.tax, ln.taxable, ln.tax, inter)
    if returns:
        sign = -1 if purchases else 1
        for r in c.p.purchase_returns:
            inter = _supplier_inter(c, r.purchase.gstin if r.purchase else "")
            if not r.known:
                unknown += 1
            for ln in r.lines:
                h.add(ln.hsn, ln.rate, ln.kind, "", ln.qty, ln.taxable + ln.tax, ln.taxable, ln.tax, inter, sign)
    return h.table(title), unknown


def _hsn_note(unknown: int) -> list:
    return ([f"{unknown} supplier bill(s) / return(s) have no line-wise taxable values and are not in this HSN "
             "table (their totals are in the registers)."] if unknown else [])


# ── GSTR-1 tables ─────────────────────────────────────────────────────────────────────────

def _nil_rows(c: Ctx) -> list:
    nil: dict = defaultdict(lambda: ZERO)
    for r in c.g1["b2cs"]:
        if _num(r["rate"]) == 0:
            nil["Inter-State supplies to unregistered persons" if r["supply"] == "INTER"
                else "Intra-State supplies to unregistered persons"] += _num(r["taxable"])
    order = ["Inter-State supplies to registered persons", "Intra-State supplies to registered persons",
             "Inter-State supplies to unregistered persons", "Intra-State supplies to unregistered persons"]
    return [[k, _f(nil.get(k, ZERO)), 0.0, 0.0] for k in order]


def nil_table(c: Ctx) -> dict:
    return section("nil (Table 8)", ["Description", "Nil Rated Supplies", "Exempted(other than nil rated/non GST supply)",
                                     "Non-GST supplies"], _nil_rows(c),
                   sums=["Nil Rated Supplies", "Exempted(other than nil rated/non GST supply)", "Non-GST supplies"])


def b2cs_report(c: Ctx) -> dict:
    sec = c.legacy("b2cs")
    rows = [r for r in sec["rows"] if _num(r[3]) != 0]
    cols = sec["columns"]
    return {"sections": [section("b2cs", cols, rows, sums=["Taxable Value", "Cess Amount", "IGST", "CGST", "SGST"],
                                 nums=["Rate"]), nil_table(c)],
            "notes": ["One row per place of supply and rate, net of returns on these bills.",
                      "0% sales go in Table 8 (nil rated), not in B2CS."]}


def _legacy_table(c: Ctx, title: str, sums: list, nums=("Rate",), rename: str = "") -> dict:
    sec = c.legacy(title)
    return section(rename or title, sec["columns"], sec["rows"], sums=sums, nums=list(nums))


def b2b_report(c: Ctx) -> dict:
    return {"sections": [_legacy_table(c, "b2b", ["Taxable Value", "Cess Amount", "IGST", "CGST", "SGST"],
                                       nums=["Rate", "Invoice Value"])],
            "notes": ["Bills to a buyer with a GSTIN on the bill date: wholesale tax invoices keep the GSTIN and place "
                      "of supply printed on them; IGST when the buyer is in another state.",
                      "One row per invoice and rate (as the GSTR-1 offline tool's b2b sheet)."]}


def b2cl_report(c: Ctx) -> dict:
    return {"sections": [_legacy_table(c, "b2cl", ["Taxable Value", "Cess Amount", "IGST"],
                                       nums=["Rate", "Invoice Value"])],
            "notes": ["Unregistered buyers in another state, invoice value above Rs 1,00,000."]}


def cdnr_report(c: Ctx) -> dict:
    return {"sections": [_legacy_table(c, "cdnr", ["Taxable Value", "Cess Amount", "IGST", "CGST", "SGST"],
                                       nums=["Rate", "Note Value"]),
                         _legacy_table(c, "cdnur", ["Taxable Value", "Cess Amount", "IGST"],
                                       nums=["Rate", "Note Value"])],
            "notes": ["cdnr: returns on bills to registered buyers. cdnur: returns on B2CL bills.",
                      "Returns on ordinary (B2C) bills are not credit notes in GSTR-1: they lower B2CS."]}


def docs_report(c: Ctx) -> dict:
    sec = c.legacy("docs")
    return {"sections": [section("docs (Table 13)", sec["columns"], sec["rows"], sums=["Total Number", "Cancelled"])],
            "notes": ["A missing number inside a series counts as cancelled (a bill deleted after it was made)."]}


def _g1_line(table: str, desc: str, rows: list, records: Optional[int] = None, value_key: str = "value") -> list:
    taxable = sum((_num(r.get("taxable")) for r in rows), ZERO)
    i = sum((_num(r.get("igst")) for r in rows), ZERO)
    cg = sum((_num(r.get("cgst")) for r in rows), ZERO)
    sg = sum((_num(r.get("sgst")) for r in rows), ZERO)
    return [table, desc, len(rows) if records is None else records, _f(taxable), _f(i), _f(cg), _f(sg), 0.0,
            _f(i + cg + sg)]


CLUB_COLS = ["Table", "Description", "Records", "Taxable Value", "Integrated Tax", "Central Tax", "State/UT Tax",
             "Cess", "Total Tax"]


def gstr1_club(c: Ctx) -> dict:
    g1 = c.g1
    taxed_b2cs = [r for r in g1["b2cs"] if _num(r["rate"]) != 0]
    rows = [
        _g1_line("4A", "B2B Invoices (registered buyers)", g1["b2b"], len({r["invoice_no"] for r in g1["b2b"]})),
        _g1_line("5A", "B2C (Large) Invoices", g1["b2cl"], len({r["invoice_no"] for r in g1["b2cl"]})),
        _g1_line("7", "B2C (Others)", taxed_b2cs),
        ["8", "Nil rated (B2C)", sum(1 for r in _nil_rows(c) if r[1]), _f(sum((_num(r[1]) for r in _nil_rows(c)), ZERO)),
         0.0, 0.0, 0.0, 0.0, 0.0],
        _g1_line("9B", "Credit Notes (Registered)", g1["cdnr"], len({r["note_no"] for r in g1["cdnr"]})),
        _g1_line("9B", "Credit Notes (Unregistered)", g1["cdnur"], len({r["note_no"] for r in g1["cdnur"]})),
        _g1_line("12", "HSN-wise summary (B2B)", g1["hsn_b2b"]),
        _g1_line("12", "HSN-wise summary (B2C)", g1["hsn_b2c"]),
        ["13", "Documents issued", sum(d["net_issued"] for d in g1["docs"]), "", "", "", "", "", ""],
    ]
    out = [_num(x) for x in rows[0][3:9]]
    for r in rows[1:3]:
        out = [a + _num(x) for a, x in zip(out, r[3:9])]
    for r in rows[4:6]:
        out = [a - _num(x) for a, x in zip(out, r[3:9])]
    rows.append(["", "Net taxable outward supplies (4A + 5A + 7 - 9B)", "", *[_f(x) for x in out]])
    secs = [section("GSTR-1 summary (club)", CLUB_COLS, rows, nums=CLUB_COLS[2:])]
    secs.append(_month_wise_sales(c))
    return {"sections": secs, "notes": ["Every GSTR-1 table of the period in one place, and month by month below. "
                                        "The HSN rows repeat the same sales by HSN: they are not added again."]}


def _months(c: Ctx) -> list:
    from core.gst_extras import _month_slices, _sub

    return [(a, b, _sub(c.p, a, b)) for a, b in _month_slices(c.p.start, c.p.end)]


def _month_wise_sales(c: Ctx) -> dict:
    cols = ["Month", "Bills", "Credit Notes", "Sales Value", "Taxable Value", "Integrated Tax", "Central Tax",
            "State/UT Tax", "Nil Rated", "Total Tax"]
    rows = []
    for a, _, sub in _months(c):
        t = gr.gstr3b(sub)
        o = t["3.1(a) Outward taxable supplies"]
        rows.append([a[:7], len(sub.bills), len(sub.notes), _f(sum((b.total for b in sub.bills), ZERO)),
                     o["taxable"], o["igst"], o["cgst"], o["sgst"], t["3.1(c) Nil rated / exempted"]["taxable"],
                     _f(_num(o["igst"]) + _num(o["cgst"]) + _num(o["sgst"]))])
    return section("Month-wise (net of credit notes)", cols, rows, sums=cols[1:])


# ── purchase side (GSTR-2) ────────────────────────────────────────────────────────────────

PB2B_COLS = ["GSTIN of Supplier", "Supplier Name", "Invoice Number", "Invoice date", "Invoice Value", "Place Of Supply",
             "Reverse Charge", "Invoice Type", "Rate", "Taxable Value", "Integrated Tax Paid", "Central Tax Paid",
             "State/UT Tax Paid", "Cess Paid", "Eligibility For ITC", "Availed ITC Integrated Tax",
             "Availed ITC Central Tax", "Availed ITC State/UT Tax", "Availed ITC Cess"]


def purchase_b2b(c: Ctx) -> dict:
    rows, unsplit = [], 0
    pos = gr._pos(c.shop_state)
    for pu in c.p.purchases:
        if not pu.gstin:
            continue
        buckets, known = purchase_buckets(c, pu)
        parts = sorted(buckets.items()) if known else [("", [pu.taxable, *_purchase_totals(c, pu)[1:]])]
        if not known:
            unsplit += 1
        for rate, (t, i, cg, sg) in parts:
            rows.append([pu.gstin, pu.supplier, pu.bill_number or pu.no, gr._dmy(pu.date), _f(pu.value), pos, "N",
                         "Regular", "" if rate == "" else gr._rate_out(rate), _f(t), _f(i), _f(cg), _f(sg), 0.0,
                         "Inputs", _f(i), _f(cg), _f(sg), 0.0])
    sums = PB2B_COLS[9:14] + PB2B_COLS[15:]
    notes = ["Supplier bills with a GSTIN (input tax credit), one row per bill and rate. Unregistered suppliers' "
             "bills are in D / U, with no ITC."]
    if unsplit:
        notes.append(f"{unsplit} bill(s) have no rate-wise split (lines saved without taxable values): one row "
                     "with the bill's totals and the Rate empty.")
    return {"sections": [section("b2b (GSTR-2)", PB2B_COLS, rows, sums=sums, nums=["Rate", "Invoice Value"])],
            "notes": notes}


DBNR_COLS = ["GSTIN of Supplier", "Supplier Name", "Note Number", "Note Date", "Invoice Number", "Invoice Date",
             "Document Type", "Supply Type", "Note Value", "Rate", "Taxable Value", "Integrated Tax Paid",
             "Central Tax Paid", "State/UT Tax Paid", "Cess Paid"]


def dbnr(c: Ctx) -> dict:
    rows, unknown = [], 0
    for r in c.p.purchase_returns:
        pu = r.purchase
        if not pu or not pu.gstin:
            continue
        inter = _supplier_inter(c, pu.gstin)
        buckets = sorted(_return_buckets(c, r).items())
        if not r.known:
            unknown += 1
            buckets = buckets or [("", [r.taxable, *gr._split(r.tax, inter)])]
        for rate, (t, i, cg, sg) in buckets:
            rows.append([pu.gstin, pu.supplier, r.no, gr._dmy(r.date), pu.bill_number or pu.no, gr._dmy(pu.date), "D",
                         "Inter State" if inter else "Intra State", _f(r.value),
                         "" if rate == "" else gr._rate_out(rate), _f(t), _f(i), _f(cg), _f(sg), 0.0])
    notes = ["Goods returned to registered suppliers: the debit notes that lower ITC (GSTR-3B 4(B)(2))."]
    if unknown:
        notes.append(f"{unknown} return(s) come from purchase lines without taxable values (see Checks).")
    return {"sections": [section("dbnr (GSTR-2)", DBNR_COLS, rows, sums=DBNR_COLS[10:], nums=["Rate", "Note Value"])],
            "notes": notes}


def purchase_party(c: Ctx) -> dict:
    cols = ["Supplier", "GSTIN", "Bills", "Taxable Value", "Integrated Tax", "Central Tax", "State/UT Tax",
            "Total Tax", "Bill Value", "Returns Taxable", "Returns Tax", "Net Taxable", "Net Tax (ITC)"]
    by: dict = {}

    def row(name, gstin):
        return by.setdefault((name, gstin), [0, ZERO, ZERO, ZERO, ZERO, ZERO, ZERO, ZERO])

    for pu in c.p.purchases:
        t, i, cg, sg = _purchase_totals(c, pu)
        r = row(pu.supplier, pu.gstin)
        r[0] += 1
        for k, v in enumerate((t, i, cg, sg, pu.value), start=1):
            r[k] += v
    for pr in c.p.purchase_returns:
        if not pr.purchase:
            continue
        r = row(pr.purchase.supplier, pr.purchase.gstin)
        r[6] += pr.taxable
        r[7] += pr.tax
    rows = []
    for (name, gstin), r in sorted(by.items(), key=lambda kv: kv[0][0].upper()):
        tax = r[2] + r[3] + r[4]
        rows.append([name, gstin, r[0], _f(r[1]), _f(r[2]), _f(r[3]), _f(r[4]), _f(tax), _f(r[5]), _f(r[6]), _f(r[7]),
                     _f(r[1] - r[6]), _f(tax - r[7])])
    return {"sections": [section("Purchase party-wise", cols, rows, sums=cols[2:])],
            "notes": ["Each supplier's bills and the goods sent back to them in the period."]}


def gstr2_club(c: Ctx) -> dict:
    reg = [pu for pu in c.p.purchases if pu.gstin]
    urd = [pu for pu in c.p.purchases if not pu.gstin]

    def line(table, desc, bills, returns=()):
        t = i = cg = sg = ZERO
        for pu in bills:
            a, b, d, e = _purchase_totals(c, pu)
            t, i, cg, sg = t + a, i + b, cg + d, sg + e
        for r in returns:
            inter = _supplier_inter(c, r.purchase.gstin)
            x = gr._split(r.tax, inter)
            t, i, cg, sg = t + r.taxable, i + x[0], cg + x[1], sg + x[2]
        return [table, desc, len(bills) or len(returns), _f(t), _f(i), _f(cg), _f(sg), 0.0, _f(i + cg + sg)]

    rets = [r for r in c.p.purchase_returns if r.purchase and r.purchase.gstin]
    t3b = c.t3b
    net = t3b["4(C) Net ITC available"]
    exempt = _exempt_inward(c)
    rows = [line("B2B", "Bills from registered suppliers (ITC)", reg),
            line("B2BUR", "Bills from unregistered suppliers (no ITC)", urd),
            line("CDN", "Goods returned to registered suppliers (debit notes)", [], rets),
            ["NIL", "Nil rated / exempt inward (0% lines)", "", _f(exempt[0] + exempt[1]), "", "", "", "", ""],
            ["ITC", "Net ITC (B2B - CDN)", "", "", net["igst"], net["cgst"], net["sgst"], 0.0,
             _f(_num(net["igst"]) + _num(net["cgst"]) + _num(net["sgst"]))]]
    month_cols = ["Month", "Bills", "Bill Value", "Taxable Value", "Integrated Tax", "Central Tax", "State/UT Tax",
                  "Returns Tax", "Net ITC"]
    mrows = []
    for a, _, sub in _months(c):
        t3 = gr.gstr3b(sub)
        itc, rev = t3["4(A)(5) All other ITC"], t3["4(B)(2) Reversed: goods returned to suppliers"]
        n = t3["4(C) Net ITC available"]
        mrows.append([a[:7], len(sub.purchases), _f(sum((pu.value for pu in sub.purchases), ZERO)),
                      _f(sum((pu.taxable for pu in sub.purchases), ZERO)), itc["igst"], itc["cgst"], itc["sgst"],
                      _f(sum((_num(rev[k]) for k in ("igst", "cgst", "sgst")), ZERO)),
                      _f(sum((_num(n[k]) for k in ("igst", "cgst", "sgst")), ZERO))])
    return {"sections": [section("GSTR-2 summary (club)", CLUB_COLS, rows, nums=CLUB_COLS[2:]),
                         section("Month-wise purchases", month_cols, mrows, sums=month_cols[1:])],
            "notes": ["Check the ITC against GSTR-2B on the portal before filing GSTR-3B."]}


def _exempt_inward(c: Ctx) -> tuple:
    """(inter-state, intra-state) value of 0% supplier lines."""
    inter = intra = ZERO
    for pu in c.p.purchases:
        if not pu.lines_known:
            continue
        v = sum((ln.taxable for ln in pu.lines if ln.rate == 0), ZERO)
        if _supplier_inter(c, pu.gstin):
            inter += v
        else:
            intra += v
    return inter, intra


# ── GSTR-3B ───────────────────────────────────────────────────────────────────────────────

def gstr3b_report(c: Ctx) -> dict:
    from core import gst_extras

    t = c.t3b
    o, nil = t["3.1(a) Outward taxable supplies"], t["3.1(c) Nil rated / exempted"]
    cols31 = ["Nature of Supplies", "Total Taxable Value", "Integrated Tax", "Central Tax", "State/UT Tax", "Cess"]
    s31 = [["(a) Outward taxable supplies (other than zero rated, nil rated and exempted)", o["taxable"], o["igst"],
            o["cgst"], o["sgst"], 0.0],
           ["(b) Outward taxable supplies (zero rated)", 0.0, 0.0, "", "", 0.0],
           ["(c) Other outward supplies (Nil rated, exempted)", nil["taxable"], "", "", "", ""],
           ["(d) Inward supplies (liable to reverse charge)", 0.0, 0.0, 0.0, 0.0, 0.0],
           ["(e) Non-GST outward supplies", 0.0, "", "", "", ""]]
    s32 = [[gr._pos(r["pos"]), r["taxable"], r["igst"]] for r in t["3.2 Inter-state to unregistered"]]
    itc, rev, net = (t["4(A)(5) All other ITC"], t["4(B)(2) Reversed: goods returned to suppliers"],
                     t["4(C) Net ITC available"])
    z = [0.0, 0.0, 0.0, 0.0]
    s4 = [["(A) ITC Available (whether in full or part)", "", "", "", ""],
          ["   (1) Import of goods", *z], ["   (2) Import of services", *z],
          ["   (3) Inward supplies liable to reverse charge (other than 1 & 2 above)", *z],
          ["   (4) Inward supplies from ISD", *z],
          ["   (5) All other ITC", itc["igst"], itc["cgst"], itc["sgst"], 0.0],
          ["(B) ITC Reversed", "", "", "", ""],
          ["   (1) As per rules 38, 42 & 43 of CGST Rules and section 17(5)", *z],
          ["   (2) Others (goods returned to suppliers)", rev["igst"], rev["cgst"], rev["sgst"], 0.0],
          ["(C) Net ITC Available (A) - (B)", net["igst"], net["cgst"], net["sgst"], 0.0],
          ["(D) Other Details", "", "", "", ""],
          ["   (1) ITC reclaimed which was reversed under Table 4(B)(2) in earlier tax period", *z],
          ["   (2) Ineligible ITC under section 16(4) & ITC restricted due to PoS rules", *z]]
    ex_inter, ex_intra = _exempt_inward(c)
    s5 = [["From a supplier under composition scheme, Exempt and Nil rated supply", _f(ex_inter), _f(ex_intra)],
          ["Non GST supply", 0.0, 0.0]]
    # 6.1: the tax due against this period's ITC plus what was left from earlier months of the year
    liab = {h: _num(o[h]) for h in ("igst", "cgst", "sgst")}
    credit = {h: _num(net[h]) for h in ("igst", "cgst", "sgst")}
    brought = {h: ZERO for h in credit}
    note61 = "Credit used in the order the law fixes (IGST first; CGST never pays SGST)."
    try:
        months = gst_extras.monthly_ledger(c.conn, c.p.end)
        first = next((m for m in months if m["from"] == c.p.start), None)
        if first is not None and months:
            idx = months.index(first)
            if idx > 0:
                brought = {h: _num(months[idx - 1]["carried"][h]) for h in credit}
                note61 += " Credit brought forward from the earlier months of the financial year is added."
    except Exception:
        pass
    so = gst_extras.set_off(liab, {h: credit[h] + brought[h] for h in credit})
    s61 = []
    for h, name in (("igst", "Integrated Tax"), ("cgst", "Central Tax"), ("sgst", "State/UT Tax")):
        cash = _num(so["cash"][h])
        s61.append([name, _f(liab[h]), _f(liab[h] - cash), _f(cash), _f(brought[h]), so["carried"][h]])
    s61.append(["Cess", 0.0, 0.0, 0.0, 0.0, 0.0])
    return {"sections": [
        section("3.1 Details of Outward Supplies and inward supplies liable to reverse charge", cols31, s31,
                nums=cols31[1:]),
        section("3.2 Inter-State supplies made to unregistered persons",
                ["Place of Supply (State/UT)", "Total Taxable value", "Amount of Integrated Tax"], s32,
                sums=["Total Taxable value", "Amount of Integrated Tax"]),
        section("4. Eligible ITC", ["Details", "Integrated Tax", "Central Tax", "State/UT Tax", "Cess"], s4,
                nums=["Integrated Tax", "Central Tax", "State/UT Tax", "Cess"]),
        section("5. Values of exempt, nil-rated and non-GST inward supplies",
                ["Nature of supplies", "Inter-State supplies", "Intra-State supplies"], s5,
                sums=["Inter-State supplies", "Intra-State supplies"]),
        section("6.1 Payment of tax", ["Description", "Tax payable", "Paid through ITC", "Tax paid in cash",
                                       "ITC brought forward", "ITC carried forward"], s61,
                sums=["Tax payable", "Paid through ITC", "Tax paid in cash", "ITC brought forward",
                      "ITC carried forward"]),
    ], "notes": [
        "3.1(a) is the sales tax net of credit notes; 4(A)(5) the GST on registered suppliers' bills.",
        note61,
        "Interest and late fee (5.1) are not worked out here.",
    ]}


# ── composition ───────────────────────────────────────────────────────────────────────────

def _comp_lines(c: Ctx, sub: gr.Period) -> tuple:
    """(sales value, of it 0%, returns value, of it 0%) of a period, line by line after the
    bill discount: the turnover a composition dealer pays on."""
    sales = nil_s = rets = nil_r = ZERO
    for b in sub.bills:
        for lg in gr.bill_tax(b).lines:
            sales += lg.net
            if lg.rate == 0:
                nil_s += lg.net
    for n in sub.notes:
        if n.problem or not n.bill or not n.lines:
            continue
        for lg in gr.note_tax(n).lines:
            rets += lg.net
            if lg.rate == 0:
                nil_r += lg.net
    return sales, nil_s, rets, nil_r


def _comp_tax(c: Ctx, turnover: Decimal) -> tuple:
    tax = (turnover * c.comp_rate / 100).quantize(CENT, rounding=ROUND_HALF_UP)
    return split_tax(tax)


COMP_COLS = ["Period", "Bills", "Sales Value", "of which Nil / Exempt", "Sale Returns", "Returns Nil / Exempt",
             "Net Taxable Turnover", "Rate", "Central Tax", "State/UT Tax", "Total Tax"]


def _comp_row(c: Ctx, label: str, sub: gr.Period) -> list:
    s, ns, r, nr = _comp_lines(c, sub)
    turnover = (s - ns) - (r - nr)
    cg, sg = _comp_tax(c, turnover)
    return [label, len(sub.bills), _f(s), _f(ns), _f(r), _f(nr), _f(turnover), gr._rate_out(gr._rate(c.comp_rate)),
            _f(cg), _f(sg), _f(cg + sg)]


def _quarters(c: Ctx) -> list:
    from core.gst_extras import _sub

    out: dict = {}
    for a, b, _ in _months(c):
        m = int(a[5:7])
        q = {4: 1, 5: 1, 6: 1, 7: 2, 8: 2, 9: 2, 10: 3, 11: 3, 12: 3, 1: 4, 2: 4, 3: 4}[m]
        fy = int(a[:4]) if m >= 4 else int(a[:4]) - 1
        key = f"FY {fy}-{str(fy + 1)[2:]} Q{q}"
        lo, hi = out.get(key, (a, b))
        out[key] = (min(lo, a), max(hi, b))
    return [(k, _sub(c.p, a, b)) for k, (a, b) in out.items()]


def composition_sale(c: Ctx) -> dict:
    months = [_comp_row(c, a[:7], sub) for a, _, sub in _months(c)]
    quarters = [_comp_row(c, k, sub) for k, sub in _quarters(c)]
    sums = COMP_COLS[1:7] + COMP_COLS[8:]
    return {"sections": [section("Month-wise (Sale - Sale Return)", COMP_COLS, months, sums=sums, nums=["Rate"]),
                         section("Quarter-wise (CMP-08)", COMP_COLS, quarters, sums=sums, nums=["Rate"])],
            "notes": [f"Composition tax at {gr._rate_out(gr._rate(c.comp_rate))}% of the turnover of taxable goods "
                      "(half Central, half State); 0% (nil / exempt) goods are left out, as for a trader.",
                      "Turnover is each bill line after the bill discount, less what was refunded on returns.",
                      "A composition dealer issues a Bill of Supply and collects no GST from the buyer."]}


def _inward_by_supplier(c: Ctx) -> tuple:
    reg: dict = {}
    urd: dict = {}
    for pu in c.p.purchases:
        t, i, cg, sg = _purchase_totals(c, pu)
        if pu.gstin:
            r = reg.setdefault((pu.gstin, pu.supplier), [0, ZERO, ZERO, ZERO, ZERO, ZERO])
            r[0] += 1
            for k, v in enumerate((t, i, cg, sg, pu.value), start=1):
                r[k] += v
        else:
            r = urd.setdefault(pu.supplier, [0, ZERO, ZERO])
            r[0] += 1
            r[1] += t
            r[2] += pu.value
    reg_rows = [[g, n, r[0], _f(r[1]), _f(r[2]), _f(r[3]), _f(r[4]), 0.0, _f(r[5])]
                for (g, n), r in sorted(reg.items(), key=lambda kv: kv[0][1].upper())]
    urd_rows = [[n, r[0], _f(r[1]), _f(r[2])] for n, r in sorted(urd.items(), key=lambda kv: kv[0].upper())]
    return reg_rows, urd_rows


REG_COLS = ["GSTIN of Supplier", "Supplier Name", "Bills", "Taxable Value", "Integrated Tax", "Central Tax",
            "State/UT Tax", "Cess", "Bill Value"]
URD_COLS = ["Supplier Name", "Bills", "Taxable Value", "Bill Value"]


def composition_purchase(c: Ctx) -> dict:
    reg, urd = _inward_by_supplier(c)
    return {"sections": [section("4A Inward supplies from registered suppliers", REG_COLS, reg, sums=REG_COLS[2:]),
                         section("4C Inward supplies from unregistered suppliers", URD_COLS, urd, sums=URD_COLS[1:])],
            "notes": ["A composition dealer takes no input tax credit: the GST on these bills is a cost.",
                      "4B (reverse charge from registered) and 4D (import of service) are not kept by the software: "
                      "0 unless entered on the portal."]}


def gstr4_annual(c: Ctx) -> dict:
    reg, urd = _inward_by_supplier(c)
    quarters = [_comp_row(c, k, sub) for k, sub in _quarters(c)]
    q_cols = ["Quarter", "Net Taxable Turnover", "Central Tax", "State/UT Tax", "Total Tax"]
    q_rows = [[r[0], r[6], r[8], r[9], r[10]] for r in quarters]
    whole = _comp_row(c, "Year", c.p)
    rate_rows = [[whole[7], whole[6], whole[8], whole[9], whole[10]]]
    return {"sections": [
        section("Table 4A: Inward supplies from registered suppliers", REG_COLS, reg, sums=REG_COLS[2:]),
        section("Table 4C: Inward supplies from unregistered suppliers", URD_COLS, urd, sums=URD_COLS[1:]),
        section("Table 5: Summary of self-assessed liability as per CMP-08", q_cols, q_rows, sums=q_cols[1:]),
        section("Table 6: Tax rate-wise details of outward supplies", ["Rate of tax", "Turnover", "Central Tax",
                                                                        "State/UT Tax", "Total Tax"], rate_rows,
                nums=["Rate of tax", "Turnover", "Central Tax", "State/UT Tax", "Total Tax"]),
    ], "notes": ["GSTR-4 is filed once a year (by 30 April) by a composition dealer; pick the financial year above.",
                 "Tables 4B, 4D (reverse charge, import of service), 7 (TDS/TCS) and 8 (interest / late fee) are "
                 "not kept by the software: fill them on the portal if any."]}


# ── annual return (GSTR-9) ────────────────────────────────────────────────────────────────

def gstr9(c: Ctx) -> dict:
    from core import gst_extras

    g1, t = c.g1, c.t3b

    def sums(rows, rate_filter=lambda r: _num(r["rate"]) != 0):
        rows = [r for r in rows if rate_filter(r)]
        return [sum((_num(r.get(k)) for r in rows), ZERO) for k in ("taxable", "cgst", "sgst", "igst")]

    b2c = [a + b for a, b in zip(sums(g1["b2cs"]), sums(g1["b2cl"]))]
    b2b = sums(g1["b2b"])
    cn = [a + b for a, b in zip(sums(g1["cdnr"]), sums(g1["cdnur"]))]
    zero4 = [ZERO] * 4
    h = [a + b for a, b in zip(b2c, b2b)]
    n = [a - b for a, b in zip(h, cn)]

    def r4(label, v):
        return [label, *[_f(x) for x in v], 0.0]

    t4 = [r4("4A Supplies made to un-registered persons (B2C)", b2c),
          r4("4B Supplies made to registered persons (B2B)", b2b),
          r4("4C Zero rated supply (Export) on payment of tax", zero4),
          r4("4D Supply to SEZs on payment of tax", zero4),
          r4("4E Deemed Exports", zero4),
          r4("4F Advances on which tax has been paid but invoice has not been issued", zero4),
          r4("4G Inward supplies on which tax is to be paid on reverse charge basis", zero4),
          r4("4H Sub-total (A to G above)", h),
          r4("4I Credit Notes issued in respect of transactions specified in (B) to (E) above (-)", cn),
          r4("4J Debit Notes issued (+)", zero4),
          r4("4K Supplies / tax declared through Amendments (+)", zero4),
          r4("4L Supplies / tax reduced through Amendments (-)", zero4),
          r4("4M Sub-total (I to L above)", [-x for x in cn]),
          r4("4N Supplies and advances on which tax is to be paid (H + M) above", n)]
    nil = _num(t["3.1(c) Nil rated / exempted"]["taxable"])
    t5 = [["5D Exempted", 0.0], ["5E Nil Rated", _f(nil)], ["5F Non-GST supply", 0.0],
          ["5G Sub-total (A to F above)", _f(nil)], ["5N Supplies and advances on which tax is not to be paid", _f(nil)],
          ["5O Total Turnover (including advances) (4N + 5N - 4G)", _f(n[0] + nil)]]
    itc, rev, net = (t["4(A)(5) All other ITC"], t["4(B)(2) Reversed: goods returned to suppliers"],
                     t["4(C) Net ITC available"])
    t6 = [["6B Inward supplies (other than imports and inward supplies liable to reverse charge) - Inputs",
           itc["cgst"], itc["sgst"], itc["igst"], 0.0],
          ["6O Total ITC availed", itc["cgst"], itc["sgst"], itc["igst"], 0.0],
          ["7H Other reversals (goods returned to suppliers)", rev["cgst"], rev["sgst"], rev["igst"], 0.0],
          ["7J Net Total ITC Available for Utilization (6O - 7I)", net["cgst"], net["sgst"], net["igst"], 0.0]]
    payable = {"cgst": n[1], "sgst": n[2], "igst": n[3]}
    cash = {h_: ZERO for h_ in payable}
    try:
        for m in gst_extras.monthly_ledger(c.conn, c.p.end):
            if c.p.start <= m["from"] <= c.p.end:
                for h_ in cash:
                    cash[h_] += _num(m["cash"][h_])
    except Exception:
        pass
    t9 = [[name, _f(payable[k]), _f(cash[k]), _f(payable[k] - cash[k])]
          for k, name in (("igst", "Integrated Tax"), ("cgst", "Central Tax"), ("sgst", "State/UT Tax"))]
    t9.append(["Cess", 0.0, 0.0, 0.0])
    hsn_out = hsn_sales(c, True, True, "Table 17: HSN wise summary of outward supplies")
    hsn_in, unknown = hsn_purchases(c, True, True, "Table 18: HSN wise summary of inward supplies")
    cols4 = ["Nature of Supplies", "Taxable Value", "Central Tax", "State Tax / UT Tax", "Integrated Tax", "Cess"]
    cols6 = ["Description", "Central Tax", "State Tax / UT Tax", "Integrated Tax", "Cess"]
    return {"sections": [
        section("Pt. II Table 4: Outward supplies on which tax is payable", cols4, t4, nums=cols4[1:]),
        section("Table 5: Outward supplies on which tax is not payable", ["Nature of Supplies", "Taxable Value"], t5,
                nums=["Taxable Value"]),
        section("Pt. III Tables 6 & 7: ITC availed and reversed", cols6, t6, nums=cols6[1:]),
        section("Pt. IV Table 9: Tax paid", ["Description", "Tax Payable", "Paid through cash",
                                            "Paid through ITC"], t9,
                sums=["Tax Payable", "Paid through cash", "Paid through ITC"]),
        hsn_out, hsn_in,
    ], "notes": ["Pick the financial year above (FY). Filled from the shop's own bills for the year.",
                 "Table 8 (ITC as per GSTR-2B), Tables 10-14 (amendments of the next year, differential tax) and "
                 "15-16 (refunds, demands, job work) need the portal / the CA.",
                 "Table 9 'paid through cash' is the month-by-month set-off of the year's own figures."]
                + _hsn_note(unknown)}


# ── A: the month's upload ─────────────────────────────────────────────────────────────────

def upload_process(c: Ctx) -> dict:
    g1 = c.g1
    nil = sum((_num(r[1]) for r in _nil_rows(c)), ZERO)
    hsn_bad = sum(1 for t in ("hsn_b2b", "hsn_b2c") for r in g1[t] if r["problem"])

    def line(table, rows, status, records=None):
        x = _g1_line(table, "", rows, records)
        return [table, x[2], x[3], x[4], x[5], x[6], status]

    cols = ["GSTR-1 table", "Records", "Taxable Value", "Integrated Tax", "Central Tax", "State/UT Tax", "In the JSON"]
    rows = [line("B2B (4A)", g1["b2b"], "Yes", len({r["invoice_no"] for r in g1["b2b"]})),
            line("B2CL (5A)", g1["b2cl"], "Yes", len({r["invoice_no"] for r in g1["b2cl"]})),
            line("B2CS (7)", [r for r in g1["b2cs"] if _num(r["rate"]) != 0], "Yes"),
            ["Nil rated (8)", sum(1 for r in _nil_rows(c) if r[1]), _f(nil), 0.0, 0.0, 0.0, "Yes"],
            line("CDNR (9B)", g1["cdnr"], "Yes", len({r["note_no"] for r in g1["cdnr"]})),
            line("CDNUR (9B)", g1["cdnur"], "Yes", len({r["note_no"] for r in g1["cdnur"]})),
            line("HSN B2B (12)", g1["hsn_b2b"], "Yes" + (f" (rows without HSN left out: {hsn_bad})" if hsn_bad else "")),
            line("HSN B2C (12)", g1["hsn_b2c"], "Yes"),
            ["Documents (13)", sum(d["net_issued"] for d in g1["docs"]), "", "", "", "", "Yes"],
            ["Exports (6A), SEZ (6B)", 0, "", "", "", "", "Not kept by the software"],
            ["Advances (11A / 11B)", 0, "", "", "", "", "Not kept by the software"],
            ["Amendments (9A, 9C, 10)", 0, "", "", "", "", "Enter on the portal if any"],
            ["E-commerce (14 / 15)", 0, "", "", "", "", "Not applicable"]]
    checks = len(c.p.checks)
    steps = [["1", "Look at Checks", f"{checks} thing(s) to look at" if checks else "Nothing to check"],
             ["2", "Fill missing HSN codes (Tools > HSN bhara)", "Done" if not hsn_bad else f"{hsn_bad} HSN row(s)"],
             ["3", "Download the GSTR-1 JSON (button above) and upload it on the portal: Returns > GSTR-1 > "
                   "Prepare Offline > Upload", ""],
             ["4", "Give the CA the Sale register (C) and Purchase register (D) Excel", ""],
             ["5", "File GSTR-1 by the 11th (monthly) / IFF by the 13th (QRMP)", ""],
             ["6", "File GSTR-3B (B) by the 20th (monthly) / 22nd-24th (QRMP), after checking ITC with GSTR-2B", ""],
             ["7", "Save the filed figures (Tools > Filed / julvni) to see later changes", ""]]
    return {"sections": [section("GSTR-1 tables for the period", cols, rows, nums=cols[1:6]),
                         section("Steps", ["Step", "What", "Status"], steps)],
            "notes": ["The JSON holds B2B, B2CL, B2CS, nil, CDNR, CDNUR, HSN and documents, made from the same bills "
                      "as these tables. Open it in the GST offline tool or upload it on the portal."]}


# ── the list the screen shows ─────────────────────────────────────────────────────────────

Builder = Callable[[Ctx], dict]

CATALOGUE: list = [
    # key, letter, group, title, form, builder
    ("upload", "A", "Month", "GST upload process for the month", "GSTR-1 JSON + registers", upload_process),
    ("gstr3b", "B", "GSTR-3B", "GSTR-3B", "GSTR-3B", gstr3b_report),
    ("b2b", "", "GSTR-1", "Sale (GSTR-1) b2b - registered buyers", "GSTR-1 Table 4A", b2b_report),
    ("sale", "C", "GSTR-1", "Sale (GSTR-1) Detail", "Sale register (GST upload)", sale_register),
    ("sale_return", "E", "GSTR-1", "Sale Return (GSTR-1) Detail", "Sale return register", sale_return_register),
    ("b2cs", "G", "GSTR-1", "Sale (GSTR-1) b2cs", "GSTR-1 Table 7 + 8", b2cs_report),
    ("b2cl", "", "GSTR-1", "Sale (GSTR-1) b2cl", "GSTR-1 Table 5A", b2cl_report),
    ("hsn_sale", "H", "GSTR-1", "HSN Sale (GSTR-1) hsn", "HSN of sales", None),
    ("cdnr", "I", "GSTR-1", "Sale Return (GSTR-1) cdnr", "GSTR-1 Table 9B", cdnr_report),
    ("hsn_sale_return", "J", "GSTR-1", "HSN Sale Return (GSTR-1) hsn", "HSN of sale returns", None),
    ("hsn_sale_net", "K", "GSTR-1", "HSN Sale - Sale Return (GSTR-1) hsn", "GSTR-1 Table 12", None),
    ("docs", "", "GSTR-1", "Documents issued (GSTR-1) docs", "GSTR-1 Table 13", docs_report),
    ("gstr1_club", "T", "GSTR-1", "Sale GSTR-1 (club)", "GSTR-1 summary", gstr1_club),
    ("purchase", "D", "GSTR-2", "Purchase (GSTR-2) Detail", "Purchase register (GST upload)", purchase_register_marg),
    ("purchase_return", "F", "GSTR-2", "Purchase Return (GSTR-2) Detail", "Purchase return register",
     purchase_return_register),
    ("purchase_b2b", "L", "GSTR-2", "Purchase (GSTR-2) b2b", "GSTR-2 b2b", purchase_b2b),
    ("hsn_purchase", "M", "GSTR-2", "HSN Purchase (GSTR-2) hsn", "HSN of purchases", None),
    ("dbnr", "N", "GSTR-2", "Purchase Return (GSTR-2) dbnr", "Debit notes", dbnr),
    ("hsn_purchase_return", "O", "GSTR-2", "HSN Purchase Rtn (GSTR-2) hsn", "HSN of purchase returns", None),
    ("hsn_purchase_net", "P", "GSTR-2", "HSN Purchase - Purchase Return (GSTR-2) hsn", "HSN of inward supplies", None),
    ("purchase_party", "Q", "GSTR-2", "Purchase (GSTR-2) Party Wise", "Supplier-wise", purchase_party),
    ("gstr2_club", "U", "GSTR-2", "Purchase GSTR-2 (club)", "GSTR-2 summary", gstr2_club),
    ("composition_sale", "R", "Composition", "Sale - Sale Return Composition GST", "CMP-08", composition_sale),
    ("composition_purchase", "S", "Composition", "Purchase Composition GST", "GSTR-4 Table 4", composition_purchase),
    ("gstr9", "V", "Annual", "Annual GSTR (GSTR-9)", "GSTR-9", gstr9),
    ("gstr4", "W", "Annual", "Annual GSTR4", "GSTR-4", gstr4_annual),
    ("julvni", "", "Tools", "Bills vs GST (julvni)", "Bills total = taxable + tax", None),
    ("checks", "", "Tools", "Checks", "Bills to look at", None),
]


def julvni(c: Ctx) -> dict:
    return {"sections": [section("Bills vs GST", ["Hisob", "Rakkam"], [[r["line"], r["amount"]] for r in gr.tally(c.p)],
                                 nums=["Rakkam"])],
            "notes": ["The bills' own totals against the GST tables: PHARAK must be 0."]}


def checks(c: Ctx) -> dict:
    return {"sections": [section("Checks", ["What", "Bill / Record", "Detail"], [list(x) for x in c.p.checks])],
            "notes": ["Nothing here is guessed: a bill named here is shown as printed, or left out and said so."]}

# Reports filled only partly from the shop's books: why, said on the screen.
PARTIAL = {
    "upload": "Exports, advances and amendments are not kept by the software.",
    "gstr3b": "Reverse charge, imports and ITC reversals under rules 42/43 are not kept: 0 unless entered on the portal.",
    "gstr9": "Tables 8 (GSTR-2B), 10-16 need the portal / CA.",
    "gstr4": "Tables 4B, 4D, 7 and 8 are not kept by the software.",
    "composition_sale": "Composition rate is the one picked (traders: 1%).",
    "composition_purchase": "Reverse charge and import of service are not kept.",
}


def _named_builder(key: str) -> Builder:
    if key in ("julvni", "checks"):
        return julvni if key == "julvni" else checks

    def build(c: Ctx) -> dict:
        if key == "hsn_sale":
            return {"sections": [hsn_sales(c, True, False, "hsn (sales)")],
                    "notes": ["Every sale line by HSN and rate, before returns."]}
        if key == "hsn_sale_return":
            return {"sections": [hsn_sales(c, False, True, "hsn (sale returns)")],
                    "notes": ["The goods taken back on returns, by HSN and rate."]}
        if key == "hsn_sale_net":
            secs = [_legacy_table(c, "hsn(b2b)", HSN_SUMS, nums=["Rate"]),
                    _legacy_table(c, "hsn(b2c)", HSN_SUMS, nums=["Rate"]),
                    hsn_sales(c, True, True, "hsn (all, net)")]
            return {"sections": secs, "notes": ["Table 12 as filed: sales less returns, B2B and B2C apart "
                                                "(the portal asks for them separately)."]}
        purchases = key in ("hsn_purchase", "hsn_purchase_net")
        returns = key in ("hsn_purchase_return", "hsn_purchase_net")
        title = {"hsn_purchase": "hsn (purchases)", "hsn_purchase_return": "hsn (purchase returns)",
                 "hsn_purchase_net": "hsn (purchases, net of returns)"}[key]
        sec, unknown = hsn_purchases(c, purchases, returns, title)
        return {"sections": [sec], "notes": _hsn_note(unknown)}
    return build


def full_catalogue() -> list:
    """Every report the engine can make (Marg's A..W list); the screen shows SIMPLE."""
    return [{"key": k, "letter": letter, "group": g, "title": t, "form": f, "partial": PARTIAL.get(k, "")}
            for k, letter, g, t, f, _ in CATALOGUE]


# ── the screen: a short list in plain words (owner, 22 Oct 2026: "simple, not that much options") ──

def gstr1_summary(c: Ctx) -> dict:
    """GSTR-1 on one page: B2B, B2CL (when there is any), B2CS with nil, HSN, credit notes, documents."""
    secs = [_legacy_table(c, "b2b", ["Taxable Value", "Cess Amount", "IGST", "CGST", "SGST"],
                          nums=["Rate", "Invoice Value"], rename="B2B (buyers with GSTIN)")]
    if c.g1["b2cl"]:
        secs.append(_legacy_table(c, "b2cl", ["Taxable Value", "Cess Amount", "IGST"],
                                  nums=["Rate", "Invoice Value"], rename="B2CL (other state, above Rs 1 lakh)"))
    b2cs = b2cs_report(c)["sections"]
    b2cs[0]["title"] = "B2CS (other sales, by rate)"
    b2cs[1]["title"] = "Nil rated (0%)"
    secs += b2cs
    secs.append(_legacy_table(c, "hsn(b2b)", HSN_SUMS, nums=["Rate"], rename="HSN (B2B)"))
    secs.append(_legacy_table(c, "hsn(b2c)", HSN_SUMS, nums=["Rate"], rename="HSN (B2C)"))
    secs.append(_legacy_table(c, "cdnr", ["Taxable Value", "Cess Amount", "IGST", "CGST", "SGST"],
                              nums=["Rate", "Note Value"], rename="Credit notes (buyers with GSTIN)"))
    if c.g1["cdnur"]:
        secs.append(_legacy_table(c, "cdnur", ["Taxable Value", "Cess Amount", "IGST"], nums=["Rate", "Note Value"],
                                  rename="Credit notes (B2CL bills)"))
    secs.append(docs_report(c)["sections"][0])
    secs[-1]["title"] = "Documents issued"
    bad = sum(1 for t in ("hsn_b2b", "hsn_b2c") for r in c.g1[t] if r["problem"])
    notes = ([f"{bad} HSN row(s) have no HSN code: they are left out of the GSTR-1 JSON (Check bills)."]
             if bad else [])
    return {"sections": secs, "notes": notes}


def gstr3b_summary(c: Ctx) -> dict:
    out = gstr3b_report(c)
    return {"sections": out["sections"],
            "notes": ["Reverse charge, imports, interest and late fee are not kept by the software: shown as 0."]}


def hsn_summary(c: Ctx) -> dict:
    return {"sections": [hsn_sales(c, True, True, "HSN of sales (less returns)")], "notes": []}


def composition_summary(c: Ctx) -> dict:
    quarters = [_comp_row(c, k, sub) for k, sub in _quarters(c)]
    months = [_comp_row(c, a[:7], sub) for a, _, sub in _months(c)]
    sums = COMP_COLS[1:7] + COMP_COLS[8:]
    return {"sections": [section("By quarter (CMP-08)", COMP_COLS, quarters, sums=sums, nums=["Rate"]),
                         section("By month", COMP_COLS, months, sums=sums, nums=["Rate"])],
            "notes": [f"Tax at {gr._rate_out(gr._rate(c.comp_rate))}% of sales less returns (trader); "
                      "0% goods are left out."]}


def check_bills(c: Ctx) -> dict:
    return {"sections": [section("Bills to check", ["What", "Bill / Record", "Detail"],
                                 [list(x) for x in c.p.checks])], "notes": []}


# key, title, hint (the GST form, small and grey), builder, shown for: "regular" | "composition" | "both"
SIMPLE = [
    ("composition", "Composition summary", "CMP-08 / GSTR-4", composition_summary, "composition"),
    ("sale", "Sale register (GST)", "Marg upload layout", sale_register, "both"),
    ("sale_return", "Sale return register", "credit notes", sale_return_register, "both"),
    ("purchase", "Purchase register", "input tax credit", purchase_register_marg, "both"),
    ("purchase_return", "Purchase return register", "debit notes", purchase_return_register, "both"),
    ("gstr1", "GSTR-1 summary", "GSTR-1", gstr1_summary, "regular"),
    ("gstr3b", "GSTR-3B summary", "GSTR-3B", gstr3b_summary, "regular"),
    ("hsn", "HSN summary", "GSTR-1 Table 12", hsn_summary, "both"),
]
HIDDEN = {"check_bills": ("Check bills", "problems", check_bills)}


def is_composition(conn) -> bool:
    """The shop prints a Bill of Supply as a composition dealer (Settings > Bill Template)."""
    try:
        from core.bill_heading import load_bill_heading

        h = load_bill_heading(conn)
        return bool(h.get("bill_of_supply")) and bool(h.get("supply_declaration"))
    except Exception:
        return False


def catalogue(conn=None) -> list:
    comp = is_composition(conn) if conn is not None else False
    want = {"both", "composition" if comp else "regular"}
    return [{"key": k, "title": t, "hint": h} for k, t, h, _, who in SIMPLE if who in want]


def _find(key: str):
    key = str(key or "").strip()
    for row in CATALOGUE:
        if row[0] == key or (row[1] and row[1] == key.upper()):
            return row
    raise ValueError(f"GST report '{key}' nahi")


def build(conn, key: str, start: str, end: str, comp_rate: float = 1.0) -> dict:
    """One report of the screen for one period (a SIMPLE key, or any of Marg's list)."""
    simple = {k: (t, h, b) for k, t, h, b, _ in SIMPLE}
    simple.update(HIDDEN)
    if str(key) in simple:
        k = str(key)
        title, form, builder = simple[k]
        letter, group = "", ""
    else:
        k, letter, group, title, form, builder = _find(key)
        builder = builder or _named_builder(k)
    p = period(conn, start, end)
    c = Ctx(conn, p, comp_rate)
    out = builder(c)
    return {"key": k, "letter": letter, "group": group, "title": title, "form": form,
            "partial": PARTIAL.get(k, ""), "period": {"from": p.start, "to": p.end}, "shop": p.shop,
            "sections": out["sections"], "notes": out.get("notes", []),
            "counts": {"bills": len(p.bills), "credit_notes": len(p.notes), "purchases": len(p.purchases),
                       "purchase_returns": len(p.purchase_returns)},
            "checks": len(p.checks)}


# ── files: Excel / CSV in the screen's (and Marg's) layout, PDF / paper ────────────────────

def _cell(v, blank_zero: bool = False):
    if isinstance(v, float) or isinstance(v, int) and not isinstance(v, bool):
        if blank_zero and not v:
            return ""
        return v
    return "" if v is None else str(v)


def _title_lines(rep: dict, heading: str) -> list:
    shop = rep.get("shop") or {}
    return [[shop.get("name") or ""], [shop.get("gstin") or ""],
            [f"Date : {gr._dmy(rep['period']['from'])} - {gr._dmy(rep['period']['to'])}"], [heading], []]


def _igst_cols(sec: dict) -> set:
    # Marg leaves an IGST cell empty when there is none (an intra-state bill).
    return {j for j, col in enumerate(sec["columns"]) if col.endswith(" IGST") or col == "Integrated Tax Amount"} \
        if sec.get("marg") else set()


def table_lines(rep: dict) -> list:
    """The report as (kind, cells) lines, the way the CSV / Excel file lays it out: the title
    block (shop, GSTIN, "Date : from - to", the heading), then each table -- a Marg register
    with its two header rows, a blank line and the totals row under ",,,,,".
    kind: title | head | row | total | blank."""
    secs = rep["sections"]
    heading = secs[0]["title"] if len(secs) == 1 and secs[0].get("marg") else rep["title"]
    lines = [("title" if cells else "blank", cells) for cells in _title_lines(rep, heading)]
    for k, sec in enumerate(secs):
        blank = _igst_cols(sec)
        if not sec.get("marg"):
            if k:
                lines.append(("blank", []))
            lines.append(("title", [sec["title"]]))
        for h in (sec["header"] if sec.get("header") else [sec["columns"]]):
            lines.append(("head", list(h)))
        for r in sec["rows"]:
            lines.append(("row", [_cell(v, j in blank) for j, v in enumerate(r)]))
        if sec.get("total"):
            if sec.get("marg"):
                lines.append(("blank", []))
            lines.append(("total", [_cell(v, j in blank) for j, v in enumerate(sec["total"])]))
        if not sec["rows"]:
            lines.append(("row", ["(nil)"]))
    return lines


def _csv_bytes(rep: dict) -> bytes:
    import csv
    import io

    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    for _, line in table_lines(rep):
        w.writerow([f"{v:.2f}" if isinstance(v, float) else v for v in line])
    return buf.getvalue().encode("utf-8-sig")


def _xlsx_bytes(rep: dict) -> bytes:
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "".join(ch for ch in rep["title"] if ch not in "[]:*?/\\")[:31] or "GST"
    bold = Font(bold=True)
    widths: dict = {}
    for n, (kind, line) in enumerate(table_lines(rep), start=1):
        for j, v in enumerate(line, start=1):
            cell = ws.cell(row=n, column=j, value=None if v == "" else v)
            if isinstance(v, float):
                cell.number_format = "0.00"
            if kind in ("title", "head", "total"):
                cell.font = bold
            if kind in ("head", "row", "total") and v not in ("", None):
                widths[j] = min(max(widths.get(j, 8), len(str(v)) + 2), 40)
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _flat_sections(rep: dict) -> list:
    out = []
    for s in rep["sections"]:
        rows = [list(r) for r in s["rows"]]
        if s.get("total"):
            t = list(s["total"])
            if not t[0]:
                t[0] = "Total"
            rows.append(t)
        out.append({"title": s["title"], "columns": s["columns"], "rows": rows})
    return out


def file_stem(rep: dict) -> str:
    a, b = gr._dmy(rep["period"]["from"]), gr._dmy(rep["period"]["to"])
    marg = {"sale": "GST_Sale_Upload_Excel_File", "purchase": "GST_Purchase_Upload_Excel_File",
            "sale_return": "GST_Sale_Return_Upload_Excel_File",
            "purchase_return": "GST_Purchase_Return_Upload_Excel_File"}
    if rep["key"] in marg:
        return f"{marg[rep['key']]}_{a}_{b}"
    name = "".join(ch if ch.isalnum() else "_" for ch in rep["title"]).strip("_")
    while "__" in name:
        name = name.replace("__", "_")
    return f"GST_{(rep.get('letter') + '_') if rep.get('letter') else ''}{name}_{a}_{b}"


def export(conn, body: dict) -> dict:
    """The screen's buttons: Excel / CSV (the layout above), PDF / Print (the shared report
    printer), and the GSTR-1 JSON for the portal."""
    import base64

    fmt = str(body.get("format") or "").lower()
    if fmt == "json":
        return gr.export(conn, {"from": body.get("from"), "to": body.get("to"), "format": "json"})
    rep = build(conn, str(body.get("report") or ""), str(body.get("from") or ""), str(body.get("to") or ""),
                float(body.get("comp_rate") or 1))
    stem = file_stem(rep)
    print_to = str(body.get("print_to") or "").strip().lower()
    if fmt in ("csv", "xlsx", "excel") and not print_to:
        if fmt == "csv":
            data, mime, ext = _csv_bytes(rep), "text/csv", "csv"
        else:
            data, ext = _xlsx_bytes(rep), "xlsx"
            mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        return {"ok": True, "filename": f"{stem}.{ext}", "mime": mime, "format": ext,
                "content_base64": base64.b64encode(data).decode("ascii")}
    from core.desktop_export_service import export_alert_sections

    shop = rep.get("shop") or {}
    title = (f"{rep['title']} -- {shop.get('name') or ''} {('GSTIN ' + shop['gstin']) if shop.get('gstin') else ''} "
             f"{gr._dmy(rep['period']['from'])} - {gr._dmy(rep['period']['to'])}")
    res = export_alert_sections({"title": " ".join(title.split()), "sections": _flat_sections(rep),
                                 "format": fmt or "pdf", "print_to": print_to,
                                 "page_layout": body.get("page_layout") or "landscape", "paper": body.get("paper")})
    if res.get("ok") and res.get("filename"):
        res["filename"] = f"{stem}.{res.get('format') or fmt}"
    return res
