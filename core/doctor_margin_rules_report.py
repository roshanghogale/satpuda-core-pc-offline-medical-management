"""Doctor Margin Report by the margin rules (owner, 22 + 23 Oct 2026) -- the screen, the files,
the pick lists of the rules screen.

Reads the doctor report's saved lines (core/doctor_margin_store.py) after bringing the period
up to date with its bills; the rules are core/doctor_margin_rules.py. Per doctor, two parts:

  On profit       only the products picked into a group. Layout 'doctor' (the owner's Excel:
                  product / batch / quantity / MRP / rate / margin per unit / total value and a
                  TOTAL) or 'detail' (customer-wise product-batch-wise, with sub-totals).
  On total sales  the types set to "On total sales", laid out like the owner's other software
                  ("Customer-Wise Product-Batch Wise Sales From .. To .."): a heading per product
                  (name + pack), one row per bill line -- customer, bill no, date, batch, qty,
                  rate, discount, value, doctor, type -- a Sub Total per product, then the total
                  sale value and the commission (X% of it).
  Doctor's margin = the on-profit TOTAL + the total-sales commission.
"""
from __future__ import annotations

import re
from datetime import date, datetime
from typing import Any

from core import doctor_margin_rules as dmr
from core import doctor_margin_store as store

VIEWS = ("doctor", "detail")
_DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")


from core.doctor_margin_report import ReportError  # noqa: E402  (the route answers it with HTTP 400)


def _dmy(iso: str, sep: str = "-") -> str:
    s = str(iso or "")[:10]
    return f"{s[8:10]}{sep}{s[5:7]}{sep}{s[0:4]}" if len(s) == 10 else s


def parse_filters(body: dict[str, Any]) -> dict[str, Any]:
    today = date.today().isoformat()
    fd = str(body.get("from") or body.get("from_date") or today).strip()[:10]
    td = str(body.get("to") or body.get("to_date") or today).strip()[:10]
    for label, v in (("From", fd), ("To", td)):
        if not _DATE_RE.match(v):
            raise ReportError(f"{label} date is not a date (YYYY-MM-DD).")
        try:
            date.fromisoformat(v)
        except ValueError:
            raise ReportError(f"{label} date is not a date (YYYY-MM-DD).") from None
    if fd > td:
        raise ReportError("From date is after To date.")
    raw = body.get("doctors")
    if raw is None and body.get("doctor"):
        raw = [body.get("doctor")]
    if isinstance(raw, str):
        raw = raw.split(",")
    doctors = sorted({dmr.norm_name(d) for d in (raw or []) if dmr.norm_name(d)})
    basis = str(body.get("rate_basis") or "").strip().lower()
    view = str(body.get("view") or "doctor").strip().lower()
    return {
        "from": fd, "to": td, "doctors": doctors,
        "rate_basis": basis if basis in dmr.RATE_BASES else "",
        "view": view if view in VIEWS else "doctor",
    }


def _is_online() -> bool:
    try:
        from core.sync_prefs import is_online_mode

        return bool(is_online_mode())
    except Exception:
        return False


def _flag_counts(rows: list[dict[str, Any]]) -> dict[str, int]:
    out = {k: 0 for k in dmr.FLAG_TEXT}
    for r in rows:
        for f in r.get("flags") or []:
            if f in out:
                out[f] += 1
    return out


def by_product(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """The on-profit detail layout: per product (name), its lines by batch / date, and a sub-total."""
    out: dict[str, dict[str, Any]] = {}
    for r in sorted(rows, key=lambda r: (r["name"], r["batch"], r["date"], r["bill_no"])):
        g = out.get(r["name"])
        if g is None:
            g = out[r["name"]] = {"name": r["name"], "type": r["type"], "lines": [], "qty": 0.0, "value": 0.0}
        g["lines"].append(r)
        g["qty"] = round(g["qty"] + r["qty"], 3)
        g["value"] = dmr.r2(g["value"] + r["value"])
    return list(out.values())


def build_report(conn, body: dict[str, Any]) -> dict[str, Any]:
    if _is_online():
        raise ReportError(
            "The Doctor Margin Report reads this PC's own books; an Online store's bills are "
            "on the server (Sales History on the web)."
        )
    f = parse_filters(body)
    if body.get("recalculate"):
        done = store.recalculate(conn, f["from"], f["to"])
        sync = {"recalculated": True, **done}
    else:
        sync = store.reconcile(conn, f["from"], f["to"])
    lines = store.read_lines(conn, f["from"], f["to"], f["doctors"])
    if f["rate_basis"]:
        lines = [dmr.with_basis(r, f["rate_basis"]) for r in lines]
    rep = dmr.assemble(lines)
    rules = dmr.load_rules(conn)
    return {
        "ok": True,
        "mode": "rules",
        "filters": f,
        "has_rules": dmr.has_rules(rules),
        "rules_rate_basis": rules["rate_basis"],
        "rate_basis_text": dmr.RATE_BASIS_TEXT.get(f["rate_basis"], "As saved with each bill (rules)"),
        "flag_counts": _flag_counts(rep["rows"]),
        "products": by_product(rep["rows"]),
        "sync": sync,
        **rep,
    }


def has_lines(rep: dict[str, Any]) -> bool:
    return bool(rep["rows"] or rep["sales"]["rows"])


# ── the rules screen's pick lists ───────────────────────────────────────────────────────────

def rule_products(conn, med_type: str = "", find: str = "") -> dict[str, Any]:
    """Products of a type for the pick list: name, stock (sale units), pack, company, contents,
    and its batches (batch, stock, expiry, MRP). Hidden / deleted medicine rows left out."""
    have = store._cols(conn, "medicines")
    where = ["TRIM(COALESCE(m.name,''))<>''"]
    args: list[Any] = []
    if "deleted" in have:
        where.append("COALESCE(m.deleted,0)=0")
    if med_type:
        where.append("LOWER(TRIM(COALESCE(m.type,'')))=?")
        args.append(dmr.norm_type(med_type))
    if find:
        where.append("UPPER(m.name) LIKE ?")
        args.append(f"%{find.strip().upper()}%")
    company = "m.manufacturer" if "manufacturer" in have else "''"
    contents = "m.content_drug" if "content_drug" in have else "''"
    products: dict[str, dict[str, Any]] = {}
    for name, mtype, batch, stock, unit, comp, cont, expiry, mrp in conn.execute(
        f"""
        SELECT m.name, COALESCE(m.type,''), COALESCE(m.batch_no,''), COALESCE(m.stock_qty,0),
               COALESCE(m.unit,''), COALESCE({company},''), COALESCE({contents},''),
               COALESCE(m.expiry_date,''), COALESCE(m.mrp,0)
        FROM medicines m WHERE {' AND '.join(where)}
        ORDER BY UPPER(m.name), m.batch_no
        """,
        args,
    ):
        key = dmr.norm_name(name)
        p = products.get(key)
        if p is None:
            p = products[key] = {"product": key, "type": str(mtype), "stock": 0.0, "pack": str(unit or ""),
                                 "company": str(comp or ""), "contents": str(cont or ""), "batches": []}
        p["stock"] += float(stock or 0)
        b = dmr.norm_batch(batch)
        if b and all(x["batch"] != b for x in p["batches"]):
            p["batches"].append({"batch": b, "stock": float(stock or 0), "expiry": str(expiry or "")[:10],
                                 "mrp": float(mrp or 0)})
        elif b:
            for x in p["batches"]:
                if x["batch"] == b:
                    x["stock"] += float(stock or 0)
        if not p["company"] and comp:
            p["company"] = str(comp)
        if not p["contents"] and cont:
            p["contents"] = str(cont)
    return {"ok": True, "type": med_type, "products": list(products.values())}


def rules_payload(conn) -> dict[str, Any]:
    from core.doctor_margin_report import report_options

    opts = report_options(conn)
    return {
        "ok": True,
        "rules": dmr.load_rules(conn),
        "doctors": [d for d in opts["doctors"] if d != dmr.NO_DOCTOR],
        "types": opts["types"],
        "online": opts["online"],
    }


# ── files ──────────────────────────────────────────────────────────────────────────────────

DOCTOR_COLS = ["Product name", "Batch", "Quantity", "MRP", "Rate", "Margin / unit", "Total value"]
DETAIL_COLS = ["Product / Customer", "Bill no", "Date", "Batch", "Qty", "MRP", "Rate", "Margin / unit",
               "Value", "Doctor", "Type"]
SALES_COLS = ["Customer Name", "Bill No.", "Date", "Batch No.", "Qty", "Rate", "Disc.", "Value", "Doctor Name",
              "Type"]
SUMMARY_COLS = ["Doctor", "On profit", "Total sales", "Commission", "Doctor's margin"]
SALES_TITLE = "Customer-Wise Product-Batch Wise Sales"


def _q(v: float) -> Any:
    return int(v) if float(v) == int(v) else v


def _pct(p: float) -> str:
    return f"{_q(round(float(p), 4))}%"


def profit_table(rep: dict[str, Any], doctor: str) -> tuple[list[str], list[tuple[str, list[Any]]]]:
    """(columns, [(kind, cells)]) of one doctor's on-profit lines. kind: head | line | sub | total."""
    out: list[tuple[str, list[Any]]] = []
    mine = [r for r in rep["rows"] if r["doctor"] == doctor]
    d = next((x for x in rep["doctors"] if x["doctor"] == doctor), None)
    if not mine or d is None:
        return DOCTOR_COLS, out
    if rep["filters"]["view"] == "detail":
        for p in by_product(mine):
            out.append(("head", [p["name"]] + [""] * (len(DETAIL_COLS) - 1)))
            for r in p["lines"]:
                out.append(("line", [r["patient"], r["bill_no"], _dmy(r["date"], "/"), r["batch"], _q(r["qty"]),
                                     r["mrp"], r["rate"], r["margin_unit"], r["value"], r["doctor"], r["type"]]))
            out.append(("sub", ["Sub total", "", "", "", _q(p["qty"]), "", "", "", p["value"], "", ""]))
        out.append(("total", ["TOTAL value (on profit)", "", "", "", _q(d["qty"]), "", "", "", d["value"], "", ""]))
        return DETAIL_COLS, out
    for r in rep["summary"]:
        if r["doctor"] == doctor:
            out.append(("line", [r["name"], r["batch"], _q(r["qty"]), r["mrp"], r["rate"], r["margin_unit"],
                                 r["value"]]))
    out.append(("total", ["", "", _q(d["qty"]), "", "", "TOTAL value", d["value"]]))
    return DOCTOR_COLS, out


def sales_table(rep: dict[str, Any], doctor: str) -> tuple[list[str], list[tuple[str, list[Any]]]]:
    """(columns, [(kind, cells)]) of one doctor's total-sales lines, the other software's layout."""
    out: list[tuple[str, list[Any]]] = []
    d = next((x for x in rep["sales"]["doctors"] if x["doctor"] == doctor), None)
    if d is None:
        return SALES_COLS, out
    blank = [""] * len(SALES_COLS)
    for p in d["products"]:
        out.append(("head", [" ".join(x for x in (p["name"], p["pack"]) if x)] + blank[1:]))
        for r in p["lines"]:
            out.append(("line", [r["patient"], r["bill_no"], _dmy(r["date"], "/"), r["batch"], _q(r["qty"]),
                                 r["sale_rate"], r["disc"] if r["disc"] else "-", r["value"], r["doctor"],
                                 str(r["type"]).upper()]))
        out.append(("sub", ["", "", "", "Sub Total:", _q(p["qty"]), "", "", p["value"], "", ""]))
    types = ", ".join(str(t).upper() for t in d["types"])
    out.append(("total", [f"Total sales value ({types})", "", "", "", _q(d["qty"]), "", "", d["value"], "", ""]))
    for b in d["by_pct"]:
        out.append(("total", [f"Commission {_pct(b['pct'])} of {b['value']:.2f}", "", "", "", "", "", "",
                              b["commission"], "", ""]))
    return SALES_COLS, out


def blocks(rep: dict[str, Any]) -> list[tuple]:
    """The report as stacked blocks, every file is made from these:
         ('doctor', name)
         ('table', title, columns, [(kind, cells)])
         ('margin', label, value)       the doctor's margin, profit + commission"""
    f = rep["filters"]
    out: list[tuple] = []
    period = f"From {_dmy(f['from'], '/')} To {_dmy(f['to'], '/')}"
    for g in rep["grand"]["doctors"]:
        doc = g["doctor"]
        out.append(("doctor", doc))
        cols, rows = profit_table(rep, doc)
        if rows:
            out.append(("table", "On profit (products in the groups)", cols, rows))
        cols, rows = sales_table(rep, doc)
        if rows:
            out.append(("table", f"On total sales - {SALES_TITLE} {period}", cols, rows))
        parts = []
        if rep["doctors"] and any(d["doctor"] == doc for d in rep["doctors"]):
            parts.append(f"on profit {g['profit']:.2f}")
        if any(d["doctor"] == doc for d in rep["sales"]["doctors"]):
            parts.append(f"commission {g['commission']:.2f}")
        out.append(("margin", f"{doc} - doctor's margin ({' + '.join(parts)})", g["total"]))
    if len(rep["grand"]["doctors"]) > 1:
        rows = [("line", [g["doctor"], g["profit"], g["sales_value"], g["commission"], g["total"]])
                for g in rep["grand"]["doctors"]]
        t = rep["grand"]["totals"]
        rows.append(("total", ["ALL DOCTORS", t["profit"], t["sales_value"], t["commission"], t["total"]]))
        out.append(("table", "All doctors", SUMMARY_COLS, rows))
    return out


def table(rep: dict[str, Any]) -> tuple[list[str], list[tuple[str, list[Any]]]]:
    """The on-profit part as ONE table (all doctors), kept for older callers."""
    rows: list[tuple[str, list[Any]]] = []
    cols = DETAIL_COLS if rep["filters"]["view"] == "detail" else DOCTOR_COLS
    for d in rep["doctors"]:
        if rep["filters"]["view"] != "detail":
            rows.append(("head", [d["doctor"]] + [""] * (len(cols) - 1)))
        rows.extend(profit_table(rep, d["doctor"])[1])
    return cols, rows


def header_lines(rep: dict[str, Any]) -> list[tuple[str, str]]:
    f = rep["filters"]
    return [
        ("Period", f"{_dmy(f['from'])} to {_dmy(f['to'])}"),
        ("Doctor", ", ".join(f["doctors"]) or "All doctors"),
        ("Rate (on profit)", rep["rate_basis_text"]),
        ("Layout (on profit)", "Doctor-wise (product, qty, MRP, rate, total value)" if f["view"] == "doctor"
         else "Customer-wise product-batch-wise"),
        ("Generated at", datetime.now().strftime("%d/%m/%Y %H:%M")),
    ]


def _shop(conn=None) -> dict[str, str]:
    try:
        from core.pharmacy_profile_io import load_pharmacy_profile

        p = load_pharmacy_profile(conn) or {}
        return {"name": str(p.get("name") or ""), "address": " ".join(str(p.get("address") or "").split())}
    except Exception:
        return {"name": "", "address": ""}


def _shop_name(conn=None) -> str:
    return _shop(conn)["name"]


def _xlsx(title: str, heads, bl: list[tuple]) -> bytes:
    import io

    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Doctor Margin"
    bold = Font(bold=True)
    ws.append([title])
    ws.cell(row=1, column=1).font = Font(bold=True, size=13)
    for k, v in heads:
        ws.append([k, v])
        ws.cell(row=ws.max_row, column=1).font = bold
    band = PatternFill("solid", fgColor="1F4E79")
    line = Side(style="thin", color="999999")
    widths: dict[int, int] = {}
    for b in bl:
        if b[0] == "doctor":
            ws.append([])
            ws.append([b[1]])
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=12, color="1F4E79")
            continue
        if b[0] == "margin":
            ws.append([b[1], b[2]])
            for c in ws[ws.max_row]:
                c.font = Font(bold=True, size=11)
            continue
        _, sub, cols, rows = b
        ws.append([sub])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, italic=True)
        ws.append(cols)
        for c in ws[ws.max_row]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = band
        for i, h in enumerate(cols, 1):
            widths[i] = max(widths.get(i, 0), len(str(h)))
        for kind, cells in rows:
            ws.append(list(cells))
            r = ws[ws.max_row]
            if kind == "head":
                r[0].font = Font(bold=True, size=11, color="1F4E79")
            elif kind in ("sub", "total"):
                for c in r:
                    c.font = bold
                    c.border = Border(top=line)
            if kind != "head":
                for i, v in enumerate(cells, 1):
                    widths[i] = max(widths.get(i, 0), len(str(v)))
    for row in ws.iter_rows(min_row=len(heads) + 2):
        for c in row:
            if isinstance(c.value, float):
                c.number_format = "0.00"
                c.alignment = Alignment(horizontal="right")
    for i, w in widths.items():
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = min(42, max(9, w + 2))
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _csv(title: str, heads, bl: list[tuple]) -> bytes:
    import csv
    import io

    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([title])
    for k, v in heads:
        w.writerow([k, v])
    for b in bl:
        if b[0] == "doctor":
            w.writerow([])
            w.writerow([b[1]])
        elif b[0] == "margin":
            w.writerow([b[1], f"{b[2]:.2f}"])
        else:
            w.writerow([b[1]])
            w.writerow(b[2])
            for _, cells in b[3]:
                w.writerow([f"{v:.2f}" if isinstance(v, float) else v for v in cells])
    return buf.getvalue().encode("utf-8-sig")


def _css_str(v: str) -> str:
    return '"' + str(v).replace("\\", "\\\\").replace('"', '\\"').replace("\n", " ") + '"'


def report_html(title: str, heads, bl: list[tuple], shop: dict[str, str] | None = None,
                period: str = "") -> str:
    """A4 portrait, small type -- close to the other software's sales report: the shop's name
    and address, Printed on and Page n of N on every page (Edge / Chrome print the @page
    boxes), then per doctor his two parts and his margin."""
    shop = shop or {"name": "", "address": ""}

    def esc(v: Any) -> str:
        return str(v).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    def cell(v: Any) -> str:
        if isinstance(v, float):
            return f"<td class=num>{v:.2f}</td>"
        if isinstance(v, int) and not isinstance(v, bool) or v == "-":
            return f"<td class=num>{v}</td>"
        return f"<td>{esc(v)}</td>"

    def numeric(rows, i: int) -> bool:
        """A column of figures: its heading goes to the right like the figures."""
        vals = [c[i] for k, c in rows if k == "line" and i < len(c) and c[i] not in ("", None, "-")]
        return bool(vals) and all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in vals)

    body: list[str] = []
    for b in bl:
        if b[0] == "doctor":
            body.append(f"<h3 class=doc>{esc(b[1])}</h3>")
            continue
        if b[0] == "margin":
            body.append(f"<div class=margin><span>{esc(b[1])}</span><b>{b[2]:.2f}</b></div>")
            continue
        _, sub, cols, rows = b
        trs = []
        for kind, cells in rows:
            if kind == "head":
                trs.append(f"<tr class=head><td colspan={len(cols)}>{esc(cells[0])}</td></tr>")
            else:
                trs.append(f"<tr class={kind}>" + "".join(cell(v) for v in cells) + "</tr>")
        hdr = "".join(f"<th class=num>{esc(h)}</th>" if numeric(rows, i) else f"<th>{esc(h)}</th>"
                      for i, h in enumerate(cols))
        body.append(f"<div class=sub>{esc(sub)}</div><table class=grid><thead><tr>{hdr}</tr></thead>"
                    f"<tbody>{''.join(trs)}</tbody></table>")
    meta = "".join(f"<tr><th>{esc(k)}</th><td>{esc(v)}</td></tr>" for k, v in heads)
    printed = datetime.now().strftime("%d/%m/%Y")
    css = (
        "@page{size:A4 portrait;margin:14mm 9mm 12mm 9mm;"
        f"@top-left{{content:{_css_str(shop['name'])};font:600 8pt Arial;color:#333}}"
        f"@top-right{{content:{_css_str('Printed on : ' + printed)};font:8pt Arial;color:#333}}"
        "@bottom-right{content:'Page ' counter(page) ' of ' counter(pages);font:8pt Arial;color:#333}}"
        "@page :first{@top-left{content:none}@top-right{content:none}}"  # page 1 has the big heading
        "body{font-family:Arial,Segoe UI,sans-serif;font-size:7.5pt;color:#000;margin:0}"
        ".top{display:flex;justify-content:space-between;align-items:flex-start;margin-bottom:4px}"
        ".shop{font-size:13pt;font-weight:600}.addr{font-size:7.5pt;color:#555;margin-top:1px}"
        ".printed{font-size:7.5pt;text-align:right}"
        "h2{font-size:9.5pt;font-weight:600;margin:6px 0 3px;color:#1f3b57}"
        "table{width:100%;border-collapse:collapse}"
        "table.meta{width:auto;margin-bottom:4px}table.meta th{text-align:left;padding:0 10px 0 0;font-weight:600}"
        "table.meta td{padding:0}"
        "h3.doc{font-size:10pt;margin:12px 0 2px;padding:3px 5px;background:#e8eef4;border-left:3px solid #1f4e79;"
        "break-after:avoid}"
        ".sub{font-weight:600;font-size:8pt;margin:6px 0 2px;color:#1f3b57;break-after:avoid}"
        "table.grid th{border-top:1px solid #444;border-bottom:1px solid #444;padding:3px;text-align:left;"
        "font-weight:600;background:#f2f2f2}"
        "table.grid td{padding:2px 3px;border-bottom:.5px solid #e2e2e2}"
        "td.num{text-align:right;font-variant-numeric:tabular-nums;white-space:nowrap}table.grid th.num{text-align:right}"
        "thead{display:table-header-group}tr{break-inside:avoid}"
        "tr.head td{font-weight:600;font-size:8pt;padding-top:7px;background:#f7f7f7;border-bottom:1px solid #999}"
        "tr.sub td{font-weight:600;border-top:1px solid #777;border-bottom:1px solid #777}"
        "tr.total td{font-weight:700;border-top:1px solid #333}"
        ".margin{display:flex;justify-content:space-between;margin:5px 0 4px;padding:4px 6px;"
        "border-top:1.5px solid #000;border-bottom:1.5px solid #000;font-size:9pt;font-weight:600}"
    )
    top = (f"<div class=top><div><div class=shop>{esc(shop['name'])}</div>"
           f"<div class=addr>{esc(shop['address'])}</div></div>"
           f"<div class=printed>Printed on : {printed}</div></div>") if shop["name"] else ""
    return (f"<!DOCTYPE html><html><head><meta charset=UTF-8><title>{esc(title)}</title><style>{css}</style>"
            f"</head><body>{top}<h2>{esc(title)}{(' ' + esc(period)) if period else ''}</h2>"
            f"<table class=meta>{meta}</table>{''.join(body)}</body></html>")


def _dot_matrix_sections(title: str, bl: list[tuple]) -> list[tuple[str, list, list]]:
    out: list[tuple[str, list, list]] = []
    doc = ""
    for b in bl:
        if b[0] == "doctor":
            doc = b[1]
        elif b[0] == "table":
            out.append((f"{title} | {doc} | {b[1]}" if doc and b[1] != "All doctors" else f"{title} | {b[1]}",
                        b[2], [c for _, c in b[3]]))
        elif b[0] == "margin" and out:
            out[-1][2].append([b[1]] + [""] * (len(out[-1][1]) - 2) + [b[2]])
    return out


def export_report(conn, body: dict[str, Any]) -> dict[str, Any]:
    import base64

    from core.doctor_margin_report import _pdf_bytes

    rep = build_report(conn, body)
    if not has_lines(rep):
        return {"ok": False, "error": "No doctor margin in this period for these doctors."}
    f = rep["filters"]
    title = "Doctor Margin Report"
    period = f"From {_dmy(f['from'], '/')} To {_dmy(f['to'], '/')}"
    heads = header_lines(rep)
    bl = blocks(rep)
    base = f"doctor_margin_{f['view']}_{f['from']}_{f['to']}"
    fmt = str(body.get("format") or "xlsx").strip().lower()
    if fmt == "excel":
        fmt = "xlsx"
    if fmt == "print":
        target = str(body.get("print_to") or "").strip().lower()
        if target not in ("dot_matrix", "printer"):
            try:
                from core.printer_manager import PrinterManager

                target = "dot_matrix" if PrinterManager.is_dot_matrix_mode() else "printer"
            except Exception:
                target = "printer"
        if target == "dot_matrix":
            from core.dot_matrix_print import print_dot_matrix_reports_combined

            try:
                print_dot_matrix_reports_combined(_dot_matrix_sections(f"{title} {period}", bl), paper="A4",
                                                  landscape=True)
            except Exception as exc:
                return {"ok": False, "printed": False, "print_error": str(exc),
                        "error": f"Dot matrix var print zala nahi: {exc}"}
            return {"ok": True, "printed": True, "message": "Sent to the dot matrix."}
        from core import saved_files
        from core.desktop_export_service import print_pdf_on_printer

        payload, ext = _pdf_bytes(report_html(title, heads, bl, _shop(conn), period))
        if ext != "pdf":
            return {"ok": False, "printed": False, "error": "PDF could not be made for printing. Install Edge or Chrome."}
        path = saved_files.save_bytes(f"{base}.pdf", payload)
        problem = print_pdf_on_printer(path, landscape=False)
        if problem:
            return {"ok": False, "printed": False, "print_error": problem, "error": problem, "path": path}
        return {"ok": True, "printed": True, "path": path, "message": "Sent to the printer."}
    if fmt == "csv":
        payload, mime = _csv(f"{title} {period}", heads, bl), "text/csv;charset=utf-8"
    elif fmt == "xlsx":
        try:
            payload = _xlsx(f"{title} {period}", heads, bl)
        except ImportError:
            return {"ok": False, "error": "openpyxl not installed for Excel export."}
        mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    elif fmt == "pdf":
        payload, fmt = _pdf_bytes(report_html(title, heads, bl, _shop(conn), period))
        mime = "application/pdf" if fmt == "pdf" else "text/html"
    elif fmt == "html":  # the PDF's page itself (tests, a look before printing)
        payload, mime = report_html(title, heads, bl, _shop(conn), period).encode("utf-8"), "text/html"
    else:
        return {"ok": False, "error": f"Unknown export format: {fmt}"}
    return {"ok": True, "filename": f"{base}.{fmt}", "mime": mime, "format": fmt,
            "content_base64": base64.b64encode(payload).decode("ascii"),
            "row_count": len(rep["rows"]) + len(rep["sales"]["rows"]),
            "totals": rep["totals"], "grand": rep["grand"]["totals"]}
